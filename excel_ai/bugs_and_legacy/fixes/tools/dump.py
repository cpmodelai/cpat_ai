import openpyxl, sys, pickle
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=False)
ws = wb['Mitigation']
rows = {}
for i,row in enumerate(ws.iter_rows(min_row=1, max_row=15000, max_col=32, values_only=True), start=1):
    rows[i] = row
pickle.dump(rows, open(sys.argv[2],'wb'))
wb2 = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
ws2 = wb2['Mitigation']
vals={}
for i,row in enumerate(ws2.iter_rows(min_row=1, max_row=15000, max_col=32, values_only=True), start=1):
    vals[i]=row
pickle.dump(vals, open(sys.argv[3],'wb'))
