"""v0.7.2 (codes plan R0-R9): structural check of the code renames, no recalculation (about a minute).

    python check_codes_v0_7_2.py [new.xlsx|xlsm]     (default: the quick build CPAT-AI-Mitigation-Prototype-v0.7.2_quick.xlsx,
                                                    else the full build)

1. Cell by cell against the last release (Old/ or this folder): every cell equals the release's cell after the rename
   map (code_renames_v0_7_2.csv) is applied to its text; texts that name the version are compared with the version
   numbers neutralised; the version log is not compared (nor, for a quick build, the sheets the batch fills:
   StoredOutputs, BaselineCO2).
2. Codes: every Mitigation code (code variable, code slots, row codes; the map) is unique in a scenario group (R7: the
   rows repeating an earlier row's variable, sector and fuel, the section 11 copies, carry no code).
3. Every code stem the output sheets read (LiveOutputs column B) exists on Mitigation.
4. With --values (about 3 minutes): the new file recalculated in LibreOffice (LAMBDAs expanded, two calculateAll) has
   the release's values in every numeric cell of Mitigation and LiveOutputs (the renames change text only), and the
   codes it computes (code variable, code variable VCode, pair slots MSlot, row codes RowCode) are those of check 2.
"""
import os
import re
import sys

import openpyxl

import build as B

HERE = os.path.dirname(os.path.abspath(__file__))
PREV = '7.1'
OLD_CODE = re.compile(r'^=CodePrefix&\$A(\d+)&(IF\(.*\))&INDEX\(VarEnd,MATCH\(\$A\d+,VarKey,0\)\)&(.+)$')
VAR_SHEETS = {'Variables', 'Check, Variables, Lambdas'}
COPY_ROWS = set()                # the release's rows repeating an earlier row's code (filled in main)
VERSION_TXT = re.compile(r'(v?7[._][12])(?!\d)')


def prev_path():
    for d in (os.path.join(HERE, 'Old'), HERE):
        p = os.path.join(d, f'{B.NAME}-v{PREV}.xlsm')
        if os.path.exists(p):
            return p
    raise SystemExit(f'last release {PREV} not found')


def neutral(v):
    return VERSION_TXT.sub('<v>', v) if isinstance(v, str) else v


def renamed_text(sheet, r, c, v, var_sfx_rows):
    """The release's cell text as the rename map writes it."""
    if sheet in VAR_SHEETS and c == 4 and r in var_sfx_rows:   # the code suffix of a Variables row
        return B.rn_sfx(var_sfx_rows[r], v or '') or None
    if not isinstance(v, str):
        return v
    if sheet == 'Mitigation' and v.startswith('=CodePrefix&$A') and r in COPY_ROWS:   # R7: copies drop their code
        return None
    if sheet == 'Mitigation' and v.startswith('=CodePrefix&$A'):    # v0.7.2: code formula reads the code variable
        m = OLD_CODE.match(v)
        return B.code_formula(int(m.group(1)), re.match(r'([A-Z]+)\$', m.group(3)).group(1),
                              var_sfx_rows.get(r)) if m else v
    if sheet == 'Mitigation' and c == 1 and r in var_sfx_rows:   # columns B and C stay internal (v0.7.2)
        return B.rn_var(var_sfx_rows[r][0])
    if sheet in VAR_SHEETS and c == 1 and v in B.VAR_SFX:      # the Variables key
        return B.rn_var(v)
    if v.startswith('mit.'):
        return B.rn_stem(v)
    return v


def RC(code):
    """An output code country.mit.<stem>.<scenario> of the release through the map."""
    p = (code or '').split('.')
    return '.'.join([p[0], B.rn_stem('.'.join(p[1:-1])), p[-1]]) if len(p) > 3 else code


def codes(ws, wv, release=False):
    """{row: code without country and scenario} of every calculated Mitigation row (Python, as code_formula; the
    release: column A and its suffix; v0.7.2: the code variable of Variables column L, no suffix)."""
    if release:
        sfx = {wv.cell(r, 1).value: (wv.cell(r, 4).value or '') for r in range(B.VR0, wv.max_row + 1)
               if wv.cell(r, 1).value}
        tok = {k: k for k in sfx}
    else:
        tok = {wv.cell(r, 1).value: wv.cell(r, 12).value for r in range(B.VR0, wv.max_row + 1) if wv.cell(r, 1).value}
        sfx = {k: '' for k in tok}
    out = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a in tok and isinstance(ws.cell(r, B.group_cols(1)[0]).value, str):
            c_, b_ = ws.cell(r, 3).value or '', ws.cell(r, 2).value or ''
            if not release:                        # the code body (code variable VCode, pair slots MSlot, row codes RowCode)
                out[r] = 'mit.' + B.rn_body(B.RN_VAR_INV.get(a, a), c_, b_)
                continue
            out[r] = 'mit.' + tok[a] + (f'.{c_}.{b_}' if c_ else '') + (f'.{sfx[a]}' if sfx[a] else '')
    return out


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    new = args[0] if args else next(
        p for p in (B.OUT.replace('.xlsm', '_quick.xlsx'), B.OUT) if os.path.exists(p))
    wn, wo = openpyxl.load_workbook(new), openpyxl.load_workbook(prev_path())
    ok = True
    # 1. cell by cell
    mo = wo['Mitigation']
    COPY_ROWS.update(B.copy_rows(mo, [B.group_cols(g)[0] for g, *_r in B.SCENARIOS]))
    rows = {r: (mo.cell(r, 1).value, mo.cell(r, 3).value or '', mo.cell(r, 2).value or '')
            for r in range(1, mo.max_row + 1) if mo.cell(r, 1).value in B.VAR_SFX}
    wvo = B.tab(wo, 'Variables')
    vrows = {r: wvo.cell(r, 1).value for r in range(1, wvo.max_row + 1) if wvo.cell(r, 1).value in B.VAR_SFX}
    skip = {'VersionLog', 'Version Log, About, and Diff'}
    if new.endswith('_quick.xlsx'):                # the quick build has no batch: stored results and baseline CO2 empty
        skip |= {'StoredOutputs', 'BaselineCO2'}
    so = B.tab(wo, 'Mapping')                     # v0.7.2: code-slot columns G, I, J beside the member labels (new cells)
    mem0 = next(r for r in range(1, so.max_row + 1)
                if isinstance(so.cell(r, 1).value, str) and so.cell(r, 1).value.startswith('Member labels'))
    alias_col = B.MTO_C0 + len([B.BASE_YEAR] + B.YEARS) + 2
    lo_n = wn['LiveOutputs']                      # R8: the alias column holds the v0.7.1 stem of every renamed stem
    bad_alias = [r for r in range(1, lo_n.max_row + 1) if isinstance(lo_n.cell(r, 2).value, str)
                 and lo_n.cell(r, 2).value.startswith('mit.')
                 and lo_n.cell(r, alias_col).value != B.MTO_ALIAS.get(lo_n.cell(r, 2).value)]
    bad_alias += [] if lo_n.cell(4, alias_col).value == B.ALIAS_HEAD else ['header']
    diffs, n = [], 0
    so_rows = 0
    if 'StoredOutputs' not in skip:               # sorted by code stem: the renamed stems reorder the rows, so the
        def stored(ws, ren):                      # rows are compared by (scenario ID, stem)
            out = {}
            for r in range(B.SR_FIRST, ws.max_row + 1):
                v = [ws.cell(r, c).value for c in range(1, B.SR_Y0 + B.NY + 1)]
                if v[2]:
                    if ren:
                        v[2], v[3] = B.rn_stem(v[2]), RC(v[3])
                    out[(v[0], v[2])] = v
            return out
        sto, stn = stored(wo['StoredOutputs'], True), stored(wn['StoredOutputs'], False)
        so_rows = len(sto)
        diffs += [f'StoredOutputs {k}' for k in sorted(set(sto) | set(stn), key=str) if sto.get(k) != stn.get(k)][:20]
        skip = skip | {'StoredOutputs'}
    for ws_o in wo.worksheets:
        if ws_o.title in skip or ws_o.title not in wn.sheetnames:
            continue
        ws_n = wn[ws_o.title]
        for row in ws_o.iter_rows():
            for c in row:
                if c.value is None and ws_n.cell(c.row, c.column).value is None:
                    continue
                n += 1
                ctx = rows if ws_o.title == 'Mitigation' else vrows if ws_o.title in VAR_SHEETS else {}
                exp = renamed_text(ws_o.title, c.row, c.column, c.value, ctx)
                got = ws_n.cell(c.row, c.column).value
                if ws_o.title == so.title and c.value is None and c.row > mem0 and c.column in (7, 9, 10):
                    continue                       # the code tables themselves: checked by the codes the workbook computes (4)
                if ws_o.title == 'LiveOutputs' and c.column == alias_col:   # R8: the alias column (new cells)
                    continue
                if neutral(exp) != neutral(got):
                    diffs.append(f'{ws_o.title}!{c.coordinate}: {str(exp)[:50]!r} -> {str(got)[:50]!r}')
    extra = [s for s in wn.sheetnames if s not in wo.sheetnames] + [s for s in wo.sheetnames if s not in wn.sheetnames]
    good = not diffs and not extra
    ok &= good
    print(f'- cells against v{PREV} through the rename map: {n} compared (StoredOutputs: {so_rows} rows by ID and stem), '
          f'{len(diffs)} unexplained differences '
          f'{diffs[:8]}; sheets added/removed {extra} {"OK" if good else "FAIL"}')
    good = not bad_alias
    ok &= good
    print(f'- LiveOutputs alias column {B.L(alias_col)}: the v0.7.1 stem of each of the {len(B.MTO_ALIAS)} renamed stems '
          f'{"OK" if good else "FAIL rows " + str(bad_alias[:8])}')
    # 2. unique codes
    wv_n, wv_o = B.tab(wn, 'Variables'), B.tab(wo, 'Variables')
    cn, co = codes(wn['Mitigation'], wv_n), codes(mo, wv_o, release=True)
    dup = lambda cd: {k for k in cd.values() if list(cd.values()).count(k) > 1}
    dn, do = dup(cn), {B.rn_stem(k) for k in dup(co)}
    good = not dn
    ok &= good
    print(f'- Mitigation codes: {len(cn)} rows, {len(set(cn.values()))} distinct; duplicates {len(dn)} (release: '
          f'{len(do)}; R7: the {len(COPY_ROWS)} copies drop their code) {"OK" if good else "FAIL " + str(sorted(dn)[:8])}')
    # 3. output stems exist
    lo = wn['LiveOutputs']
    stems = {lo.cell(r, 2).value for r in range(B.MTO_R0, lo.max_row + 1)
             if isinstance(lo.cell(r, 2).value, str) and lo.cell(r, 2).value.startswith('mit.')}
    have = set(cn.values())
    miss = sorted(s for s in stems if s not in have)
    good = not miss
    ok &= good
    print(f'- LiveOutputs code stems found on Mitigation: {len(stems) - len(miss)} of {len(stems)} '
          f'{"OK" if good else "FAIL " + str(miss[:8])}')
    if '--values' in sys.argv:                     # 4. values: the renames change text only
        import lo_session as LO
        import check_full as C
        src = new
        if new.endswith('_quick.xlsx'):           # the quick build has no batch: the release's BaselineCO2 values
            import tempfile
            wq, wr = openpyxl.load_workbook(new), openpyxl.load_workbook(prev_path())
            for row in wr['BaselineCO2'].iter_rows():
                for c in row:
                    if not (isinstance(c.value, str) and c.value.startswith('=')):
                        wq['BaselineCO2'].cell(c.row, c.column).value = c.value
            src = os.path.join(tempfile.mkdtemp(), 'quick_bco2.xlsx')
            wq.save(src)
        data, _errs = LO.start_values(LO.expanded_zip(src, C.LAMBDA_CALL, C.expand), ['Mitigation', 'LiveOutputs'])()
        vo = openpyxl.load_workbook(prev_path(), data_only=True)
        worst, where, n = 0.0, None, 0
        for sh in ('Mitigation', 'LiveOutputs'):
            a, b = vo[sh], data[sh]
            for row in a.iter_rows():
                for c in row:
                    x, y = c.value, b.cell(c.row, c.column).value
                    if isinstance(x, (int, float)) and not isinstance(x, bool):
                        n += 1
                        d = abs(x - y) / max(1.0, abs(x)) if isinstance(y, (int, float)) else float('inf')
                        if d > worst:
                            worst, where = d, f'{sh}!{c.coordinate}'
        good = worst < 1e-9
        ok &= good
        print(f'- values (Mitigation, LiveOutputs) recalculated against v{PREV}, same cells: {n} numbers, max rel diff '
              f'{worst:.2e}{" at " + where if where else ""} {"OK" if good else "FAIL"}')
        gc = B.group_cols(1)[0]                    # the codes the workbook computes = the codes of check 2
        bad = [f'{r}: {data["Mitigation"].cell(r, gc).value} vs {k}' for r, k in cn.items()
               if not isinstance(data['Mitigation'].cell(r, gc).value, str)
               or data['Mitigation'].cell(r, gc).value.split('.', 1)[-1].rsplit('.', 1)[0] != k]
        good = not bad
        ok &= good
        print(f'- Mitigation codes computed by the workbook (VCode, MSlot, RowCode) as check 2: {len(cn) - len(bad)} '
              f'of {len(cn)} {"OK" if good else "FAIL " + str(bad[:5])}')
    print('PASS' if ok else 'FAIL')
    return ok


if __name__ == '__main__':
    sys.exit(0 if main() else 1)
