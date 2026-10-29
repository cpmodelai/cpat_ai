"""Build the standalone Air Pollution workbook input scaffold."""

from __future__ import annotations

import sys
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[1]
VERSION = "0.3"
OUTPUT = HERE / f"CPAT-AI-Air-Pollution-Standalone-v{VERSION}.xlsx"
AP_DATA_DIR = REPO_ROOT / "data" / "air_pollution" / "legacy_459"
AP_TABLES = (
    ("WB", "wb.csv", "Country-year population share, GNI and real interest rates", (4,)),
    ("GBD", "gbd.csv", "GBD exposure, incidence, population and DALY parameters", (4,)),
    ("IHME", "ihme.csv", "Health expenditure and financing shares", (3,)),
    ("ILO", "ilo.csv", "Earnings and output growth", (3,)),
    ("ExposureCoefs", "exposure_coefs.csv", "HAP exposure paths and adjustment coefficients", (4, 206, 404)),
    ("IF PP", "if_pp.csv", "Power-plant intake fractions", (3,)),
    ("ConcMat", "conc_mat.csv", "Concentration matrices by method, country, sector and pollutant", (3,)),
    ("Air_Data", "air_data.csv", "Work absence, survival, labor share, density and calorific values", (3,)),
    ("Air_meta", "air_meta.csv", "Air Pollution data definitions and sources", (3,)),
)

YEARS = tuple(range(2022, 2042))
SCENARIOS = ("Baseline", "Policy")
COUNTRY_DEFAULT = "SEN"
MITIGATION_VERSION = "CPAT-AI-Mitigation-MVP-v0.7.1"

POLLUTANTS = (
    ("nox", "NOx"),
    ("so2", "SO2"),
    ("bc", "Black carbon"),
    ("pom", "Primary organic matter"),
    ("opm", "Other PM2.5"),
    ("voc", "VOC"),
    ("ch4", "CH4"),
    ("co", "CO"),
    ("pm2", "PM2.5"),
    ("oc", "Organic carbon"),
)

SECTORS = (
    ("cpp", "Coal power plants"),
    ("gpp", "Gas power plants"),
    ("opp", "Other power plants (oil, biomass)"),
    ("rod", "Ground level road"),
    ("gro", "Ground level other (residential, services, construction)"),
    ("oth", "Other (rail, aviation, navigation, industry, other energy use)"),
    ("foo", "Food & forestry"),
    ("all", "All sectors"),
)

TITLE_FILL = PatternFill("solid", fgColor="00AE49")
SECTION_FILL = PatternFill("solid", fgColor="92D050")
INPUT_FILL = PatternFill("solid", fgColor="EBF1DE")
CODE_FILL = PatternFill("solid", fgColor="DCE6F1")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="FFFFFF")
SECTION_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
HEADER_FONT = Font(name="Arial", size=10, bold=True)
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=10, color="595959")
THIN_GREY = Side(style="thin", color="BFBFBF")
HEADER_BORDER = Border(bottom=THIN_GREY)


def input_codes() -> list[tuple[str, str, str]]:
    """Return the current Mitigation module's stable aps code block."""
    return [
        (f"aps.{pollutant}.{sector}", pollutant_label, sector_label)
        for pollutant, pollutant_label in POLLUTANTS
        for sector, sector_label in SECTORS
    ]


def style_title(ws, last_col: int, title: str) -> None:
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    cell = ws.cell(row=2, column=1, value=title)
    cell.fill = TITLE_FILL
    cell.font = TITLE_FONT
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 24


def style_section(ws, row: int, last_col: int, title: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    cell.alignment = Alignment(vertical="center")


def build_cover(wb: Workbook) -> None:
    ws = wb.create_sheet("Cover")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 110
    style_title(ws, 3, "CPAT-AI Air Pollution Standalone")

    ws["B4"] = "Status"
    ws["C4"] = "Inputs and reference data are included; Air Pollution calculations are not implemented yet."
    ws["C4"].fill = NOTE_FILL
    ws["C4"].font = NOTE_FONT
    ws["C4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 32

    style_section(ws, 6, 3, "Purpose and boundary")
    rows = [
        ("Purpose", "Standalone rebuild of the Air Pollution module in CPAT 1.0pre_459."),
        ("Mitigation handoff", "Baseline and policy emissions keyed by the existing aps.<pollutant>.<sector> codes."),
        ("Scope", "Emissions-to-concentration methods, exposure, health outcomes, economic outputs, and report-only fuel externalities."),
        ("Country", "SEN is prefilled to match the 1.0pre_459_NoPropData Dashboard selection."),
        ("Horizon", "2022-2041 input columns from Mitigation; the 459 Air Pollution sheet ends at 2040."),
        ("Data", "The nine AP reference tables from 1.0pre_459_NoPropData are embedded on their own tabs."),
        ("External links", "None. The workbook accepts pasted values and does not link to another workbook."),
    ]
    for row, (label, text) in enumerate(rows, start=7):
        ws.cell(row=row, column=2, value=label).font = HEADER_FONT
        cell = ws.cell(row=row, column=3, value=text)
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30 if len(text) > 90 else 20

    style_section(ws, 16, 3, "Build status")
    status_rows = [
        ("Implemented", "Cover, paste-ready Mitigation inputs, Data Sources index, and nine legacy AP data tabs."),
        ("Not implemented", "Calculation formulas, results, legacy regression, and coded-module numeric reconciliation."),
        ("Method source", "coded/cpat_documentation/user/06_Air_pollution.html"),
        ("Equation inventory", "excel_ai/air_pollution/Equations_v1_0.md (working; reconcile before formulas)."),
        ("Legacy target", "excel_ai/legacy/workbooks/CPAT 1.0pre_459_NoPropData.xlsb (read-only ground truth)."),
    ]
    for row, (label, text) in enumerate(status_rows, start=17):
        ws.cell(row=row, column=2, value=label).font = HEADER_FONT
        cell = ws.cell(row=row, column=3, value=text)
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    style_section(ws, 24, 3, "Version log")
    version_header = ("Version", "Date", "Description", "Legacy regression")
    for col, value in enumerate(version_header, start=2):
        cell = ws.cell(row=25, column=col, value=value)
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
    for col, value in enumerate((VERSION, date.today().isoformat(), "Senegal scaffold with embedded AP data", "Not run"), start=2):
        ws.cell(row=26, column=col, value=value).font = BODY_FONT
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 28

    ws.freeze_panes = "B4"


def csv_value(value: str):
    if value == "":
        return None
    try:
        number = float(value)
    except ValueError:
        return value
    if value.startswith("0") and len(value) > 1 and value[1].isdigit():
        return value
    return int(number) if number.is_integer() else number


def read_ap_table(file_name: str) -> list[list[str]]:
    path = AP_DATA_DIR / file_name
    if not path.is_file():
        raise FileNotFoundError(f"Missing extracted Air Pollution table: {path}")
    with path.open("r", newline="", encoding="utf-8") as stream:
        return list(csv.reader(stream))


def build_data_sources(wb: Workbook) -> None:
    ws = wb.create_sheet("Data Sources")
    ws.sheet_view.showGridLines = False
    headers = ("Workbook tab", "Legacy source tab", "CSV extract", "Source rows", "Source columns", "Contents")
    style_title(ws, len(headers), "Data Sources: legacy Air Pollution tables")
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=value)
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
    for row, (sheet_name, csv_name, description, _) in enumerate(AP_TABLES, start=5):
        source_rows = read_ap_table(csv_name)
        values = (sheet_name, sheet_name, csv_name, len(source_rows), max(map(len, source_rows), default=0), description)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in enumerate((20, 20, 26, 14, 16, 68), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A5"


def build_data_divider(wb: Workbook) -> None:
    ws = wb.create_sheet("DATA_AIR_POLL->")
    ws.sheet_view.showGridLines = False
    style_title(ws, 4, "Air Pollution reference data")
    headers = ("Tab", "Contents", "In legacy CPAT", "Legacy source tab")
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=value)
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
    for row, (sheet_name, csv_name, description, _) in enumerate(AP_TABLES, start=5):
        values = (sheet_name, description, "Yes", sheet_name)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in enumerate((20, 72, 16, 22), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A5"


def build_ap_data_tables(wb: Workbook) -> None:
    for sheet_name, csv_name, _, header_rows in AP_TABLES:
        rows = read_ap_table(csv_name)
        ws = wb.create_sheet(sheet_name)
        ws.sheet_view.showGridLines = False
        column_count = max(map(len, rows), default=1)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
        source_parts = []
        if rows:
            source_parts.extend(str(value) for value in rows[0] if value not in (None, ""))
        if rows and sheet_name not in {"ConcMat", "Air_meta"}:
            source_parts.extend(str(value) for value in rows[1] if value not in (None, ""))
        source_title = " | ".join(source_parts)
        source_line = ws.cell(
            row=1,
            column=1,
            value=f"Source: CPAT 1.0pre_459_NoPropData.xlsb | Legacy tab: {sheet_name} | {source_title}",
        )
        source_line.font = NOTE_FONT
        source_line.alignment = Alignment(wrap_text=True, vertical="center")
        style_title(ws, column_count, f"{sheet_name}: Air Pollution reference data")

        # Preserve source rows 3 onward at their legacy addresses. ConcMat and
        # Air_meta have header rows at source row 2, so those keyed headers move
        # to row 3 and later rows shift by one under the workbook title bar.
        if sheet_name in {"ConcMat", "Air_meta"} and len(rows) > 1:
            row_map = [(3, rows[1])]
            row_map.extend((source_row + 1, rows[source_row - 1]) for source_row in range(3, len(rows) + 1))
        else:
            row_map = [(source_row, rows[source_row - 1]) for source_row in range(3, len(rows) + 1)]
        for row_index, values in row_map:
            for column_index, value in enumerate(values, start=1):
                cell = ws.cell(row=row_index, column=column_index, value=csv_value(value))
                cell.font = BODY_FONT
        for header_row in header_rows:
            for cell in ws[header_row]:
                if cell.value is not None:
                    cell.font = HEADER_FONT
                    cell.border = HEADER_BORDER
        for column_index in range(1, column_count + 1):
            ws.column_dimensions[get_column_letter(column_index)].width = 18 if column_index > 1 else 28
        ws.freeze_panes = f"A{min(header_rows) + 1}"

        for header_row in header_rows:
            source_row = header_row - 1 if sheet_name in {"ConcMat", "Air_meta"} else header_row
            expected = [csv_value(value) for value in rows[source_row - 1]]
            actual = [ws.cell(row=header_row, column=column).value for column in range(1, len(expected) + 1)]
            if actual != expected:
                raise ValueError(f"Header row {header_row} of embedded tab {sheet_name!r} does not match its source.")


def build_inputs(wb: Workbook) -> None:
    ws = wb.create_sheet("Mitigation Inputs")
    ws.sheet_view.showGridLines = False
    descriptor_headers = (
        "Country",
        "Mitigation Output Code",
        "Legacy Baseline Code",
        "Legacy Policy Code",
        "Pollutant",
        "AP Sector",
        "Mitigation Unit",
        "Legacy Unit",
        "Input Source",
        "Source Version",
    )
    first_year_col = len(descriptor_headers) + 1
    last_col = len(descriptor_headers) + len(YEARS) * len(SCENARIOS)
    style_title(ws, last_col, "Mitigation Inputs: air-pollutant emissions")

    note = (
        "Paste baseline and policy values by aps code and year. Mitigation v0.7.1 defines aps values in kt; "
        "1.0pre_459 reads country.mit.em.<pollutant>.<sector>.<scenario> in Ton/year. "
        "The legacy code columns are a traceability crosswalk only; no conversion is applied in this scaffold."
    )
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    ws.cell(row=3, column=1, value=note).font = NOTE_FONT
    ws.cell(row=3, column=1).fill = NOTE_FILL
    ws.cell(row=3, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 32

    for scenario_index, scenario in enumerate(SCENARIOS):
        start = first_year_col + scenario_index * len(YEARS)
        end = start + len(YEARS) - 1
        ws.merge_cells(start_row=4, start_column=start, end_row=4, end_column=end)
        cell = ws.cell(row=4, column=start, value=f"{scenario} (scenario {scenario_index + 1})")
        cell.fill = SECTION_FILL
        cell.font = SECTION_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for col, year in enumerate(YEARS, start=start):
            year_cell = ws.cell(row=5, column=col, value=year)
            year_cell.font = HEADER_FONT
            year_cell.alignment = Alignment(horizontal="center")
            year_cell.border = HEADER_BORDER

    for col, header in enumerate(descriptor_headers, start=1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row, (code, pollutant, sector) in enumerate(input_codes(), start=6):
        _, pollutant_code, sector_code = code.split(".")
        metadata = (
            COUNTRY_DEFAULT,
            code,
            None,
            None,
            pollutant,
            sector,
            "kt",
            "Ton/year",
            "Mitigation output code",
            MITIGATION_VERSION,
        )
        for col, value in enumerate(metadata, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")
            if col in (1, 7, 8, 9, 10):
                cell.fill = INPUT_FILL
            if col == 2:
                cell.fill = CODE_FILL
            if col in (3, 4):
                scenario_number = col - 2
                cell.value = (
                    f'=LOWER(A{row})&".mit.em.{pollutant_code}."'
                    f'&"{sector_code}.{scenario_number}"'
                )
                cell.fill = CODE_FILL
                cell.font = BODY_FONT
        for col in range(first_year_col, last_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = INPUT_FILL
            cell.font = BODY_FONT
            cell.number_format = "#,##0.000"

    country_validation = DataValidation(type="textLength", operator="equal", formula1="3")
    country_validation.error = "Enter a three-character ISO country code."
    country_validation.errorTitle = "Check country code"
    country_validation.prompt = "Three-character ISO code; SEN matches the NoPropData reference selection."
    country_validation.promptTitle = "Country"
    ws.add_data_validation(country_validation)
    country_validation.add(f"A6:A{5 + len(input_codes())}")

    ws.auto_filter.ref = f"A5:{get_column_letter(last_col)}{5 + len(input_codes())}"
    ws.freeze_panes = f"{get_column_letter(first_year_col)}6"
    widths = {1: 12, 2: 25, 3: 35, 4: 35, 5: 25, 6: 64, 7: 15, 8: 15, 9: 27, 10: 34}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for col in range(first_year_col, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.row_dimensions[5].height = 30


def build_workbook() -> Path:
    codes = input_codes()
    code_names = [code for code, _, _ in codes]
    if len(codes) != 80 or len(set(code_names)) != len(code_names):
        raise ValueError("The Mitigation aps input code block must contain 80 unique codes.")

    wb = Workbook()
    wb.remove(wb.active)
    build_cover(wb)
    build_data_sources(wb)
    build_inputs(wb)
    build_data_divider(wb)
    build_ap_data_tables(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "manual"
    wb.save(OUTPUT)

    sys.path.insert(0, str(REPO_ROOT / "excel_ai" / "shared"))
    import excel_lint

    problems = excel_lint.lint(OUTPUT)
    if problems:
        raise RuntimeError("Excel lint failed:\n" + "\n".join(problems))
    return OUTPUT


if __name__ == "__main__":
    print(f"Built {build_workbook()}")