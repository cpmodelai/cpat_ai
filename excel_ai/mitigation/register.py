"""The register as one list of records, the same fields for every module and kind (v8.4, user 2026-10-28).

Source data: issues_register.py (bugs, issues, upgrades), build.py (LegacyDiff rows), make_legacy_fixes.py (priority,
fix and effect of the legacy bugs), the cell plans (excel_ai/legacy/bugfixes/data) and the check of the CPAT team's
fixes in legacy 465 (excel_ai/legacy/bugfixes/check465). Every output reads records() from here: the legacy bugs
document (make_bugfix_book.py), the tracker (make_tracker.py), the documentation appendices (make_docs.py) and the
lists BUGS.md, ISSUES.md (make_replication_lists.py).

The rule (NORMS 0): a bug is fixed, never replicated; an issue is replicated and tracked. Data problems are issues
(area Data; v8.5: the kinds DB and DI are retired): replicated, or kept by user decision; where legacy needs a formula
change for them, that change is a bug (L) with its cell plan.
"""
import functools
import json
import os
import re

import build as B
import issues_register as R

HERE = os.path.dirname(os.path.abspath(__file__))
XL = os.path.join(HERE, '..')
PLANS = os.path.join(XL, 'legacy', 'bugfixes', 'data')
CHECK465 = os.path.join(XL, 'legacy', 'bugfixes', 'check465', 'legacy465_check.json')
LEGACY_FILE = 'CPAT 1.0pre_465_NoPropData.xlsb'

KINDS = {  # prefix: (kind, plural, where the full entry is)
    'L': ('Legacy bug', 'Legacy bugs', 'Legacy CPAT bugs document'),
    'M': ('Prototype bug', 'Prototype bugs', 'Documentation, appendix'),
    'I': ('Issue', 'Issues', 'Documentation, appendix'),
    'LD': ('Departure', 'Departures (LegacyDiff)', 'Documentation, appendix'),
    'U': ('Upgrade', 'Upgrades', 'Documentation, appendix')}
ORDER = ['L', 'M', 'I', 'LD', 'U']
DEFINITIONS = [  # prefix, definition, what the prototype does
    ('L', 'Legacy bug: legacy CPAT does not compute what it means to (a wrong row, a double application, a dead '
          'reference, a blank read as 0 where a value is meant); the fix is a change of formulas, with a cell plan.',
     'Fixes it; never replicated.'),
    ('M', 'Prototype bug: a defect of the prototype itself, found and fixed during development.', 'Fixed.'),
    ('I', 'Issue: legacy does what it means to, but the choice is open to question (a method, a default, a '
          'convention, the input data: area Data), or a limitation of the prototype. Priority: Major or Minor; data '
          'problems Very high (a bug fix covers them) or Extremely high (no cell plan possible; the prototype does '
          'exactly what legacy does).', 'Replicates it and tracks it; a departure only by user decision, recorded '
                                        'with the issue.'),
    ('LD', 'Departure: a row of the prototype\'s LegacyDiff sheet that is not "Same" (a decision, a simplification, '
           'a part not built yet).', 'Departs; each says why.'),
    ('U', 'Upgrade: a possible improvement beyond legacy.', 'Not built.')]

POWER = ({f'L{k}' for k in range(1, 13)} | {f'L{k}' for k in range(52, 57)} | {'L67', 'L83', 'L84', 'L88', 'L95'}
         | {f'L{k}' for k in range(107, 112)} | {'L126', 'L127'})
TRANSPORT = {f'L{k}' for k in range(25, 42)}
DIST = {f'L{k}' for k in range(138, 169)} | {f'I{k}' for k in range(76, 85)} | {'I96'}
INPUTS = {'L137', 'I85', 'I86', 'I87', 'I88'}
DATA = R.DATA_ISSUES | {'L170'}                       # v8.5: the data problems (were DI, DB)
EXTRA_STATUS = {  # legacy bugs not in R.LEGACY: status in the prototype
    'L135': 'Not built in the prototype (Manual inputs J185:J188 not read), so the prototype gives legacy\'s values; '
            'found by the cell plans (2026-10-20), to confirm with the CPAT team.',
    'L136': 'Fixed by construction: every fuel\'s new-policy price carries the new ETS and its VAT; found by the cell '
            'plans (2026-10-20), to confirm with the CPAT team.',
    'L170': 'Fixed: a blank VAT cell takes the general VAT rate (VAT_WEO) for residential coal and gas and the oil '
            'products, 0 for power, industry and biomass (since v0.14; was departure LD3, data bug DB2).'}
SEE_ALSO = {'I98': ['L58', 'L59'], 'I97': ['L170', 'LD3'], 'L170': ['I97', 'LD3'], 'I99': ['LD4', 'LD33', 'I14'], 
            'LD3': ['L170', 'I97'], 'LD4': ['I99'], 'LD33': ['I99'], 'L58': ['I98'], 'L59': ['I98'], 'L63': [], 'L64': [],
            'L137': ['I86', 'I87', 'I88'], 'I86': ['L137'], 'I87': ['L137'], 'I88': ['L137']}


def prefix(code):
    return re.match(r'[A-Z]+', code).group(0)


def sort_key(code):
    return ORDER.index(prefix(code)), int(re.search(r'\d+', code).group(0))


@functools.lru_cache(maxsize=None)
def plans():
    out = {}
    for name in ('power_engineer', 'transport_engineer', 'other_mitigation', 'other_mitigation_2', 'distribution', 'v90'):
        p = os.path.join(PLANS, f'fixes_{name}.json')
        if os.path.exists(p):
            for b in json.load(open(p, encoding='utf-8')):
                out[b['id']] = b
    return out


@functools.lru_cache(maxsize=None)
def check465():
    return json.load(open(CHECK465, encoding='utf-8'))['bugs'] if os.path.exists(CHECK465) else {}


def legacy465(code):
    """Status of a legacy bug in legacy 465 (the CPAT team's fixes, checked cell by cell on 2026-10-22)."""
    c = check465().get(code) or check465().get(R.FORMER.get(code, ''))
    if c is None:
        return 'Open', ''
    return {'Confirmed fixed': 'Fixed', 'Fixed with caveats': 'Fixed with caveats',
            'Not fixed': 'Not fixed (marked fixed)'}[c['status']], c['note']


@functools.lru_cache(maxsize=None)
def _fixes():
    import make_legacy_fixes as F                       # priorities, fixes and effects of the legacy bugs
    return F.PRIORITY, F.FIX


def _area(code):
    if code in DATA:
        return 'Data'
    if code in INPUTS:
        return 'Inputs (ScenarioDefinition)'
    if code in DIST:
        return 'Distribution'
    if code in POWER:
        return 'Mitigation: engineer power'
    if code in TRANSPORT:
        return 'Mitigation: engineer road transport'
    return 'Mitigation'


def _rec(**k):
    code = k['id']
    p = prefix(code)
    base = dict(id=code, former=R.FORMER.get(code, ''), kind=KINDS[p][0], module=_area(code).split(':')[0],
                area=_area(code), title='', where='', description='', fix='', effect='', priority='',
                prototype='', legacy465='', legacy465_note='', status='', see_also=', '.join(SEE_ALSO.get(code, [])),
                documented=KINDS[p][2], cell_plan='')
    base.update(k)
    return base


@functools.lru_cache(maxsize=None)
def records():
    """Every item of the register, in code order (L, M, I, LD, U)."""
    pri, fix = _fixes()
    pl = plans()
    out = []
    for i, (_k, title, where, what, fx, why) in enumerate(R.BUGS, 1):           # L1-L10
        c = f'L{i}'
        out.append(_rec(id=c, title=title, where=f'Mitigation {where}', description=what, fix=fx, effect=why,
                        priority=pri.get(c, 'Low'), prototype='Fixed'))
    for c, title, where, what, st in R.LEGACY:                                 # L11-L134
        fx, eff = fix.get(c, ('', ''))
        out.append(_rec(id=c, title=title, where=where, description=what, fix=fx, effect=eff,
                        priority=pri.get(c, 'Low'), prototype=st))
    for c in ('L135', 'L136', 'L170'):                                          # found by the cell plans; L170 v8.5
        b = pl[c]
        out.append(_rec(id=c, title=b['title'], where=b.get('where', ''), description=b.get('why', ''),
                        fix=b.get('ref', ''), effect=b.get('effect', ''),
                        priority=(b.get('pri') or 'Low').split(' (')[0], prototype=EXTRA_STATUS[c]))
    for c in sorted((k for k in pl if re.fullmatch(r'L(17[1-9]|1[89]\d|2\d\d)', k)), key=sort_key):   # v9.0
        b = pl[c]
        out.append(_rec(id=c, title=b['title'], where=b.get('where', ''), description=b.get('why', ''),
                        fix=b.get('ref', ''), effect=b.get('effect', ''),
                        priority=(b.get('pri') or 'Low').split(' (')[0], prototype=b.get('prototype', ''),
                        cell_plan='Yes' if b.get('fixes') else ''))
    for c, title, where, what, st in R.INPUT_BUGS:                              # L137
        out.append(_rec(id=c, title=title, where=where, description=what, priority='Low', prototype=st,
                        fix='Name the MTInputs values without the D_ prefix and without the space.'))
    for c, title, where, what, st in R.DIST_BUGS:                               # L138-L168
        out.append(_rec(id=c, title=title, where=where, description=what, priority=pri.get(c, 'Low'), prototype=st))
    for c, title, desc, where in R.ISSUES:
        sev, st, nxt = R.ISSUE_STATUS.get(c, ('', '', ''))
        leg = R.ISSUE_LEGACY.get(c, '')
        if c.startswith('L'):                                                   # L169: an issue that is a bug
            out.append(_rec(id=c, title=title, where='Mitigation row 18449 H', description=desc, priority='Low',
                            prototype=f'Fixed (v5.7): {nxt}', fix='Read the data (GHGs row 889).'))
            continue
        proto = {'replicated': 'Replicated', 'not replicated': 'Departs (user decision)',
                 'mvp': 'Prototype limitation'}.get(leg, leg)
        if c in R.DATA_ISSUES and leg == 'replicated':
            proto = 'Replicated'
        if c == 'I97':
            proto = 'The prototype has the fix of bug L170 (a blank takes the general VAT rate)'
        out.append(_rec(id=c, title=title, description=desc, fix=nxt, priority=sev, prototype=proto, status=st,
                        documented=f'{KINDS[prefix(c)][2]}; {where}'))
    for c, title, desc, st in R.DIST_ISSUES:
        out.append(_rec(id=c, title=title, description=desc, prototype='Replicated' if st.startswith('Replicated')
                        else st, status='Accepted', priority='Minor'))
    for c, title, found, what, fx in R.MVP_BUGS:
        out.append(_rec(id=c, title=title, description=what, fix=fx, prototype=f'Fixed ({found})', status='Closed'))
    for k, (area, legacy, model, why, effect, typ) in enumerate(B.LEGACY_DIFFS, 1):
        if typ == 'Same':
            continue
        c = f'LD{k}'
        out.append(_rec(id=c, title=area, description=f'Legacy: {legacy}. Prototype: {model}.', fix=why,
                        effect=effect, priority=typ, prototype='Kept as legacy' if k in R.LD_REPLICATED else 'Departs',
                        status='Accepted'))
    for c, title, desc, p_ in R.UPGRADES:
        out.append(_rec(id=c, title=title, description=desc, priority=p_, prototype='Not built', status='Potential'))
    for r in out:
        if r['kind'] in ('Legacy bug',):
            st, note = legacy465(r['id'])
            r['legacy465'], r['legacy465_note'] = st, note
            r['status'] = 'Closed in legacy 465' if st == 'Fixed' else 'Open in legacy'
            r['cell_plan'] = 'yes' if r['id'] in pl or R.FORMER.get(r['id']) in pl else ''
    ids = [r['id'] for r in out]
    assert len(ids) == len(set(ids)), [c for c in ids if ids.count(c) > 1]
    return tuple(sorted(out, key=lambda r: sort_key(r['id'])))


def by_kind(p):
    return [r for r in records() if prefix(r['id']) == p]


def counts():
    return {p: len(by_kind(p)) for p in ORDER}


if __name__ == '__main__':
    print(counts())


TRACKER_OLD = os.path.join(XL, 'legacy', 'Tracking issues.xlsx')      # the tracker before v8.4 (tab CPAT-AI_bugs)
NEW_DATE = '2026-10-28'


@functools.lru_cache(maxsize=None)
def _tracker_rows():
    """{code: (version found in, date, raised by)} from the latest tracker (Tracking_Legacy_Bugs_v*.xlsx, tab Bugs)
    or, before v8.4, from tab CPAT-AI_bugs of legacy/Tracking issues.xlsx (keyed by the former codes)."""
    import glob
    import openpyxl
    out = {}
    new = sorted(glob.glob(os.path.join(XL, 'Tracking_Legacy_Bugs_v*.xlsx')) +
                 glob.glob(os.path.join(XL, 'releases', 'Tracking_Legacy_Bugs_v*.xlsx')))
    if new:
        ws = openpyxl.load_workbook(new[-1], read_only=True)['Bugs']
        rows = ws.iter_rows(values_only=True)
        hdr = list(next(rows))
        for r in rows:
            if r and r[0]:
                d = dict(zip(hdr, r))
                out[r[0]] = (str(d.get('CPAT Version with identified issue') or ''), d.get('Date'),
                             d.get('Issue raised by'))
        return out
    ws = openpyxl.load_workbook(TRACKER_OLD, read_only=True)['CPAT-AI_bugs']
    back = {v: k for k, v in R.FORMER.items()}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            out[back.get(str(r[0]), str(r[0]))] = (str(r[2]) if r[2] else '', r[3], r[6])
    return out


def found(code):
    """(legacy version the bug was found in, date, raised by)."""
    v = _tracker_rows().get(code)
    if v:
        return v
    if code == 'L169':                                  # registered as issue I67 (v5.7)
        return ('456', '2026-10-19', 'Stephen (CPAT-AI)')
    return ('465', NEW_DATE, 'Stephen (CPAT-AI)')


PRIORITY_ORDER = ['Extremely high', 'Very high', 'High', 'Major', 'Medium', 'Minor', 'Low']


def data_items():
    """v8.5: the data problems (area Data): the bug L170 and the issues I89-I99, most urgent first."""
    rank = lambda r: (PRIORITY_ORDER.index(r['priority']) if r['priority'] in PRIORITY_ORDER else 9, sort_key(r['id']))
    return sorted((r for r in records() if r['area'] == 'Data'), key=rank)
