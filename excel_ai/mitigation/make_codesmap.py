"""CPAT-AI_CodesMap_v<version>.xlsx (excel_ai/): the single codes map (v0.8.4, user 2026-10-28), inputs included.

The codes map of v0.7.1 (excel_ai/legacy/codes/CPAT_CodesMap_v0_7_1.xlsx: protocol, vocabularies, variables, crosswalk;
built by legacy/codes/make_codes_map_v0_7_1.py) with the current state added in front:
  - Inputs: every row of ScenarioDefinition (legacy MTInputs 465 rows 8-413, then the inputs the prototype adds): row,
    label, unit, the name (NameOfParameter, the code of the input in Excel and Python), the legacy name and what changed
    with its register code (bug L137, issue I87), the Name Manager names, the Dashboard twin and the Python key;
  - CodeRenames: the code renames applied since v0.7.2 (data/mvp_legacy_extracts/code_renames_v0_7_2.csv);
  - Outputs: every output code of the workbook (MTOutputs) with its description.
The v0.7.1 tabs keep their content (their names are as of v0.7.1; the renames tab says what changed since).

    python make_codesmap.py          (after build.py)
"""
import csv
import os
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import build as B
import issues_register as R

HERE = os.path.dirname(os.path.abspath(__file__))
XL = os.path.join(HERE, '..')
ROOT = os.path.join(XL, '..')
BASE = os.path.join(XL, 'legacy', 'codes', 'CPAT_CodesMap_v0_7_1.xlsx')
INPUTS = os.path.join(ROOT, 'data', 'mvp_legacy_extracts', 'mtinputs_465.csv')
RENAMES = os.path.join(ROOT, 'data', 'mvp_legacy_extracts', 'code_renames_v0_7_2.csv')
PY = [os.path.join(ROOT, 'coded', 'cpat_model', f) for f in ('dashboard_inputs.py', 'distribution_inputs.py')]
OUT = B.REL.versioned('CPAT-AI_CodesMap', R.REGISTER_VERSION, 'xlsx')
HEAD = PatternFill('solid', fgColor='D9E8C4')
CHANGED = PatternFill('solid', fgColor='FCE4D6')
PY_GROUPS = {'SPPEff': 'adj_eff_margins_shadow_pricing', 'AdditionalEfficiency': 'additional_eff_gains'}


def python_keys():
    keys = set()
    for p in PY:   # v0.9.2: retired in Coded CPAT 6.3 (it reads the inputs by NameOfParameter); skipped if gone
        if os.path.exists(p):
            keys |= set(re.findall(r'^\s+(\w+)\s*:', open(p, encoding='utf-8').read(), flags=re.M))
    return keys


def sheet(wb, name, header, rows, widths, pos=None, changed_col=None):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, pos)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in rows:
        ws.append(list(r))
        if changed_col is not None and r[changed_col]:
            for c in ws[ws.max_row]:
                c.fill = CHANGED
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def input_rows():
    pyk = python_keys()
    out = []
    for r in csv.DictReader(open(INPUTS, encoding='utf-8')):
        if r['kind'] != 'input':
            continue
        n = r['name']
        py = n if n in pyk else next((f'{g}[{n}]' for k, g in PY_GROUPS.items() if n.startswith(k)), '')
        out.append([int(r['row']), r['section'], r['label'], r['unit'], n, r['legacy_name'], r['change'],
                    r['register'], r['name_manager_F'], r['name_manager_other'], r['dashboard_twin'], py,
                    'legacy MTInputs 465'])
    wb = openpyxl.load_workbook(B.OUT, read_only=True)
    ws = wb['ScenarioDefinition']
    for i, row in enumerate(ws.iter_rows(min_row=414, max_col=8, values_only=True), 414):
        if row[7]:
            out.append([i, 'Prototype additions', row[1], row[2], row[7], '', 'input the prototype adds', '', '', '',
                        '', row[7] if row[7] in pyk else '', 'prototype (no legacy MTInputs row)'])
    return out


def output_rows():
    wb = openpyxl.load_workbook(B.OUT, read_only=True, data_only=True)
    ws = wb['MTOutputs']                                   # v1.0.4: legacy layout, E code, D description
    out = []
    for row in ws.iter_rows(min_row=B.MTO_R0, values_only=True):
        code = row[4] if len(row) > 4 else None
        if isinstance(code, str) and re.fullmatch(r'[a-z]{3}\.mit\.[\w.]+', code):
            out.append([code, row[3] or ''])
    return out


def main():
    wb = openpyxl.load_workbook(BASE)
    inputs = input_rows()
    sheet(wb, 'Inputs', ['Row (ScenarioDefinition = MTInputs 465)', 'Section', 'Label', 'Unit',
                         'Name (NameOfParameter: Excel and Python)', 'Legacy name (MTInputs H)', 'Change',
                         'Register', 'Name Manager (value cell F)', 'Name Manager (other cells)', 'Dashboard twin',
                         'Python key (coded)', 'Source'], inputs, [9, 22, 50, 12, 30, 30, 26, 9, 26, 30, 18, 30, 22],
          pos=1, changed_col=6)
    ren = list(csv.reader(open(RENAMES, encoding='utf-8')))
    sheet(wb, 'CodeRenames', ren[0], ren[1:], [6, 10, 22, 22, 16, 30, 10], pos=2)
    outs = output_rows()
    sheet(wb, 'Outputs', ['Output code', 'Description'], outs, [40, 90], pos=3)
    ws = wb['Cover']
    for row in ws.iter_rows():
        for c in row:
            c.value = None
    lines = [f'CPAT codes map v{R.REGISTER_VERSION} (the single codes map; CPAT-AI-Mitigation-Prototype v{B.VERSION})',
             f'Built by excel_ai/mitigation/make_codesmap.py from the v0.7.1 codes map (legacy/codes/CPAT_CodesMap_v0_7_1.xlsx), '
             f'the input names (data/mvp_legacy_extracts/mtinputs_465.csv, extract_inputs.py), the code renames since '
             f'v0.7.2 and the workbook. Do not edit: change the scripts or their sources.', None, 'Sheets',
             f'Inputs - every input row of ScenarioDefinition ({len(inputs)}): legacy MTInputs 465 rows, then the '
             'inputs the prototype adds. Name = the code of the input in Excel (column H) and Python. Orange rows: a '
             'name that changed or was added (register: bug L137, issue I87). Legacy inconsistencies kept as issues: '
             'I85 (percentages typed as whole numbers), I86 (spelling), I88 (one Dashboard cell for two inputs).',
             'CodeRenames - the output-code renames applied since v0.7.2 (code_renames_v0_7_2.csv).',
             f'Outputs - every output code of MTOutputs ({len(outs)}).',
             'Protocol, Positions, vocabularies (Tab ... CountryCode), Variables, Crosswalk, MTOutputs, MVPOnly, Checks '
             '- as of v0.7.1 (make_codes_map_v0_7_1.py); codes renamed since are in CodeRenames.', None,
             'Remediation plan of v0.7.1: legacy/codes/Codes_Remediation_Plan_v0_7_1.md. Bugs and issues: tracker '
             f'Tracking_Legacy_Bugs_v{R.REGISTER_VERSION}.xlsx.']
    for i, t in enumerate(lines, 1):
        ws.cell(i, 1, t)
    ws['A1'].font = Font(bold=True, size=13)
    ws['A4'].font = Font(bold=True)
    wb.move_sheet('Cover', -wb.index(wb['Cover']))
    wb.save(OUT)
    import sys
    sys.path.insert(0, os.path.join(XL, 'shared'))
    import excel_lint
    excel_lint.check(OUT)
    B.REL.publish(OUT)
    print(f'wrote {OUT}: {len(inputs)} inputs, {len(ren) - 1} code renames, {len(outs)} output codes')


if __name__ == '__main__':
    main()
