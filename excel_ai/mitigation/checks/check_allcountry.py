"""v0.4.5 all-country test (run on request, about 6 minutes): python check_allcountry.py [workbook]
Would the workbook work with the full (all-country) data tables? Every row of a data tab that belongs to Egypt is
cloned under the code XEG (country codes EGY/egy replaced by XEG/xeg) and appended below the last row of its tab, as
pasted all-country data would arrive; key formulas are translated. Settings C4 is set to XEG, the workbook is
recalculated (LAMBDAs expanded), and every output code of the clone (xeg.*) must equal Egypt's (egy.*) from a
recalculation of the unchanged workbook, with no error values. It fails if a lookup range is too short, if a table
with country rows is not the last on its tab, or if any formula depends on Egypt other than through the country code.
Derived tables that the build computes for one country (BaselineCO2) are cloned like data.
v0.4.6: tabs kept at their legacy cell addresses (fixed legacy ranges, already all countries) are tested by renaming
Egypt's codes to XEG in place instead of appending a clone."""
import os as _os, sys as _sys  # the excel_ai code folders and data/scripts on the path (core/tools/xlpaths.py)
_xl = _os.path.dirname(_os.path.abspath(__file__))
while _os.path.basename(_xl) != 'excel_ai':
    _xl = _os.path.dirname(_xl)
_sys.path.insert(0, _os.path.join(_xl, 'core', 'tools'))
import xlpaths  # noqa: E402,F401
import os, re, sys
sys.path.insert(0, os.getcwd())
import openpyxl
from openpyxl.formula.translate import Translator
import build as B, check_full as C

ISO = re.compile(r'(?<![A-Za-z])(EGY|egy)(?![A-Za-z])')
ROW_INDEX = ('.ghg.h.',)     # outputs that are row numbers of a data table (the clone's rows are elsewhere): only
                             # whether a row was found (> 0) must agree
ERRORS = ('#DIV/0!', '#N/A', '#NAME?', '#NULL!', '#NUM!', '#REF!', '#VALUE!', 'Err:')
CLONE = {'EGY': 'XEG', 'egy': 'xeg'}
MODULE_TABS = ['FromAP', 'FromTrans']   # 1.2.2: the module tabs are keyed by country too (typed for Egypt): cloned
IN_PLACE = {'Countries', 'Elasticities', 'ETS+LTS', 'Multipliers', 'Prices_int', 'ECP', 'Transport3', 'Power', 'CoC',
            'GHGs'}   # v0.4.6; v0.4.7: Power, CoC, GHGs


def data_tabs(wb):
    names = wb.sheetnames
    a = names.index('DATA_MITIGATION->>')          # 1.2.2: the data tabs are last (no REF group)
    return [n for n in names[a:] if n not in B.DIVIDERS]


def rename_rows(ws):
    """v0.4.6: rename Egypt's codes to XEG in place (tabs at legacy addresses); returns the rows changed."""
    n = 0
    for row in ws.iter_rows(min_row=3):
        hit = False
        for c in row:
            v = c.value
            if isinstance(v, str) and not v.startswith('=') and ISO.search(v):
                c.value = ISO.sub(lambda m: CLONE[m.group(1)], v)
                hit = True
        n += hit
    return n


def clone_rows(ws):
    """Append a copy of every Egypt row (row 4 on) below the last data row of the sheet; returns rows cloned."""
    rows = [r for r in range(4, ws.max_row + 1)
            if any(isinstance(c.value, str) and not c.value.startswith('=') and ISO.search(c.value)
                   for c in ws[r])]
    data = lambda c: c.value not in (None, '') and not (isinstance(c.value, str) and c.value.startswith('='))
    dest = max(r for r in range(1, ws.max_row + 1) if any(data(c) for c in ws[r])) + 1   # after the last data row
                                        # (key formulas filled below a table do not count; column A may be empty)
    for k, r in enumerate(rows):
        for c in ws[r]:
            v = c.value
            if isinstance(v, str) and v.startswith('='):
                v = Translator(v, origin=c.coordinate).translate_formula(f'{c.column_letter}{dest + k}')
            elif isinstance(v, str):
                v = ISO.sub(lambda m: CLONE[m.group(1)], v)
            ws.cell(dest + k, c.column, v)
    return len(rows)


def outputs(data):
    """{code: values} of every output code on Mitigation and Distribution (code columns of the scenario groups)."""
    out = {}
    for name, cols in (('Mitigation', [B.group_cols(g)[0] for g in (1, 2, 3)]),
                       ('Distribution', [B.dist_group_cols(g)[0] for g in (1, 2, 3)])):
        if name not in data.sheetnames:            # 1.2.2: Distribution is its own workbook since 1.2.0
            continue
        ws = data[name]
        for c in cols:
            for r in range(1, ws.max_row + 1):
                code = ws.cell(r, c).value
                if isinstance(code, str) and '.' in code:
                    out[code] = [ws.cell(r, c + k).value for k in range(1, 12)]
    return out


def main(path):
    base = C.recalc(C.variant(openpyxl.load_workbook(path), 'expand')).parent
    egy = outputs(base)
    wb = openpyxl.load_workbook(path)
    n = {t: rename_rows(wb[t]) if t in IN_PLACE else clone_rows(wb[t]) for t in data_tabs(wb) + MODULE_TABS}
    wb[B.TAB_CONST][B.settings_cell('C4')] = 'XEG'   # v1.2.0: the country setting (the link to MTInputs overridden)
    data = C.recalc(C.variant(wb, 'expand')).parent
    errs = [(ws.title, c.coordinate) for ws in data.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith(ERRORS)]
    xeg = outputs(data)
    num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
    worst, at, m, missing = 0.0, None, 0, 0
    for code, vals in egy.items():
        x = xeg.get('xeg' + code[3:]) if code.startswith('egy') else xeg.get(code)
        if x is None:
            missing += 1
            continue
        if any(t in code for t in ROW_INDEX):
            vals, x = ([float(v > 0) if num(v) else v for v in w] for w in (vals, x))
        for a, b in zip(vals, x):
            if num(a) and num(b):
                m += 1
                d = abs(a - b) / max(1.0, abs(a))
                if d > worst:
                    worst, at = d, code
            elif num(a) != num(b):
                worst, at = max(worst, 1.0), code
    ok = not errs and missing == 0 and worst < 1e-9
    print(f'rows cloned as XEG: {sum(n.values())} in {sum(1 for v in n.values() if v)} tabs '
          f'({", ".join(f"{t} {k}" for t, k in n.items() if k)})')
    print(f'error values with XEG: {len(errs)} {errs[:6]}')
    print(f'outputs XEG vs EGY: {m} values, codes not found {missing}, max relative difference {worst:.3g} {at or ""}')
    print('PASS' if ok else 'FAIL')


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else B.OUT)
