"""v0.2.3: imports build / check_full (previous version in Old/; method unchanged).

Compare CPAT-AI-Mitigation-Prototype (scenario 1, baseline) with the cached baseline results in the legacy workbook.

The legacy file CPAT 1.0pre_456_NoPropData.xlsb keeps the cached values of its last full run (its country: Egypt, real data; the comparison needs COUNTRY of the build to match):
fuel use and after-tax prices by subsector and fuel, CO2 by subsector. This script
  1. extracts them by output code (egy.mit.ener.<sub>.<fuel>.e.1, egy.mit.atp.<sub>.<fuel>.e.1,
     egy.mit.co2.<grp>.<sub>.1) -> data/legacy_cached_baseline.csv;
  2. recalculates the prototype twice in LibreOffice: as shipped (last historical price year 2022, as legacy; v0.3.7) and
     with the last historical price year set to 2024 (Settings C10; the LAMBDA formulas are copied over whole rows,
     so the history / projection switch follows the setting);
  3. writes legacy_comparison_v6_4.md: totals by sector, selected prices, and the fit of each variant.

    python legacy_compare.py
"""
import os as _os, sys as _sys  # the excel_ai code folders and data/scripts on the path (core/tools/xlpaths.py)
_xl = _os.path.dirname(_os.path.abspath(__file__))
while _os.path.basename(_xl) != 'excel_ai':
    _xl = _os.path.dirname(_xl)
_sys.path.insert(0, _os.path.join(_xl, 'core', 'tools'))
import xlpaths  # noqa: E402,F401
from legacy_csv import legacy_dir, save_to_tidy  # noqa: E402,F401  (data/Public01/tidy/ in the legacy CSV form)
import csv
import os
from collections import defaultdict

import openpyxl
from openpyxl.utils import get_column_letter as L

import build as B
import check_full as C

HERE = xlpaths.MIT                     # excel_ai/mitigation/ (the script sits in a subfolder)
LEGACY = os.path.join(HERE, '..', 'legacy', 'workbooks', 'CPAT 1.0pre_456_NoPropData.xlsb')
CSV_OUT = os.path.join(legacy_dir(), 'legacy_cached_baseline.csv')
REPORT = os.path.join(HERE, 'legacy_comparison_v6_4.md')
YEARS = list(range(2022, 2041))
SUBS = [s[0] for s in B.SUBSECTORS]
GRP = {s[0]: s[2] for s in B.SUBSECTORS}
LEG_GRP = {'tra': 'tra', 'bld': 'bld', 'ind': 'ind', 'oen': 'ind'}    # legacy co2 codes: oen sits in 'oth'
FUELS = [f[0] for f in B.FUELS]
P = B.COUNTRY.lower()   # v0.3.3: output-code prefix; the legacy file's cached run must be of the same country


def extract():
    from pyxlsb import open_workbook
    # v0.3.11: legacy has jet fuel in domestic aviation only; the other jfu rows count as zero
    sf = [(s, f) for s in SUBS for f in FUELS if f != 'jfu' or s == 'avi']
    want = {f'{P}.mit.ener.{s}.{f}.e.1' for s, f in sf} | \
        {f'{P}.mit.atp.{s}.{f}.e.1' for s, f in sf} | \
        {f'{P}.mit.co2.{LEG_GRP[GRP[s]]}.{s}.1' for s in SUBS if s != 'oen'} | {f'{P}.mit.co2.oth.1'}
    rows = {}
    with open_workbook(LEGACY) as wb, wb.get_sheet('Mitigation') as sh:
        for i, row in enumerate(sh.rows()):
            vals = [c.v for c in row]
            for j, v in enumerate(vals[:12]):
                if isinstance(v, str) and v in want and v not in rows:
                    rows[v] = (i + 1, [vals[11 + k] if 11 + k < len(vals) else None for k in range(len(YEARS))])
                    break
    with open(CSV_OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(['code', 'legacy_row'] + YEARS + ['source'])
        for k in sorted(rows):
            w.writerow([k, rows[k][0]] + rows[k][1] + ['CPAT 1.0pre_456_NoPropData.xlsb Mitigation (cached values)'])
    missing = sorted(want - set(rows))
    return {k: v[1] for k, v in rows.items()}, missing


class V71(dict):
    """v0.7.2 (codes R9): prototype values by output code, also found by the v0.7.1 code the comparison names (rename map)."""
    def __getitem__(self, k):
        return dict.__getitem__(self, k if dict.__contains__(self, k) else C.RC(k))

    def get(self, k, d=None):
        try:
            return self[k]
        except KeyError:
            return d


def model(last_hist):
    wb = openpyxl.load_workbook(B.OUT)
    if last_hist != B.LAST_HIST:
        wb[B.TAB_CM]['C10'] = last_hist                # v1.2.0: Settings block
        path = C.variant(wb, 'everywhere')       # LAMBDA over whole rows, expanded: the switch follows C10
    else:
        path = C.variant(wb, 'expand')
    ws = C.recalc_wb(path)['Mitigation']
    code, base, years = B.group_cols(1)
    vals = {}
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, code).value
        if isinstance(k, str):
            vals[k] = [ws.cell(r, c).value for c in [base] + years]
    return V71(vals)


def main():
    leg, missing = extract()
    leg = defaultdict(lambda: [None] * len(YEARS), leg)    # v0.3.11: rows legacy does not have (jfu outside aviation)
    ours = {2022: model(2022), 2024: model(2024)}
    units = {r['fuel']: float(r['gj_per_price_unit']) for r in csv.DictReader(open(os.path.join(legacy_dir(),
                                                                                                'fuel_units.csv')))}
    yi = {y: k for k, y in enumerate(YEARS)}
    show = [2022, 2023, 2024, 2025, 2027, 2030]

    def lsum(var, grp):
        out = []
        for y in show:
            if var == 'ener':
                out.append(sum(leg[f'{P}.mit.ener.{s}.{f}.e.1'][yi[y]] or 0 for s in SUBS if GRP[s] == grp
                               for f in FUELS))
            else:
                ks = [f'{P}.mit.co2.{LEG_GRP[GRP[s]]}.{s}.1' for s in SUBS if GRP[s] == grp and s != 'oen'] + \
                    ([f'{P}.mit.co2.oth.1'] if grp == 'oen' else [])
                out.append(sum(leg[k][yi[y]] or 0 for k in ks))
        return out

    def osum(v, var, grp):
        sfx = '.e' if var == 'ener' else '.e'
        return [v[f'{P}.mit.{var}.{grp}.all{sfx}.1'][yi[y]] for y in show]

    lines = ['# Legacy CPAT cached baseline vs CPAT-AI-Mitigation-Prototype (scenario 1)', '',
             'Legacy: cached values of the last full run stored in `CPAT 1.0pre_456_NoPropData.xlsb` (Egypt, real '
             f'data), extracted to `data/legacy_cached_baseline.csv`. prototype: v{B.VERSION} scenario 1 recalculated in '
             'LibreOffice, (a) as shipped, last historical price year 2022 (as legacy); (b) last historical price year '
             '2024, the latest IMF data (Settings C10, LAMBDAs copied over whole rows; issue I89). Single-run tables: '
             'as shipped. Fuel use ktoe, CO2 MtCO2.', '',
             f'Legacy codes not found: {", ".join(missing) or "none"}.', '']
    for var, unit in (('ener', 'ktoe'), ('co2', 'MtCO2')):
        lines += [f'## {"Fuel use" if var == "ener" else "CO2 from fuel combustion"} by sector ({unit})', '',
                  '| Sector | Run | ' + ' | '.join(map(str, show)) + ' |', '|---|---|' + '---|' * len(show)]
        for _n, grp, title in B.SECTIONS:
            fmt = (lambda x: f'{x:,.0f}') if var == 'ener' else (lambda x: f'{x:,.1f}')
            lines.append(f'| {title} | legacy | ' + ' | '.join(fmt(x) for x in lsum(var, grp)) + ' |')
            for lh in (2022, 2024):
                lines.append(f'| | prototype, prices to {lh} | ' + ' | '.join(fmt(x) for x in osum(ours[lh], var, grp))
                             + ' |')
        lines.append('')
    # v0.2.8: GHG emissions by inventory category (legacy section 15, cached rows)
    from pyxlsb import open_workbook
    grows = {7565: 'enr', 7570: 'ipr', 7575: 'agr', 7580: 'lucf', 7585: 'wst', 7590: 'oth', 7600: 'all'}
    lg = {}
    with open_workbook(LEGACY) as wb_, wb_.get_sheet('Mitigation') as sh:
        for row in sh.rows(sparse=True):
            if row and row[0].r + 1 in grows:
                v_ = {c.c: c.v for c in row if c.v is not None}
                lg[grows[row[0].r + 1]] = [v_.get(11 + k) for k in range(len(YEARS))]
            if row and row[0].r + 1 > 7600:
                break
    gshow = [2022, 2024, 2025, 2030, 2035, 2040]
    names = {'enr': 'Energy', 'ipr': 'Industrial processes', 'agr': 'Agriculture', 'lucf': 'LULUCF', 'wst': 'Waste',
             'oth': 'Other', 'all': 'Total excluding LULUCF'}
    lines += ['## GHG emissions by inventory category (Mt CO2e), scenario 1 (v0.2.8)', '',
              '| Category | Run | ' + ' | '.join(map(str, gshow)) + ' |', '|---|---|' + '---|' * len(gshow)]
    for c, nm in names.items():
        key = f'{P}.mit.ghg.{c}.all.1' if c != 'all' else f'{P}.mit.ghg.all.all.1'
        mine = ours[B.LAST_HIST].get(key)
        lines.append(f'| {nm} | legacy | ' + ' | '.join(f'{lg[c][yi[y]]:,.1f}' for y in gshow) + ' |')
        if mine:
            lines.append('| | prototype | ' + ' | '.join(f'{mine[yi[y]]:,.1f}' for y in gshow) + ' |')
    lines += ['', 'Why they differ: energy follows the fuel-use gap above (and energy CO2 has no inventory '
              'adjustment, I18); the legacy file has no combustion factor paths (zero combustion CH4), so its energy '
              'methane is fuel supply only; the legacy population path grows 2-4% a year from 2027 (L20), which '
              'raises agriculture, waste and other after 2030; legacy livestock, rice and waste methane are constant '
              '(L18); industrial processes follow industrial energy CO2 in both.', '']
    # v0.3.0: local air pollutants (legacy section 13 totals by air-pollution sector, section 14 forcing; cached rows)
    pblk = {'pm2': 6693, 'nox': 6795, 'so2': 6897, 'voc': 6999, 'bc': 7101, 'oc': 7203, 'co': 7412}
    secs = [s_ for s_, _n in B.AP_SECS]
    want = {pblk[p] + 2 + i: (p, s_) for p in pblk for i, s_ in enumerate(secs)}
    want.update({pblk[p] + 11: (p, 'all') for p in pblk})
    frc = {7516: 'bc', 7517: 'co', 7518: 'voc', 7519: 'oc', 7520: 'nox', 7521: 'so2', 7522: 'all'}
    lp, lf = {}, {}
    with open_workbook(LEGACY) as wb_, wb_.get_sheet('Mitigation') as sh:
        for row in sh.rows(sparse=True):
            r_ = row[0].r + 1 if row else 0
            if r_ in want or r_ in frc:
                v_ = {c.c: c.v for c in row if c.v is not None}
                vals = [v_.get(11 + k) or 0.0 for k in range(len(YEARS))]
                if r_ in want:
                    lp[want[r_]] = [x / 1000 for x in vals]          # legacy t/year -> kt
                else:
                    lf[frc[r_]] = vals
            if r_ > 7522:
                break
    m24 = ours[B.LAST_HIST]
    pnames = dict(B.AP_POLS)
    lines += ['## Local air pollutants: total emissions by pollutant (kt), scenario 1 (v0.3.0)', '',
              '| Pollutant | Run | ' + ' | '.join(map(str, gshow)) + ' |', '|---|---|' + '---|' * len(gshow)]
    for p in pblk:
        ssum = [sum(lp[(p, s_)][k] for s_ in secs) for k in range(len(YEARS))]
        lines.append(f'| {pnames[p]} | legacy total | ' + ' | '.join(f'{lp[(p, "all")][yi[y]]:,.1f}' for y in gshow) + ' |')
        lines.append('| | legacy, sum of sectors | ' + ' | '.join(f'{ssum[yi[y]]:,.1f}' for y in gshow) + ' |')
        mine = m24.get(f'{P}.mit.ap.all.{p}.1')
        if mine:
            lines.append('| | prototype | ' + ' | '.join(f'{mine[yi[y]]:,.1f}' for y in gshow) + ' |')
    lines += ['', '## PM2.5 by air-pollution module sector (kt), scenario 1 (v0.3.0)', '',
              '| Sector | Run | ' + ' | '.join(map(str, gshow)) + ' |', '|---|---|' + '---|' * len(gshow)]
    for s_, nm in B.AP_SECS:
        lines.append(f'| {nm} | legacy | ' + ' | '.join(f'{lp[("pm2", s_)][yi[y]]:,.1f}' for y in gshow) + ' |')
        mine = m24.get(f'{P}.mit.aps.{s_}.pm2.1')
        if mine:
            lines.append('| | prototype | ' + ' | '.join(f'{mine[yi[y]]:,.1f}' for y in gshow) + ' |')
    lines += ['', '## Forcing of local air pollutants (Mt CO2e, GWP100), scenario 1 (v0.3.0)', '',
              '| Pollutant | Run | ' + ' | '.join(map(str, gshow)) + ' |', '|---|---|' + '---|' * len(gshow)]
    for p in ['bc', 'co', 'voc', 'oc', 'nox', 'so2', 'all']:
        lines.append(f'| {pnames.get(p, "Total")} | legacy | ' + ' | '.join(f'{lf[p][yi[y]]:,.2f}' for y in gshow) + ' |')
        mine = m24.get(f'{P}.mit.apf.all.{p}.1' if p != 'all' else f'{P}.mit.apf.all.all.1')
        if mine:
            lines.append('| | prototype | ' + ' | '.join(f'{mine[yi[y]]:,.2f}' for y in gshow) + ' |')
    lines += ['', 'Why they differ: fuel use differs (tables above); the legacy file has no year-by-year factor paths '
              '(the prototype draws a straight line from the base-year to the final GAINS factor); legacy leaves navigation '
              'out of every pollutant row (L22); its total formula leaves services and food & forestry out (L23), but '
              'they emit zero in the legacy Egypt run, so its total equals the sum of its sectors. '
              'From 2024 the prototype is higher mainly because its fuel use is (see above).', '']
    lines += ['## After-tax prices, scenario 1 (legacy units: $/liter for liquids, $/GJ otherwise; real 2026 USD)', '',
              '| Row | Run | ' + ' | '.join(map(str, show)) + ' |', '|---|---|' + '---|' * len(show)]
    for s, f in (('rod', 'gso'), ('rod', 'die'), ('res', 'lpg'), ('res', 'nga'), ('cem', 'coa'), ('omn', 'oop')):
        k = f'{P}.mit.atp.{s}.{f}.e.1'
        u = units[f] if f in ('gso', 'die', 'lpg', 'ker') else (1.0 if f != 'oop' else units['oop'])
        lines.append(f'| {s} {f} | legacy | ' + ' | '.join(f'{leg[k][yi[y]]:.3f}' for y in show) + ' |')
        for lh in (2022, 2024):
            lines.append(f'| | prototype, prices to {lh} | ' + ' | '.join(f'{ours[lh][k][yi[y]] * u:.3f}' for y in show)
                         + ' |')
    lines.append('')
    # fit: mean absolute % difference of subsector x fuel fuel use, 2023-2030, rows with legacy > 1 ktoe
    lines += ['## Fit: fuel use by subsector and fuel, 2023-2030', '',
              '| Run | Mean abs. % difference vs legacy | Total fuel use 2030 vs legacy |', '|---|---|---|']
    for lh in (2022, 2024):
        diffs, tl, to = [], 0.0, 0.0
        for s in SUBS:
            for f in FUELS:
                k = f'{P}.mit.ener.{s}.{f}.e.1'
                for y in range(2023, 2031):
                    a, b = leg[k][yi[y]] or 0, ours[lh][k][yi[y]]
                    if a > 1:
                        diffs.append(abs(b / a - 1))
                tl += leg[k][yi[2030]] or 0
                to += ours[lh][k][yi[2030]]
        lines.append(f'| prototype, prices to {lh} | {sum(diffs) / len(diffs):.1%} | {to / tl - 1:+.1%} |')
    lines.append('')
    with open(REPORT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    print('\n'.join(lines))


if __name__ == '__main__':
    main()
