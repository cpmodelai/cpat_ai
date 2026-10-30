"""v0.9.0 (R10.8-R10.9): goal seek of the range-anxiety start values on the prototype's own engineer road model.

Legacy CPAT stores, by country, the range anxiety (USD per vehicle) of passenger cars, freight and 2-3 wheelers that
makes the model's EV sales share equal the actual one in the fit year (sheet Transport3, A1 = 2025; solver results in
columns H, N, T; Mitigation rows 4354-4356 read them). The prototype's road model differs in detail from legacy's
(issue I61), so legacy's stored values do not fit it. This script repeats legacy's fit on the prototype: for each class
it finds the start value RA0 with residual r(RA0) = calculated - actual EV sales share in 2025 = 0 (rows raf.pas,
raf.fre, raf.cyc of section 5b, scenario 1), by the secant method (bracketing kept, bisection fallback), each step a
LibreOffice recalculation of a working copy with the trial values written into Transport3 (names RAPas, RAFre, RACyc).
The engineer road model computes in every scenario group whatever TrModel is (TrModel only chooses which road fuel
use enters the totals), so the fit runs on scenario 1 as defined in the workbook.

A goal seek cannot be a formula of the sheet (the start value is read by every year from the base year, the residual
is a 2025 result: a formula would read a later year, and the root needs iteration), so the result is stored data:

    python ra_goalseek.py WORKBOOK.xlsx [--tol 1e-6] [--max-iter 8]
    -> range_anxiety_fit.csv in data/Public01/tidy/ (via data/scripts/legacy_csv.py) (row of the workbook's country replaced)

The builder writes the stored values over the country's legacy cells of Transport3 (bright yellow, ra_fit_apply).
A negative root is kept (as legacy, issue I145).
"""
import os as _os, sys as _sys  # the excel_ai code folders and data/scripts on the path (core/tools/xlpaths.py)
_xl = _os.path.dirname(_os.path.abspath(__file__))
while _os.path.basename(_xl) != 'excel_ai':
    _xl = _os.path.dirname(_xl)
_sys.path.insert(0, _os.path.join(_xl, 'core', 'tools'))
import xlpaths  # noqa: E402,F401
from legacy_csv import legacy_dir, save_to_tidy  # noqa: E402,F401  (data/Public01/tidy/ in the legacy CSV form)
import argparse
import csv
import os
import re

import openpyxl

import build as B
import check_full as C

CLASSES = (('passenger', 'RAPas', 'raf.pas'), ('freight', 'RAFre', 'raf.fre'), ('cycle', 'RACyc', 'raf.cyc'))
FIT_YEAR = 2025                                   # legacy Transport3 A1


def _dest(wb, name):
    """(sheet, column letter, first row) of a defined name 'Sheet!$H$5:$H$700'."""
    t = wb.defined_names[name].attr_text
    m = re.match(r"'?([^'!]+)'?!\$([A-Z]+)\$(\d+):\$[A-Z]+\$\d+", t)
    return m.group(1), m.group(2), int(m.group(3))


def _ctry_name():
    for r in csv.DictReader(open(os.path.join(B.DATA, 'range_anxiety.csv'), encoding='utf-8')):
        if r['iso'] == B.COUNTRY:
            return r['country']
    raise SystemExit(f'{B.COUNTRY} not in range_anxiety.csv')


def _country_row(wb):
    sh, col, r0 = _dest(wb, 'RAKey')
    ws, name = wb[sh], _ctry_name()
    for r in range(r0, ws.max_row + 1):
        if str(ws[f'{col}{r}'].value) == name:
            return r
    raise SystemExit(f'country {name} not in {sh}!{col}')


def evaluate(path, ra):
    """Residuals {class: calculated - actual EV sales share in FIT_YEAR} with the start values ra {class: value}."""
    wb = openpyxl.load_workbook(path)
    row = _country_row(wb)
    for cls, nm, _k in CLASSES:
        sh, col, _r0 = _dest(wb, nm)
        wb[sh][f'{col}{row}'].value = ra[cls]
    ws = C.recalc(C.variant(wb, 'expand'))
    _code, base, years = B.group_cols(1)
    c = ([base] + years)[FIT_YEAR - B.BASE_YEAR]
    out = {cls: float(ws.cell(B.R_TR[k], c).value) for cls, _n, k in CLASSES}
    return out


def solve(path, tol=1e-6, max_iter=8, log=print):
    wb = openpyxl.load_workbook(path)
    row = _country_row(wb)
    x0 = {}
    for cls, nm, _k in CLASSES:
        sh, col, _r0 = _dest(wb, nm)
        x0[cls] = float(wb[sh][f'{col}{row}'].value)
    pts = {cls: [] for cls, *_ in CLASSES}         # (x, residual) per class
    x = dict(x0)
    for it in range(1, max_iter + 1):
        res = evaluate(path, x)
        log(f'  iteration {it}: ' + ', '.join(f'{k} RA0 {x[k]:,.2f} residual {res[k]:+.6f}' for k in x))
        for k in x:
            pts[k].append((x[k], res[k]))
        if all(abs(res[k]) < tol for k in x):
            return x, res, it
        for k in x:
            if abs(res[k]) < tol:
                continue
            p = pts[k]
            if len(p) == 1:                         # first step: 20% up (anxiety lowers the EV share)
                x[k] = p[0][0] * 1.2 if p[0][0] else 1000.0
                continue
            (xa, ra), (xb, rb) = p[-2], p[-1]
            neg = [q for q in p if q[1] < 0]
            pos = [q for q in p if q[1] > 0]
            nxt = xb - rb * (xb - xa) / (rb - ra) if rb != ra else xb * 1.5
            if neg and pos:                         # keep the root bracketed
                lo, hi = max(neg, key=lambda q: q[1]), min(pos, key=lambda q: q[1])
                a_, b_ = sorted((lo[0], hi[0]))
                if not a_ < nxt < b_:
                    nxt = (lo[0] + hi[0]) / 2
            x[k] = nxt
    best = {k: min(pts[k], key=lambda q: abs(q[1])) for k in pts}   # the evaluated point nearest the root, per class
    return {k: best[k][0] for k in best}, {k: best[k][1] for k in best}, max_iter


def store(x, res, it, path):
    out = os.path.join(legacy_dir(), 'range_anxiety_fit.csv')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    head = ['iso', 'country', 'passenger', 'freight', 'cycle', 'fit_year', 'residual_passenger', 'residual_freight',
            'residual_cycle', 'iterations', 'source']
    rows = [r for r in csv.DictReader(open(out, encoding='utf-8'))] if os.path.exists(out) else []
    rows = [r for r in rows if r['iso'] != B.COUNTRY]
    rows.append({'iso': B.COUNTRY, 'country': _ctry_name(), **{k: f'{x[k]:.6f}' for k in x}, 'fit_year': FIT_YEAR,
                 **{f'residual_{k}': f'{res[k]:.3e}' for k in res}, 'iterations': it,
                 'source': f'ra_goalseek.py on {os.path.basename(path)} (scenario 1)'})
    with open(out, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, head)
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: r['iso']))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('workbook')
    ap.add_argument('--tol', type=float, default=1e-6, help='largest |residual| (share; default 1e-6)')
    ap.add_argument('--max-iter', type=int, default=8)
    a = ap.parse_args()
    x, res, it = solve(a.workbook, a.tol, a.max_iter)
    print(f'stored {store(x, res, it, a.workbook)} after {it} recalculations: {x}')


if __name__ == '__main__':
    main()
