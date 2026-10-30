"""Screenshots of the current CPAT-AI-Mitigation-Prototype workbook for the core deck (presentation/img/v<version>_*.png,
e.g. v0_7_1_) and the deck's numbers (img/v0_7_1_data.json, v0_7_1_facts.json with the register counts, v0_7_1_transport.json). Every view is rendered from the real
workbook: recalculated in LibreOffice (LAMBDA calls expanded), one sheet turned into a values copy (chosen cells show
their formula text), rows / columns / print area set, printed to PDF by LibreOffice and cut to a PNG. The legacy-bug
counts (img/v0_6_2_legacy.json) come from make_legacy_fixes.py.

    python presentation/make_screenshots.py
"""
import os as _os, sys as _sys  # the excel_ai code folders and data/scripts on the path (core/tools/xlpaths.py)
_xl = _os.path.dirname(_os.path.abspath(__file__))
while _os.path.basename(_xl) != 'excel_ai':
    _xl = _os.path.dirname(_xl)
_sys.path.insert(0, _os.path.join(_xl, 'core', 'tools'))
import xlpaths  # noqa: E402,F401
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from PIL import Image, ImageChops

HERE = os.path.dirname(os.path.abspath(__file__))     # excel_ai/mitigation/docs/presentation/
import build as B                      # noqa: E402
import check_full as C                      # noqa: E402

IMG = os.path.join(HERE, 'img')
PFX = 'v' + B.VERSION.replace('.', '_')     # v0_6_2
LO_PROFILE = 'file://' + os.path.join(tempfile.gettempdir(), 'lo_shots_profile_v0_6_2')


def recalculated():
    """Values workbook as Excel 365 shows it: the shipped file with the LAMBDA calls expanded, recalculated in
    LibreOffice."""
    return C.recalc_wb(C.variant(openpyxl.load_workbook(B.OUT), 'expand'))


def soffice(args, outdir):
    subprocess.run(['soffice', f'-env:UserInstallation={LO_PROFILE}', '--headless'] + args + ['--outdir', outdir],
                   check=True, capture_output=True, timeout=900)


def crop(png):
    im = Image.open(png).convert('RGB')
    bg = Image.new('RGB', im.size, (255, 255, 255))
    box = ImageChops.difference(im, bg).getbbox()
    if box:
        im = im.crop((max(box[0] - 8, 0), max(box[1] - 8, 0), box[2] + 8, box[3] + 8))
    im.save(png)


def render(name, sheet, values, area, rows=None, cols=None, show_formula=(), widths=None, hide_cols=(),
           show_cols=(), wrap_cols=(), heights=None):
    """rows: set of rows to show (others hidden) or None (keep the shipped outline state)."""
    wb = openpyxl.load_workbook(B.OUT)
    ws = wb[sheet]
    vs = values[sheet]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                if c.coordinate in show_formula:
                    c._value, c.data_type = c.value, 's'              # show the formula text
                    c.alignment = Alignment(wrap_text=True, vertical='top')
                    c.font = Font(name='Courier New', size=11, color=c.font.color.rgb if c.font and c.font.color
                                  and isinstance(c.font.color.rgb, str) else '000000')
                else:
                    c.value = vs[c.coordinate].value
    for other in [w for w in wb.worksheets if w.title != sheet]:
        wb.remove(other)
    wb.defined_names.clear()
    if rows is not None:                     # explicit rows: drop the outline so collapsed groups do not hide them
        ws.sheet_format.outlineLevelRow = 0
        for r in range(1, ws.max_row + 2):
            d = ws.row_dimensions[r]
            d.outlineLevel, d.collapsed, d.hidden = 0, False, r not in rows
    if show_cols:
        ws.sheet_format.outlineLevelCol = 0
        for d in list(ws.column_dimensions.values()):
            d.outline_level, d.collapsed = 0, False
    for col in hide_cols:
        ws.column_dimensions[col].hidden = True
    for col in show_cols:
        ws.column_dimensions[col].hidden = False
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    for col in wrap_cols:
        for r in range(1, ws.max_row + 1):
            ws[f'{col}{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    for r, h in (heights or {}).items():
        ws.row_dimensions[r].height = h
    ws.freeze_panes = None
    ws.print_area = area
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.headings = True
    ws.print_options.gridLines = True
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 0.2
    ws.oddHeader.center.text = ws.oddFooter.center.text = ''
    tmp = tempfile.mkdtemp()
    x = os.path.join(tmp, f'{name}.xlsx')
    wb.save(x)
    soffice(['--convert-to', 'pdf', x], tmp)
    subprocess.run(['pdftoppm', '-png', '-r', '170', '-f', '1', '-l', '1', os.path.join(tmp, f'{name}.pdf'),
                    os.path.join(tmp, name)], check=True)
    png = [f for f in os.listdir(tmp) if f.startswith(name) and f.endswith('.png')][0]
    dst = os.path.join(IMG, f'{name}.png')
    shutil.move(os.path.join(tmp, png), dst)
    crop(dst)
    print('wrote', dst)


def facts():
    """Numbers of the quality slide, from the check reports of this version (quick check; full check if present)."""
    out = {}
    q = os.path.join(xlpaths.MIT, 'checks', 'check_quick_report.md')
    txt = open(q, encoding='utf-8').read()
    out['quick'] = 'PASS' if '**PASS**' in txt else 'FAIL'
    m = re.search(r'- (\d+) formulas; plain formulas over 100 characters: \[\]', txt)
    out['formulas'] = f'{int(m.group(1)):,}' if m else None
    m = re.search(r'error values (\d+)', txt)
    out['errors'] = m.group(1) if m else None
    m = re.search(r'output codes vs (\S+): (\d+) shared, (\d+) new, (\d+) gone; (\d+) change', txt)
    if m:
        out.update(prev=m.group(1), shared=f'{int(m.group(2)):,}', codes=f'{int(m.group(2)) + int(m.group(3)):,}',
                   changed=m.group(5))
    m = re.search(r'Distribution group 2 vs the Python reference[^:]*: (\d+) values, max relative difference (\S+)',
                  txt)
    out['distn_ref'] = (m.group(1), m.group(2)) if m else None
    full = os.path.join(xlpaths.MIT, 'checks', 'check_report.md')
    if os.path.exists(full):
        ft = open(full, encoding='utf-8').read()
        m = re.search(r'Overall[^\n]*', ft)
        v = re.search(r'Prototype v(\S+)', ft)            # v1.0.0: the version the full check last ran on
        out['full_overall'] = m.group(0) if m else None
        out['full_version'] = v.group(1) if v else None
    import issues_register as IR                 # v0.7.1: counts of the register (as the management briefing)
    from collections import Counter
    st = Counter(IR.ISSUE_LEGACY.values())
    # v1.0.0: the counts of BUGS.md and ISSUES.md (generated from the whole register before the deck), so the deck
    # and the lists agree; mitigation = every area but Distribution
    top = xlpaths.CORE                                  # BUGS.md, ISSUES.md
    rows = [r.split('|') for r in open(os.path.join(top, 'BUGS.md'), encoding='utf-8').read().split('## 2.')[0]
            .splitlines() if re.match(r'\| L\d+ ', r)]
    dist = [r for r in rows if r[3].strip().startswith('Distribution')]
    out['bugs'] = {'mit': len(rows) - len(dist), 'dist': len(dist)}
    out['bug_priority'] = dict(Counter(r[4].strip() for r in rows if r not in dist))
    irows = [r.split('|') for r in open(os.path.join(top, 'ISSUES.md'), encoding='utf-8').read()
             .split('## 2. Issues (I)')[1].split('## 3.')[0].splitlines() if re.match(r'\| I\d+ ', r)]
    idist = sum(r[3].strip().startswith('Distribution') for r in irows)
    out['issues'] = {'mit': len(irows) - idist, 'dist': idist,
                     'replicated': sum(r[7].strip() == 'Replicated' for r in irows)}
    out['links'] = [t for t, *_x in B.DATA_LINKS]
    with open(os.path.join(IMG, f'{PFX}_facts.json'), 'w', encoding='utf-8') as f:
        json.dump(out, f, indent=1)
    print('facts', out)


def transport():
    """Egypt engineer road baseline, legacy behaviour vs all fixes (Python reference; as make_transport_doc)."""
    import transport_model_ref as T
    inp = T.legacy_inputs('base')
    leg = T.summary(T.engine(inp, fix=frozenset()))
    fixed = T.summary(T.engine(inp, fix=T.ALL_FIXES))
    ys = (2030, 2035, 2041)
    road = lambda s: [round(s[y]['gso'] + s[y]['die'] + s[y]['lpg'] + s[y]['nga']) for y in ys]
    deck = {'years': list(ys), 'nfix': len(T.FIXES), 'road_legacy': road(leg), 'road_fixed': road(fixed),
            'ev_legacy': [leg[y]['ev_pass'] for y in ys], 'ev_fixed': [fixed[y]['ev_pass'] for y in ys]}
    with open(os.path.join(IMG, f'{PFX}_transport.json'), 'w', encoding='utf-8') as f:
        json.dump(deck, f, indent=1)


def main():
    os.makedirs(IMG, exist_ok=True)
    facts()
    transport()
    values = recalculated()
    (k1, b1, y1), (k2, b2, y2) = B.group_cols(1), B.group_cols(2)
    yr = lambda ys, y: L(ys[y - B.BASE_YEAR - 1])
    last = L(y2[-1] + 1)

    # 1. Overview: the sheet as shipped (rolled up), scenario groups 1 and 2
    render(f'{PFX}_01_overview', 'Mitigation', values, f'A1:{last}{B.R_EMI["co2.chg"] + 1}')

    # 1b. Scenarios left to right: top rows and carbon price, code columns shown
    top = set(range(1, B.R_GDP + 1)) | {B.B_POL, B.R_CP}
    render(f'{PFX}_01b_scenarios', 'Mitigation', values, f'A1:{last}{B.R_CP}', rows=top, show_cols=['K'],
           widths={'K': 22, L(k2): 22})

    # 2. Short plain formulas: road gasoline rows, formula text in 2040 (same formula as every year)
    h = B.SUB_HEAD['rod']
    vars_ = ['ctxnew', 'nce', 'atp', 'shp', 'ener']
    rws = [h + B.VOFF[v] + 2 for v in vars_]
    c40 = L(y1[-2])
    cells = {f'{c40}{r}' for r in rws}
    hide = [L(c) for c in range(CI('I'), y1[-2]) if c not in (y1[-3],)] + [L(y1[-1])]
    render(f'{PFX}_02_formulas', 'Mitigation', values, f'A1:{c40}{max(rws)}', rows={1, 4, 5} | set(rws),
           show_formula=cells, hide_cols=hide + ['B', 'D', 'E', 'F', 'G'], show_cols=['C', L(y1[-3]), c40],
           widths={c40: 95, 'H': 34, L(y1[-3]): 11}, heights={r: 36 for r in rws})

    # 3. Power: generation by type, VRE, capacity, investment, electricity prices (scenario 1)
    keep = [('gen', m) for m, _n in B.PW_TECH] + [('gtot', ''), ('vre', ''), ('captot', ''), ('nimtot', ''),
                                                    ('usdtot', ''), ('rp', 'res'), ('rp', 'ind'), ('demtot', '')]
    rows = {1, 4, 5, B.B_POW} | {B.R_PW[k] for k in keep if k in B.R_PW}
    render(f'{PFX}_03_power', 'Mitigation', values, f'A1:{L(y1[-1])}{max(rows)}', rows=rows, hide_cols=['I', 'J', 'K'])

    # 5. ScenarioCompare (v0.7.1): the entries, the scenario slots and the key table (six slots)
    k1 = B.SC_KEY_R0 + 3 * len(B.SC_KEY) - 1
    rows = {1, 4, 5, 6, B.SC_R_ID, B.SC_R_SHORT, B.SC_R_NAME} | set(range(B.SC_KEY_R0 - 1, k1 + 1))
    last_c = L(B.SC_C0 + 5)
    render(f'{PFX}_05_compare', 'ScenarioCompare', values, f'A1:{last_c}{k1}', rows=rows, hide_cols=['A'],
           widths={'B': 40, 'C': 12, **{L(B.SC_C0 + j): 13 for j in range(6)}}, heights={B.SC_R_NAME: 52})

    # 6. ScenarioDefinition: short names, Run?, scenario definitions J:O
    n = B.MT_NAMES
    keys = ['CPIntro', 'CPLevelStart', 'CPLevelTarget', 'CPOutro', 'Feb_Level_Start_Trans',
            'Feb_Level_Target_Trans', 'NewETS', 'ETSChangeRelStart', 'ETSChangeRelTarget']
    rows = {B.MT_ROW_SHORT, 4, 5, 6} | {n[k] for k in keys}
    render(f'{PFX}_06_definitions', 'ScenarioDefinition', values, f'B1:O{max(rows)}', rows=rows,
           hide_cols=list('CDEFGHI'), widths={'B': 40, **{c: 15 for c in 'JKLMNO'}}, heights={6: 64, 4: 26})

    # 8. Distribution: the sheet as shipped (rolled up), groups 1 and 2
    if not B.DSPEC:
        wm, mitrow = values['Mitigation'], {}
        for r_, (a_, f_, s_) in enumerate(wm.iter_rows(min_row=1, max_col=3, values_only=True), 1):
            if isinstance(a_, str) and a_:
                mitrow.setdefault((a_, f_ or '', s_ or ''), r_)
        B._dist_layout(mitrow)
    wd = values['Distribution']
    _c2, bb2, dd2, sp2 = B.dist_group_cols(2)
    rmax = max(r for r in range(1, wd.max_row + 1) if wd.cell(r, 2).value not in (None, ''))
    render(f'{PFX}_08_distribution', 'Distribution', values, f'A1:{L(sp2)}{rmax}')

    # chart data from StoredOutputs and Distribution
    st, names = C.stored_blocks(values['StoredOutputs'])
    years = list(range(B.BASE_YEAR, B.LAST_YEAR + 1))
    get = lambda k, s: st[k].get(s)
    data = {'names': {str(k): v for k, v in names.items()}, 'years': years}
    for key, stem in (('co2', 'mit.co2.all.all.e'), ('co2_pow', 'mit.co2.pow.all.e'), ('vre', 'mit.vre'),
                      ('gtot', 'mit.gtot'), ('ghg', 'mit.ghg.all.all'), ('ch4', 'mit.ghg.all.ch4'),
                      ('rtot', 'mit.rtot.all.all'), ('gdp_eff', 'mit.gdp.eff.all.all'),
                      ('ndc_lvl', 'mit.ndc.lvl.all.sel'), ('ndc_gap', 'mit.ndc.gap.all.all'),
                      ('cov', 'mit.cov.all.all.all'), ('cpi', 'mit.cpi.all.all'), ('xnim', 'mit.xnim.all.all'),
                      ('scc', 'mit.scc.all.nat'), ('wcli', 'mit.wcli.all.nat'), ('eff', 'mit.eff.all.all'),
                      ('wel_usd', 'mit.wel.all.usd'), ('wel_pct', 'mit.wel.all.pct'), ('cptraj', 'mit.cptraj')):
        stem = getattr(B, 'rn_stem', lambda x: x)(stem)    # v0.7.2: the stems through the code rename map
        data[key] = {str(k): get(k, stem) for k in st}
    rs = getattr(B, 'rn_stem', lambda x: x)
    data.update(ets_p=st[5][rs('mit.ets.p')], co2_ets=st[5][rs('mit.co2.ets.all.all')], co2_cap=st[5][rs('mit.co2.cap.all.all')])
    dv = lambda key, j: wd.cell(B.DROW[key], bb2 + j).value
    data['distn'] = {'year': dv('YearDistn', 0), 'scenario': wd.cell(5, _c2).value,
                     'tef': [dv('tef.ove', j) for j in range(1, 11)], 'net': [dv('net.ove', j) for j in range(1, 11)],
                     'gini_rec': dv('pct_gini_rec.ove', 0), 'gini_norec': dv('pct_gini_norec.ove', 0)}
    with open(os.path.join(IMG, f'{PFX}_data.json'), 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=1, default=str)
    print(f'wrote {PFX}_data.json')


if __name__ == '__main__':
    main()
