"""Workbook formulas of CPAT-AI-Mitigation, generated from the workbook (Formulas.md): the implementation
reference. The model equations in mathematical form are in Equations.md (written by hand, no cell detail).

Every equation is read from the Mitigation sheet itself (scenario 1, years 2022-2041), so the list always matches
the built workbook. Cell references become variable codes: code[fuel,sector](year), where f and s stand for the
fuel and sector of the row being defined, t for its year; <label> is a parameter of the row (columns D:G).
Rows and years with the same equation are shown once. Labels, units, descriptions and legacy references come from
the variable dictionary of the builder; changes against legacy CPAT from the LegacyDiff list and the register of
bugs, issues and upgrades.

Run after the build: python make_equations.py
"""
import os as _os, sys as _sys  # the excel_ai code folders and data/scripts on the path (core/tools/xlpaths.py)
_xl = _os.path.dirname(_os.path.abspath(__file__))
while _os.path.basename(_xl) != 'excel_ai':
    _xl = _os.path.dirname(_xl)
_sys.path.insert(0, _os.path.join(_xl, 'core', 'tools'))
import xlpaths  # noqa: E402,F401
import os
import re
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L

import build as B
import issues_register as R

HERE = xlpaths.MIT                     # excel_ai/mitigation/ (the script sits in a subfolder)
OUT = os.path.join(HERE, 'docs', 'Formulas.md')   # v0.3.2: was Equations; v0.7.1: the notes keep the release name v7_0
_, BASE_COL, PROJ = B.group_cols(1)
YCOLS = [BASE_COL] + PROJ                                  # L .. AE = 2022 .. 2041
YEAR = {c: B.BASE_YEAR + i for i, c in enumerate(YCOLS)}
TOP = {B.R_YEAR: 'year', B.R_LHCOL: 'lastHistCol', B.R_SCEN: 'scenario', B.R_NAME: 'scenarioName'}

# Changes against legacy CPAT that touch a variable (LegacyDiff ids LDn, register ids L/M): variable codes
MODS = {
    'LD1': ['ener'], 'LD2': ['sp', 'txo'], 'LD3': ['vat'], 'LD4': ['sp', 'txo'], 'LD5': ['vat'],
    'LD6': ['ets.est', 'ets.p'], 'LD7': ['obr', 'etsb'], 'LD8': ['ets.vadj', 'ets.vol'], 'LD9': ['reg', 'shps'],
    'LD10': ['cptraj', 'fpr', 'fb'], 'LD12': ['fb'], 'LD13': ['ctcov'], 'LD14': ['rsub', 'esub', 'esubpu'],
    'LD16': ['sp'], 'LD17': ['etx', 'esub'], 'LD18': ['fpr'], 'LD20': ['txo'], 'LD21': ['ener'],
    'LD23': ['mult', 'dem', 'lgi', 'nim', 'reb', 'cmin', 'early', 'reteff', 'need', 'usd', 'msc'],
    'LD24': ['shock', 'dem'], 'LD25': ['dem'], 'LD26': ['pret', 'ret'], 'LD27': ['dem'], 'LD29': ['lvc', 'dsum'],
    'LD30': ['fix', 'gnc'], 'LD31': ['fpp'], 'LD32': ['sc', 'rppre', 'rp'], 'LD33': ['rppre'], 'LD35': ['prnew'],
    'LD36': ['co2', 'ghg'], 'LD37': ['rtx', 'rsub', 'rnew'], 'LD38': ['xcp', 'xcpd', 'xck'], 'LD39': ['ener', 'co2', 'prtx'],
    'L1': ['mult'], 'L2': ['mult'], 'L3': ['mult'], 'L4': ['msc', 'cph'], 'L5': ['usd'], 'L6': ['lgi'],
    'L7': ['early'], 'L8': ['reteff'], 'L9': ['reb', 'tvc'], 'L10': ['need'], 'L11': ['gen'], 'L12': ['nimtot'],
    'L13': ['sp', 'txo'], 'L14': ['txo'], 'L15': ['ets.est'], 'L16': ['ets.est'],
    'M2': ['rnvat'], 'M3': ['msc'], 'M4': ['vre'], 'M6': ['prnew'],
    'LD43': ['ch4.act', 'ch4.abt', 'ch4.em'], 'LD44': ['ghg.flr'], 'LD45': ['pop', 'pop.g'], 'LD46': ['mfee', 'ch4.abt'],
    'LD47': ['ch4'], 'L18': ['ch4.act', 'ch4.em'], 'L19': ['ghg.flr'], 'L20': ['pop'], 'L21': ['ch4.abt'],   # v0.2.8
    'LD49': ['ap'], 'LD50': ['ap'], 'L22': ['ap'],   # v0.2.9
    'LD51': ['aps'], 'LD52': ['ap', 'aps'], 'L23': ['ap'],   # v0.3.0
    'LD54': ['gdp.lvl', 'gdp.pos.pct'], 'LD55': ['fis.rev', 'fis.sp'], 'LD56': ['fis.eff', 'fis.key'],
    'L24': ['gdp.lvl', 'gdp.pos.pct'],   # v0.3.1
    'LD57': ['co2', 'ctxnew', 'ets'],   # v0.3.2: calibration
}

REF = re.compile(r"(?:(?:'[^']+'|[A-Za-z_][\w.]*)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?(?![\w(])")
STR = re.compile(r'"[^"]*"')


SETTINGS = {}


def load():
    wb = openpyxl.load_workbook(B.OUT, read_only=True)
    for i, row in enumerate(wb[B.TAB_CM].iter_rows(max_col=2, max_row=90, values_only=True), 1):
        if isinstance(row[1], str):
            SETTINGS[i] = row[1].split(' (')[0][:50]
    rows = [list(r) for r in wb['Mitigation'].iter_rows(max_col=YCOLS[-1] + 1, values_only=True)]
    return [[None] * 0] + rows                                   # 1-based


def key(rows, r):
    a, f, s = (rows[r][i] if r < len(rows) else None for i in range(3))
    ok = isinstance(a, str) and a and not a.startswith('=') and ' ' not in a and not a[:1].isdigit()
    return (a if ok else None, f if isinstance(f, str) and not f.startswith('=') else None,
            s if isinstance(s, str) and not s.startswith('=') else None)


RN = getattr(B, 'rn_var', lambda v: v)          # v0.7.2: the dictionary key as the rename map writes it (column A)
VARINFO = {RN(v[0]): v for v in B.VARIABLES}
_TPL = [list(r) for r in openpyxl.load_workbook(B.MT_TEMPLATE, read_only=True).active.iter_rows(max_col=8,
                                                                                    values_only=True)]


def label_unit(code):
    """Label and unit of a variable; MTInputs inputs take the text of their MTInputs row."""
    v = VARINFO.get(code)
    if v is None:
        return code, ''
    lab, unit = v[1], v[2]
    m = re.search(r'MTInputs row (\d+)', str(v[5]))
    if (not isinstance(lab, str) or lab.startswith('=')) and m and int(m.group(1)) <= len(_TPL):
        row = _TPL[int(m.group(1)) - 1]
        texts = [x for x in row if isinstance(x, str) and x.strip() and x != 'X']
        lab = texts[0] if texts else code
        unit = texts[1] if len(texts) > 2 else ''
    return lab, unit


def plabel(code, k):
    v = VARINFO.get(code)
    lab = v[6 + k] if v is not None and len(v) > 6 + k and v[6 + k] else f'p{k + 1}'
    return lab


def rowsym(rows, rr, r):
    if rr in TOP:
        return TOP[rr]
    a, f, s = key(rows, rr)
    if not a:
        return f'row{rr}'
    _, f0, s0 = key(rows, r)
    idx = [('f' if f == f0 else f) if f else None, ('s' if s == s0 else s) if s else None]
    idx = [i for i in idx if i]
    return a + (f'[{",".join(idx)}]' if idx else '')


def tsym(cc, c, absolute):
    if absolute:
        return str(YEAR[cc]) if cc in YEAR else None
    if not (YCOLS[0] - 2 <= cc <= YCOLS[-1] + 2):               # year columns, and spans just past either end
        return None
    d = cc - c
    return 't' if d == 0 else f't{d:+d}'


def cellsym(rows, r, c, ref):
    m = re.fullmatch(r'(\$?)([A-Z]{1,3})(\$?)(\d+)', ref)
    ca, cl, _ra, rn = m.groups()
    cc, rr = CI(cl), int(rn)
    if cc <= 3 and rr == r:                                      # the row's own code, fuel and sector
        return ('code', 'f', 's')[cc - 1]
    if 4 <= cc <= 7:                                             # parameter columns D:G
        lab = plabel(key(rows, rr)[0], cc - 4)
        return f'<{lab}>' if rr == r else f'<{lab} of {rowsym(rows, rr, r)}>'
    if rr == B.R_YEAR and cc in YEAR:
        return 'year' if not ca else str(YEAR[cc])
    t = tsym(cc, c, bool(ca))
    if t is not None:
        return f'{rowsym(rows, rr, r)}({t})'
    if cc > YCOLS[-1]:
        return f'{rowsym(rows, rr, r)}(scenario {(cc - B.COL_G0) // B.GW + 1})'
    return ref.replace('$', '')


def rangesym(rows, r, c, a, b):
    m1 = re.fullmatch(r'\$?([A-Z]{1,3})\$?(\d+)', a)
    m2 = re.fullmatch(r'\$?([A-Z]{1,3})\$?(\d+)', b)
    c1, r1, c2, r2 = CI(m1.group(1)), int(m1.group(2)), CI(m2.group(1)), int(m2.group(2))
    if r1 == r2:
        rs = rowsym(rows, r1, r)
    else:
        codes = OrderedDict((key(rows, x)[0], 1) for x in range(r1, r2 + 1) if key(rows, x)[0])
        n = sum(1 for x in range(r1, r2 + 1) if key(rows, x)[0])
        first = key(rows, r1)[0]
        rs = (f'{next(iter(codes))}[{n} rows]' if len(codes) == 1
              else f'block@{first or ""}{{{", ".join(list(codes)[:4])}'
                   f'{f", ... {len(codes)} variables" if len(codes) > 4 else ""}}}')
    if c1 == c2:
        t = tsym(c1, c, a.startswith('$')) or L(c1)
        return f'{rs}({t})'
    t1, t2 = tsym(c1, c, a.startswith('$')), tsym(c2, c, b.startswith('$'))
    if c1 < YCOLS[0] and t2:                                     # the row's history up to t2 (e.g. txo floor)
        return f'{rs}(..{t2})'
    if t1 and t2:
        return f'{rs}({t1}..{t2})'
    return f'{rs}[{L(c1)}:{L(c2)}]'


def symbolic(rows, r, c, f):
    if not isinstance(f, str) or not f.startswith('='):
        return 'data' if f is not None else None
    strings = STR.findall(f)
    g = STR.sub('"§"', f[1:])
    g = g.replace('_xlfn.', '').replace('_xlpm.', '')

    def sub(m):
        ref = m.group(0)
        if '!' in ref:
            return ref.replace('$', '')
        if ':' in ref:
            a, b = ref.split(':')
            return rangesym(rows, r, c, a, b)
        return cellsym(rows, r, c, ref)
    g = REF.sub(sub, g)
    g = re.sub(rf'INDEX\(ScenarioDefinition!J1:AZ{B.MT_END},<MTInputs row[^>]*>,MATCH\(scenario\(t\),ScenarioDefinition!J5:AZ5,0\)\)',
               'MTInputs[code](scenario)', g)
    # INDEX(block starting at variable X, position p) picks variable X of block p: X[p]
    g = re.sub(r'INDEX\(block@([\w.]+)\{[^}]*\}(\([^)]*\)),(<[^>]*>)(?:,0)?\)', r'\1[\3]\2', g)
    g = re.sub(r'block@\{', 'block{', g)
    g = re.sub(r'Settings!C(\d+)', lambda m: f'Settings[{SETTINGS.get(int(m.group(1)), "C" + m.group(1))}]', g)
    for s_ in strings:
        g = g.replace('"§"', s_, 1)
    return g


def segments(forms):
    out = []
    for y, fm in forms:
        if out and out[-1][2] == fm and out[-1][1] == y - 1:
            out[-1][1] = y
        else:
            out.append([y, y, fm])
    return tuple((a, b, fm) for a, b, fm in out)


def build():
    rows = load()
    sections, cur = [], None
    eqs = OrderedDict()                                          # code -> list of (row, segments, params)
    for r in range(2, len(rows)):
        a = rows[r][0] if rows[r] else None
        if r in B.BAND_ROWS and r > 2 and isinstance(a, str):
            cur = [a, []]
            sections.append(cur)
            continue
        code, f0, s0 = key(rows, r)
        if not code or cur is None:
            if code and cur is None:
                cur = ['Scenario assumptions', []]
                sections.append(cur)
            if not code:
                continue
        forms = [(YEAR[c], symbolic(rows, r, c, rows[r][c - 1])) for c in YCOLS]
        forms = [(y, fm) for y, fm in forms if fm is not None]
        pars = [(plabel(code, k), symbolic(rows, r, 4 + k, rows[r][3 + k])) for k in range(4)
                if rows[r][3 + k] is not None]
        if code not in eqs:
            eqs[code] = []
            cur[1].append(code)
        eqs[code].append((f0, s0, segments(forms), tuple(pars)))
    merged = []                                                  # sections without new variables (sector
    for title, codes in sections:                                # blocks 6-8 reuse section 5's) join the previous
        if not codes and merged:
            merged[-1][0] += f'; {title}'
        else:
            merged.append([title, codes])
    return merged, eqs


def fmt_members(ms):
    labs = [','.join(x for x in m if x) or '-' for m in ms]
    return ', '.join(labs[:12]) + (f', ... ({len(labs)} rows)' if len(labs) > 12 else '')


def mods_of(code):
    return [m for m, codes in MODS.items() if code in codes]


def write(sections, eqs):
    lines = [f'# Workbook formulas of CPAT-AI-Mitigation v{B.VERSION} (implementation reference)', '',
             'Generated from the workbook by `make_equations.py` (do not edit by hand): every formula of the '
             'Mitigation sheet, scenario 1, with cell references replaced by variable codes, the changes against '
             'legacy CPAT (section B) and the LAMBDA library (section C). For the model itself read the equations in '
             'mathematical form, `Equations.md`; this file is for checking and editing the workbook.', '',
             '**Notation.** `code[f,s](t)`: variable `code` for fuel f and sector s in year t, where f, s and t are '
             'those of the row being defined (other fuels or sectors are named, e.g. `ener[gso,rod](t-1)`); '
             '`code[n rows](t)`: the block of n rows of that variable (a lookup or sum over it); `(t-2..t+2)`: a span '
             'of years; `<label>`: a parameter of the row (columns D:G, defined under *Parameters*); `data`: a '
             'number from the data step or an input; `year`: the year of the column; named constants (e.g. '
             '`MWyMkWh`) and defined names of data tables are as on sheet `Constants` and in the name manager. '
             '2041 is the LAMBDA year: its cells call the named LAMBDAs of section C with the same arguments as the '
             'plain formulas of the other years. `(B)` after a block: the fuel column of those rows. `SecVar`, `SecSub`, '
             '`SecFuel`: the variable, sector and fuel columns of the sector blocks (sections 5-8), so '
             '`SUMIFS(block{...}(t),SecVar,"co2",SecFuel,f)` is the sum of `co2` over every subsector for fuel f. '
             '`X[<price position>](t)`: variable X of the price block of the row\'s price fuel and sector (section 2).',
             '']
    n = 0
    index = {}
    lines.append('## A. Equations in model order')
    lines.append('')
    for title, codes in sections:
        if not codes:
            continue
        lines.append(f'### {title}')
        lines.append('')
        for code in codes:
            n += 1
            index[code] = n
            v = VARINFO.get(code)
            lab, unit = label_unit(code)
            lines.append(f'**E{n}. `{code}`: {lab}** ({unit or "-"})')
            lines.append('')
            if v and v[4]:
                lines.append(f'- Description: {v[4]}')
            if v and v[5]:
                lines.append(f'- Legacy: {v[5]}')
            ms = mods_of(code)
            if ms:
                lines.append(f'- Changes against legacy: {", ".join(ms)} (section B)')
            groups = OrderedDict()
            for f0, s0, segs, pars in eqs[code]:
                groups.setdefault(segs, []).append(((f0, s0), pars))
            shown = list(groups.items())
            for segs, members in shown[:6]:
                if len(groups) > 1 or any(m[0] != (None, None) for m in members):
                    lines.append(f'- Rows: {fmt_members([m[0] for m in members])}')
                if not segs:
                    lines.append('  - No values in the year columns: an input row left blank (or a row of '
                                 'parameters only, columns D:G)')
                for y0, y1, fm in segs:
                    yrs = f'{y0}' if y0 == y1 else f'{y0}-{y1}'
                    lines.append(f'  - {yrs}: `{code}(t) = {fm}`')
                pars = members[0][1]
                if pars:
                    lines.append('  - Parameters: ' + '; '.join(f'`<{p}>` = `{q}`' for p, q in pars))
            if len(shown) > 6:
                lines.append(f'- ... and {len(shown) - 6} further row patterns (see the workbook)')
            lines.append('')
    lines += ['## B. Changes against legacy CPAT', '',
              'Each change is described in the workbook sheet `LegacyDiff` (LD) or in the register of bugs, issues '
              'and upgrades (`Legacy_CPAT_Bugs_v<version>.docx` and the tracker: L legacy bug fixed, M bug of an '
              'earlier prototype version). Equations: the numbers of section A it touches.', '',
              '| ID | Area | Type | Equations |', '|---|---|---|---|']
    for k, d in enumerate(B.LEGACY_DIFFS, 1):
        i = f'LD{k}'
        e = ', '.join(f'E{index[c]}' for c in MODS.get(i, []) if c in index) or '-'
        lines.append(f'| {i} | {d[0]}: {d[2]} | {d[5]} | {e} |')
    reg = [(f'L{i + 1}', t, 'Legacy bug (fixed)') for i, (_k, t, *_r) in enumerate(R.BUGS)]
    reg += [(i, t, 'Legacy bug') for i, t, *_r in R.LEGACY]
    reg += [(i, t, 'prototype bug (fixed)') for i, t, *_r in R.MVP_BUGS]
    for i, t, typ in reg:
        e = ', '.join(f'E{index[c]}' for c in MODS.get(i, []) if c in index) or '-'
        lines.append(f'| {i} | {t} | {typ} | {e} |')
    lines += ['', '## C. LAMBDA library (the 2041 cells)', '',
              'Each 2041 cell calls one of these named LAMBDAs (sheet `Lambdas`, name manager).', '',
              '| LAMBDA | Parameters | Body | Note |', '|---|---|---|---|']
    for name, (params, body) in B.LAMBDAS.items():
        note = B.LAMBDA_NOTES.get(name, '').replace('|', '/')
        lines.append(f'| `{name}` | {", ".join(params)} | `{body.replace("|", "/")}` | {note} |')
    lines += ['', '## D. Index of variables', '',
              ' · '.join(f'`{c}` E{index[c]}' for c in sorted(index, key=str.lower)), '']
    with open(OUT, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(lines))
    return n, index


def main():
    sections, eqs = build()
    n, index = write(sections, eqs)
    missing = [RN(v[0]) for v in B.VARIABLES if RN(v[0]) not in index]
    print(f'wrote {OUT}: {n} variables; dictionary codes without a Mitigation row: {len(missing)}')
    return n, missing


if __name__ == '__main__':
    main()
