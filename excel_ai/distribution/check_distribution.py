"""Check of the standalone distribution module (Plan_Modules M3). Report: check_distribution_report.md.

1. Distribution: every cell of the module (recalculated by LibreOffice when it was built) equals the prototype's
   Distribution tab (its shipped values): same formulas, same values.
2. FromMitigation = the prototype's section 19a (ExportDistn) as Update from Mitigation copies it (the macro's row
   mapping emulated: years to row 1, scenario numbers to row 5, the data rows from row 8), and the variables of
   column A match row by row (the macro's layout check).
3. The kept tabs other than FromMitigation and Settings have the prototype's formulas; no formula reads a dropped tab;
   every defined name a formula reads exists.
4. The VBA project: module CPATDistribution = CPATDistribution.bas; identifiers declared (as the prototype's check).
5. Lint (Excel would not repair the file).
6. End to end: the live scenario of the prototype changed (carbon price 60 USD/t instead of 20), the prototype
   recalculated (LibreOffice), the module updated from it (the macro's copy, emulated) and recalculated: the module's
   Distribution equals the prototype's, and differs from the shipped values (the test moves the results).
"""
import os as _os, sys as _sys  # the excel_ai code folders and data/scripts on the path (core/tools/xlpaths.py)
_xl = _os.path.dirname(_os.path.abspath(__file__))
while _os.path.basename(_xl) != 'excel_ai':
    _xl = _os.path.dirname(_xl)
_sys.path.insert(0, _os.path.join(_xl, 'core', 'tools'))
import xlpaths  # noqa: E402,F401
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_distribution as D                         # noqa: E402
P = D.P
REPORT = os.path.join(HERE, 'check_distribution_report.md')
LINES = []
FAILS = []


def log(line, ok=None):
    if ok is not None:
        line += ' OK' if ok else ' FAIL'
        if not ok:
            FAILS.append(line)
    print(line, flush=True)
    LINES.append(line)


def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))
    return (a if a not in ('', None) else None) == (b if b not in ('', None) else None)


def end_to_end(mf, mv, pf, price=60):
    import tempfile
    import check_full as C
    import lo_session as LO
    sd = pf['ScenarioDefinition']
    rows = [r for r in range(1, sd.max_row + 1) if sd.cell(r, 8).value in ('CPLevelStart', 'CPLevelTarget')]
    for r in rows:
        sd.cell(r, 11).value = price                   # column K: the live scenario (Mitigation group 2)
    tmp = tempfile.mkdtemp()
    proto = os.path.join(tmp, 'proto.xlsx')
    pf.save(proto)
    pvals, _ = LO.start_values(LO.expanded_zip(proto, C.LAMBDA_CALL, C.expand), {'Mitigation', 'Distribution'})()

    class Book:                                        # export_values reads wbv[sheet].cell() and its defined names
        defined_names = pf.defined_names

        def __getitem__(self, k):
            return pvals[k]
    mod = openpyxl.load_workbook(D.OUT)
    fm = mod['FromMitigation']
    for r, vals in D.export_values(Book()).items():
        for c, v in enumerate(vals, P.COL_G0):
            fm.cell(r, c).value = v
    path = os.path.join(tmp, 'module.xlsx')
    mod.save(path)
    mvals, _ = LO.start_values(LO.expanded_zip(path, C.LAMBDA_CALL, C.expand), {'Distribution'})()
    a, b, s0 = mvals['Distribution'], pvals['Distribution'], mv['Distribution']
    diff, moved = [], 0
    for r in range(1, s0.max_row + 1):
        for c in range(1, s0.max_column + 1):
            x, y = a.cell(r, c).value, b.cell(r, c).value
            if not same(x, y):
                diff.append((s0.cell(r, c).coordinate, x, y))
            moved += isinstance(x, (int, float)) and not same(x, s0.cell(r, c).value)
    log(f'- end to end (live carbon price {price} USD/t on rows {rows} of ScenarioDefinition): module updated from '
        f'the recalculated prototype = the prototype\'s Distribution: {len(diff)} differ {diff[:5]}; {moved} values '
        'moved from the shipped ones', not diff and moved > 100)


def main():
    log(f'# Check of {D.NAME} v{D.VERSION}')
    log(f'Module `{os.path.basename(D.OUT)}`; prototype `{os.path.basename(D.SOURCE)}` v{P.VERSION}.')
    log('')
    mf, mv = openpyxl.load_workbook(D.OUT), openpyxl.load_workbook(D.OUT, data_only=True)
    pf, pv = openpyxl.load_workbook(D.SOURCE), openpyxl.load_workbook(D.SOURCE, data_only=True)
    log(f'- tabs: {", ".join(mf.sheetnames)}', mf.sheetnames == ['Cover'] + D.TABS)
    # 1. Distribution values
    a, b = mv['Distribution'], pv['Distribution']
    n, diff = 0, []
    for r in range(1, max(a.max_row, b.max_row) + 1):
        for c in range(1, max(a.max_column, b.max_column) + 1):
            x, y = a.cell(r, c).value, b.cell(r, c).value
            n += x is not None
            if not same(x, y):
                diff.append((a.cell(r, c).coordinate, x, y))
    log(f'- Distribution: {n} cells, values equal to the prototype\'s: {len(diff)} differ {diff[:5]}', not diff)
    # 2. FromMitigation = ExportDistn, mapped as the macro does
    rng = pf.defined_names['ExportDistn'].attr_text
    m = re.match(r"(\w+)!\$A\$(\d+):\$([A-Z]+)\$(\d+)", rng)
    ws, r0, last, r1 = pv[m.group(1)], int(m.group(2)), P.CI_(m.group(3)), int(m.group(4))
    fm, fmf = mv['FromMitigation'], mf['FromMitigation']
    rows = [1, 5] + list(range(P.FM_R0, P.FM_R0 + r1 - r0 - 1))
    bad = []
    for i, fr in enumerate(rows):
        for c in range(P.COL_G0, P.CI_('IZ') + 1):
            want = ws.cell(r0 + i, c).value if c <= last else None
            if not same(fm.cell(fr, c).value, want):
                bad.append((fm.cell(fr, c).coordinate, fm.cell(fr, c).value, want))
        if i >= 2 and ws.cell(r0 + i, 1).value != fmf.cell(fr, 1).value:
            bad.append((f'A{fr}', fmf.cell(fr, 1).value, ws.cell(r0 + i, 1).value))
    log(f'- FromMitigation = ExportDistn ({rng}, {len(rows)} rows) through the macro\'s row mapping: {len(bad)} '
        f'differ {bad[:5]}', not bad and len(rows) - 2 == max(P.FM_ROW.values()) - P.FM_R0 + 1)
    nf = sum(isinstance(fmf.cell(r, c).value, str) and fmf.cell(r, c).value.startswith('=')
             for r in range(1, fmf.max_row + 1) for c in range(1, fmf.max_column + 1))
    log(f'- FromMitigation holds values only: {nf} formulas', nf == 0)
    # 3. formulas
    keep = set(mf.sheetnames)
    names = set(mf.defined_names.keys())
    changed, reads, missing = [], [], set()
    known = set(P.LAMBDAS) | {'TRUE', 'FALSE'} | keep
    for t in mf.sheetnames:
        for row in mf[t].iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith('=')):
                    continue
                if D.sheets_in(v) - keep:
                    reads.append(f'{t}!{c.coordinate}')
                for tok in D.tokens(v):
                    if re.fullmatch(r'[A-Z]{1,3}\d*|_xlfn\.\w+|_xlws\.\w+', tok) or tok in names or tok in known:
                        continue
                    if not re.search(rf'(?<![\w.]){re.escape(tok)}\s*\(', v):
                        missing.add(tok)
                if t not in ('Cover', 'Settings', 'FromMitigation') and pf[t][c.coordinate].value != v:
                    changed.append(f'{t}!{c.coordinate}')
    log(f'- kept tabs (but Cover, Settings, FromMitigation): formulas as the prototype\'s: {len(changed)} differ '
        f'{changed[:5]}', not changed)
    log(f'- formulas reading a dropped tab: {len(reads)} {reads[:5]}', not reads)
    lam = set(P.LAMBDAS)
    missing = {t for t in missing if t not in lam and not t.lower().startswith(('lambda', 'let'))}
    log(f'- names read by formulas and not defined: {sorted(missing) or "none"}', not missing)
    # 4. VBA
    from oletools.olevba import VBA_Parser
    vp = VBA_Parser(D.OUT)
    mods = {nm.rsplit('.', 1)[0]: code for (_f, _s, nm, code) in vp.extract_macros()}
    src = open(D.BAS, encoding='cp1252').read().replace('\r\n', '\n')
    got = mods.get('CPATDistribution', '').replace('\r\n', '\n')
    codes = {w.sheet_properties.codeName for w in mf.worksheets}
    log(f'- VBA project: module CPATDistribution = CPATDistribution.bas ({len(src.splitlines())} lines), document '
        f'modules match the sheet code names:', got.strip() == src.strip() and codes | {'ThisWorkbook'} ==
        set(mods) - {'CPATDistribution'})
    import check_quick as Q
    Q.B.BAS, Q.log = D.BAS, log
    Q.VBA_KNOWN |= {'empty', 'filedatetime'}             # VBA built-ins the prototype's module does not use
    Q.VBA_NAMED_TYPES |= {'filename', 'readonly', 'savechanges', 'updatelinks'}   # named arguments
    Q.vba_names()
    procs = re.findall(r'^Public (?:Sub|Function) (\w+)', src, flags=re.M)
    vbanames = set(re.findall(r'Names\("(\w+)"\)', src))
    log(f'- public macros: {", ".join(procs)}; workbook names the macros read ({", ".join(sorted(vbanames))}) '
        'defined:', vbanames <= names)
    log(f'- buttons on Cover call public macros:', {m_ for _l, m_ in D.MACRO_BUTTONS} <= set(procs))
    end_to_end(mf, mv, pf)
    # 5. lint
    import excel_lint
    problems = excel_lint.lint(D.OUT)
    log(f'- lint: {problems or "no problem"}', not problems)
    log('')
    log(f'**Result: {"PASS" if not FAILS else "FAIL (" + str(len(FAILS)) + ")"}**')
    open(REPORT, 'w').write('\n'.join(LINES) + '\n')
    return not FAILS


if __name__ == '__main__':
    sys.exit(0 if main() else 1)
