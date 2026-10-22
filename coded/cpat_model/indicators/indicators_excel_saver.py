from typing import Dict, List
from io import BytesIO

import pandas as pd
from matplotlib.figure import Figure

from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.drawing.image import Image

from cpat_model.indicators.indicators import MitIndicators
from cpat_model.indicators.utils.generate_excel_columns import index_to_excel_column

from cpat_model.constants import BASELINE_SCENARIO

from cpat_shiny.mapping import ALL_COUNTRIES
from cpat_shiny.server_logic.visualization.graphs_processors.constants import (
    XLSX_HEADERS, GRAPH_X_OFFSET, GRAPH_Y_OFFSET
)

from cpat_shiny.server_logic.visualization.graphs_titles import GRAPHS_TITLES


class IndicatorsExcelSaver:
    def __init__(
            self,
            indicators: MitIndicators,
            graphs: Dict[str, Figure],
            file_path: str,
            module_name: str
            ) -> None:
        self.__selected_countries = indicators.selected_countries
        self.__user_scenarios = [s for s in indicators.scenarios if s != BASELINE_SCENARIO]

         # pylint: disable=abstract-class-instantiated
        with pd.ExcelWriter(
            file_path,
            engine="openpyxl"
        ) as results_writer:
            indicators.grouped_indicators.to_excel(results_writer, sheet_name='Indicators ' + module_name)

            wb = results_writer.book
            self.__ws: Worksheet = wb.create_sheet(title=module_name)
            self.__create_headers()
            self.__save_graphs(indicators, graphs)

    def __create_headers(
            self
            ) -> None:
        header_font = Font(bold=True, color="FFFFFF")
        for cell_no, title in XLSX_HEADERS.items():
            cell = self.__ws[cell_no]
            cell.value = title
            cell.font = header_font

        fill = PatternFill(start_color='92D050', end_color='92D050',fill_type="solid")
        rows_to_fill = [int(''.join(filter(str.isdigit, key))) for key in XLSX_HEADERS]
        for row_number in rows_to_fill:
            for col in range(1, 101): # colors first 100 cols
                self.__ws.cell(row=row_number, column=col).fill = fill

    def __save_graphs(
            self,
            indicators: MitIndicators,
            graphs: Dict[str, Figure]
            ):
        if 'ec' in indicators.selected_indicators and 'dec' in indicators.selected_indicators:
            graph_types = ['ec_by_sector', 'ec_by_fuel', 'ec_change_by_fuel']
            self.__save_graphs_group(graphs, graph_types, graphs_rows_tracker=2)
        if 'p' in indicators.selected_indicators:
            graph_types = ['domestic_p']
            self.__save_graphs_group(graphs, graph_types, graphs_rows_tracker=82)
        if 'em' in indicators.selected_indicators:
            graph_types = ['em_by_sector']
            self.__save_graphs_group(graphs, graph_types, graphs_rows_tracker=110)

    def __save_graphs_group(
            self,
            graphs: Dict[str, Figure],
            graph_types: List[str],
            graphs_rows_tracker: int
            ) -> None:
        column_no = 1
        for country in self.__selected_countries:
            for scenario_name in self.__user_scenarios:
                format_title_params = {
                    'scenario_name': scenario_name,
                    'country': ALL_COUNTRIES[country.upper()]
                }
                column = index_to_excel_column(column_no)
                for i, graph_type in enumerate(graph_types):
                    if graph_type == 'domestic_p':
                        title = GRAPHS_TITLES[graph_type].format(
                            **format_title_params,
                            reference_year=2022 #TODO softcode it
                        )
                    else:
                        title = GRAPHS_TITLES[graph_type].format(**format_title_params)
                    self.__save_plot(
                        graphs[title],
                        column,
                        (graphs_rows_tracker + i * GRAPH_Y_OFFSET)
                    )
                column_no += GRAPH_X_OFFSET

    def __save_plot(
            self,
            fig: Figure,
            column: str,
            starting_row: int
            ) -> None:
        img_data = BytesIO()
        fig.savefig(img_data, format='png')
        img_data.seek(0)
        img = Image(img_data)
        self.__ws.add_image(img, f'{column}{starting_row}')
