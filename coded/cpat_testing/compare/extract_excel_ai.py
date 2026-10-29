"""Extract from the shipped Excel-AI workbook (its stored values; no recalculation, D8 as decided 2026-10-29):

  data/inputs/scenarios_excel_ai.csv                 its ScenarioDefinition: every MTInputs row, scenarios 1-6,
                                                     plus the Python rows below (cpat_model.inputs.scenarios)
  cpat_testing/compare/reference/excel_ai_<v>.csv    every coded row of sheet Mitigation, scenario groups 1 and 2,
                                                     years as columns: the reference values of the comparison
  data/mvp_legacy_extracts/module_<Tab>_cells.csv    the typed module values of tabs FromAirPollution and
                                                     FromTransportCB (row, col, value), for the placeholder modules

Run from coded/:  python cpat_testing/compare/extract_excel_ai.py
"""
import csv
import sys
from pathlib import Path

import openpyxl

CODED = Path(__file__).resolve().parents[2]
REPO = CODED.parent
sys.path.insert(0, str(CODED))
from cpat_model import EXCEL_AI_VERSION  # noqa: E402
from cpat_model.inputs.scenarios import PYTHON_ROWS  # noqa: E402

WORKBOOK = REPO / 'excel_ai' / f'CPAT-AI-Mitigation-Prototype_v{EXCEL_AI_VERSION}.xlsm'
SCENARIOS = REPO / 'data' / 'inputs' / 'scenarios_excel_ai.csv'
REFERENCE = CODED / 'cpat_testing' / 'compare' / 'reference' / f'excel_ai_{EXCEL_AI_VERSION}.csv'
SD_FIRST, SD_LAST, SD_COLS = 8, 600, range(10, 16)          # ScenarioDefinition rows, columns J-O (scenarios 1-6)
GROUPS = [(11, 1), (33, 2)]                                  # Mitigation: code column K, AG; years follow (20)


def text(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)
    return str(v)


def scenarios(wb):
    ws = wb['ScenarioDefinition']
    rows = list(ws.iter_rows(min_row=1, max_row=SD_LAST, values_only=True))
    names = [text(rows[5][c - 1]) for c in SD_COLS]            # row 6: scenario names
    out = [['', '_ScenarioName', 'Scenario name', ''] + names,
           ['', '_Countries', 'Countries', 'ISO3'] + ['EGY'] * len(SD_COLS)]
    last = SD_FIRST
    for i in range(SD_FIRST - 1, len(rows)):
        r = rows[i]
        if any(v not in (None, '') for v in r[:15]):
            last = i + 1
    for i in range(SD_FIRST - 1, last):
        r = rows[i]
        out.append([str(i + 1), text(r[7]), text(r[1]), text(r[2])] + [text(r[c - 1]) for c in SD_COLS])
    for name, label, unit, default in PYTHON_ROWS:
        out.append(['', name, label, unit] + [default] * len(SD_COLS))
    head = ['row', 'name', 'label', 'unit'] + [str(n) for n in range(1, len(SD_COLS) + 1)]
    with open(SCENARIOS, 'w', encoding='utf-8', newline='') as f:
        csv.writer(f, lineterminator='\n').writerows([head] + out)
    return len(out)


def reference(wb):
    ws = wb['Mitigation']
    rows = list(ws.iter_rows(values_only=True))
    years = [int(y) for y in rows[0][11:31]]
    out, seen = [], set()
    for r in rows:
        for col, _n in GROUPS:
            code = r[col - 1]
            if isinstance(code, str) and code.count('.') >= 2 and code not in seen:   # section 19 repeats codes
                seen.add(code)
                out.append([code] + [text(v) for v in r[col:col + 20]])
    with open(REFERENCE, 'w', encoding='utf-8', newline='') as f:
        csv.writer(f, lineterminator='\n').writerows([['code'] + years] + out)
    return len(out)


def module_tabs(wb):
    for tab in ('FromAirPollution', 'FromTransportCB'):
        cells = [[r, c, text(v)] for r, row in enumerate(wb[tab].iter_rows(values_only=True), 1)
                 for c, v in enumerate(row, 1) if v not in (None, '')]
        out = REPO / 'data' / 'mvp_legacy_extracts' / f'module_{tab}_cells.csv'
        with open(out, 'w', encoding='utf-8', newline='') as f:
            csv.writer(f, lineterminator='\n').writerows([['row', 'col', 'value']] + cells)
        print(f'{tab}: {len(cells)} cells -> {out}')


if __name__ == '__main__':
    book = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    print('scenario rows:', scenarios(book), '->', SCENARIOS)
    print('reference codes:', reference(book), '->', REFERENCE)
    module_tabs(book)
