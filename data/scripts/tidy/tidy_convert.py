"""Convert CPAT data between the legacy layout and tidy CSV, both ways (data/INTERFACE.md).

    legacy layout   the data tabs of legacy CPAT, as Excel-AI is built from them: either a folder of legacy CSV files
                    (data/mvp_legacy_extracts/: legacy_<Tab>.csv tables, legacy_<Tab>_cells.csv tabs at their cell
                    addresses) or a legacy-style workbook (.xlsx, one sheet per legacy tab, each table at its legacy
                    header row, each cells tab at its legacy addresses, so it pastes onto legacy CPAT)
    tidy CSV        the input of Coded CPAT: one CSV per table, long, one observation per row, standard key names
                    (`country` = ISO3 upper case, `year`, `code`, ...), plus `_layout/` (what is not data: headers,
                    names, row order, labels and notes of the legacy tabs) and `_manifest.csv`

legacy -> tidy -> legacy gives back every cell exactly (the same text in the CSV form; the same values in the
workbook form): `check` proves it, and data/tests/test_tidy_roundtrip.py runs it. Rows added in tidy form (a new
country) are written after the rows the layout knows. Table specs: tidy_specs.py; cells-tab blocks: tidy_blocks.py.

Run from the repository root:
    python data/scripts/tidy/tidy_convert.py to-tidy   --legacy data/mvp_legacy_extracts --out data/tidy
    python data/scripts/tidy/tidy_convert.py to-legacy --tidy data/tidy --out <folder>          # legacy CSV files
    python data/scripts/tidy/tidy_convert.py to-legacy --tidy data/tidy --xlsx <file.xlsx>      # legacy-style workbook
    python data/scripts/tidy/tidy_convert.py to-tidy   --legacy <file.xlsx> --out <folder>      # from a workbook
    python data/scripts/tidy/tidy_convert.py check     --legacy data/mvp_legacy_extracts       # round trip, both forms
Real (proprietary) data: point --legacy/--out at folders outside this repository (root NORMS section 8).
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from tidy_specs import CELLS, LAYOUT_FILES, TABLES  # noqa: E402

csv.field_size_limit(10**8)


# ------------------------------------------------------------------------------------------------ small helpers
def read_csv(path):
    with open(path, encoding='utf-8', newline='') as f:
        return list(csv.reader(f))


def write_csv(path, header, rows):
    os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
    with open(path, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(header)
        w.writerows(rows)


def tidy_country(v, spec):
    return v.upper() if spec.get('country') else v


def legacy_country(v, spec):
    return v.lower() if spec.get('country', ('', ''))[1] == 'lower' else v


def cell_text(v):
    """A workbook value as the text the legacy CSV files hold (whole numbers without .0, as the extractors write)."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)
    return str(v)


def cell_value(s):
    """Text -> number where the text is a number (for the workbook form), else the text."""
    if s == '':
        return None
    try:
        if s.lstrip('-').isdigit() and not (len(s.lstrip('-')) > 1 and s.lstrip('-')[0] == '0'):
            return int(s)
        f = float(s)
        return f if repr(f) == s or str(f) == s or s.lower().count('e') == 1 else s
    except ValueError:
        return s


# ------------------------------------------------------------------------------------------------ tables
def table_to_tidy(spec, header, rows):
    """A legacy table -> (tidy header, tidy rows, layout)."""
    pos = {h: j for j, h in enumerate(header)}
    for name in list(spec['keys']) + spec.get('attrs', []):
        assert name in pos, f"{spec['tidy']}: column {name!r} not in the legacy header"
    assert len(set(header)) == len(header), f"{spec['tidy']}: duplicate legacy column names"
    key_pos = [pos[k] for k in spec['keys']]
    attr_pos = [pos[a] for a in spec.get('attrs', [])]
    val_pos = [j for j in range(len(header)) if j not in key_pos and j not in attr_pos]
    tkeys = list(spec['keys'].values())
    skip = set(spec.get('skip', []))
    melt = spec.get('melt')
    rename = spec.get('rename', {})
    if melt:
        thead = tkeys + list(melt[:-1] if len(melt) == 3 else melt) + ['value']
    else:
        thead = tkeys + [rename.get(header[j], header[j]) for j in val_pos]
    trows, lay_rows, skipped, seen = [], [], {}, {}
    for i, r in enumerate(rows):
        r = r + [''] * (len(header) - len(r))
        if r[key_pos[0]] in skip:
            skipped[i] = r
            continue
        k = [r[j] for j in key_pos]
        tk = [tidy_country(v, spec) if spec.get('country') and header[j] == spec['country'][0] else v
              for v, j in zip(k, key_pos)]
        if tuple(tk) in seen:                                # an exact repeat of an earlier row: layout only
            assert r == seen[tuple(tk)], f"{spec['tidy']}: key {tk} is not unique"
            skipped[i] = r
            continue
        seen[tuple(tk)] = r
        lay_rows.append([i] + k + [r[j] for j in attr_pos])
        if melt:
            for j in val_pos:
                if r[j] != '':
                    lab = header[j].split(melt[2], 1) if len(melt) == 3 else [header[j]]
                    trows.append(tk + lab + [r[j]])
        else:
            trows.append(tk + [r[j] for j in val_pos])
    layout = dict(kind='table', header=header, key_pos=key_pos, attr_pos=attr_pos, val_pos=val_pos,
                  skipped={str(i): r for i, r in skipped.items()}, nrows=len(rows))
    return thead, trows, layout, lay_rows


def table_from_tidy(spec, thead, trows, layout, lay_rows):
    """Tidy rows + layout -> the legacy table (header, rows)."""
    header = layout['header']
    key_pos, attr_pos, val_pos = layout['key_pos'], layout['attr_pos'], layout['val_pos']
    nk = len(key_pos)
    melt = spec.get('melt')
    rename = spec.get('rename', {})
    col_of = {}
    if melt:
        for j in val_pos:
            col_of[tuple(header[j].split(melt[2], 1)) if len(melt) == 3 else (header[j],)] = j
    else:
        tcols = thead[nk:]
        back = {rename.get(header[j], header[j]): j for j in val_pos}
        assert set(tcols) == set(back), f"{spec['tidy']}: tidy columns differ from the layout"
    values = {}                                              # tidy key -> {legacy column: text}
    order = []
    for t in trows:
        k = tuple(t[:nk])
        if k not in values:
            values[k] = {}
            order.append(k)
        if melt:
            lab = tuple(t[nk:-1])
            assert lab in col_of, f"{spec['tidy']}: {lab} is not a column of the legacy table"
            values[k][col_of[lab]] = t[-1]
        else:
            for name, v in zip(thead[nk:], t[nk:]):
                values[k][back[name]] = v
    out = [None] * layout['nrows']
    for i, r in layout['skipped'].items():
        out[int(i)] = r
    known = set()
    for lr in lay_rows:
        i, k, attrs = int(lr[0]), lr[1:1 + nk], lr[1 + nk:]
        tk = tuple(tidy_country(v, spec) if spec.get('country') and header[j] == spec['country'][0] else v
                   for v, j in zip(k, key_pos))
        known.add(tk)
        row = [''] * len(header)
        for v, j in zip(k, key_pos):
            row[j] = v
        for v, j in zip(attrs, attr_pos):
            row[j] = v
        for j, v in values.get(tk, {}).items():
            row[j] = v
        out[i] = row
    for k in order:                                          # rows added in tidy form
        if k in known:
            continue
        row = [''] * len(header)
        for v, j in zip(k, key_pos):
            row[j] = legacy_country(v, spec) if spec.get('country') and header[j] == spec['country'][0] else v
        for j, v in values[k].items():
            row[j] = v
        out.append(row)
    return header, [r for r in out if r is not None]


# ------------------------------------------------------------------------------------------------ cells tabs
def _row_key(at, r, b):
    """The row keys of a block row; a row with blank key cells is named by its position (row<r>)."""
    if not b['key_cols']:                            # a one-row block: no row keys
        return []
    key = [at[(r, kc)][2] if (r, kc) in at else '' for kc in b['key_cols']]
    return key if any(key) else [f'row{r}'] + [''] * (len(key) - 1)


def _col_label(at, c, b):
    """The column label of a block column: the block's own name for it (`labels`, where the legacy header is blank or
    repeated), else the label cell; a blank label cell is named by its position (col<c>)."""
    if c in b.get('labels', {}):
        return b['labels'][c]
    return at[(b['label_row'], c)][2] if (b['label_row'], c) in at else f'col{c}'


def cells_to_tidy(spec, head, cells):
    """A legacy tab at its addresses (cells: [row, col, value, ...extra]) -> {tidy name: (header, rows)}, layout."""
    at = {(int(c[0]), int(c[1])): c for c in cells}
    used, tables = set(), {}
    labels = set()                                   # key and label cells of every block stay in the layout
    for b in spec.get('blocks', []):
        labels |= {(r, kc) for r in range(b['rows'][0], b['rows'][1] + 1) for kc in b['key_cols']}
        labels |= {(b['label_row'], c) for c in range(b['cols'][0], b['cols'][1] + 1)}
    for b in spec.get('blocks', []):
        (r0, r1), (c0, c1) = b['rows'], b['cols']
        trows = []
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if c in b['key_cols'] or (r, c) not in at or (r, c) in labels or (r, c) in used:
                    continue
                key = _row_key(at, r, b)
                lab = _col_label(at, c, b)
                trows.append(key + [lab, at[(r, c)][2]])
                used.add((r, c))
        tables[b['name']] = (b['keys'] + [b['label'], 'value'], trows)
    residual = [c for c in cells if (int(c[0]), int(c[1])) not in used]
    extra = {f'{r},{c}': at[(r, c)][3:] for (r, c) in used if len(at[(r, c)]) > 3 and any(at[(r, c)][3:])}
    order = [[int(c[0]), int(c[1])] for c in cells]
    layout = dict(kind='cells', header=head, extra=extra,
                  order=None if order == sorted(order) else order)
    return tables, layout, residual


def cells_from_tidy(spec, tables, layout, residual):
    head = layout['header']
    width = len(head)
    at = {(int(c[0]), int(c[1])): c for c in residual}
    out = list(residual)
    for b in spec.get('blocks', []):
        (r0, r1), (c0, c1) = b['rows'], b['cols']
        row_of = {}
        for r in range(r0, r1 + 1):
            row_of.setdefault(tuple(_row_key(at, r, b)), r)
        col_of = {_col_label(at, c, b): c for c in range(c0, c1 + 1) if c not in b['key_cols']}
        _thead, trows = tables[b['name']]
        nk = len(b['key_cols'])
        for t in trows:
            r, c = row_of.get(tuple(t[:nk])), col_of.get(t[nk])
            assert r and c, f"{b['name']}: no place in the legacy tab for {t[:nk + 1]}"
            extra = layout['extra'].get(f'{r},{c}', [''] * (width - 3))
            out.append([str(r), str(c), t[-1]] + list(extra)[:width - 3])
    if layout.get('order'):
        pos = {(r, c): i for i, (r, c) in enumerate(layout['order'])}
        out.sort(key=lambda x: pos.get((int(x[0]), int(x[1])), len(pos)))
    else:
        out.sort(key=lambda x: (int(x[0]), int(x[1])))
    return head, out


# ------------------------------------------------------------------------------------------------ legacy side I/O
def read_legacy(src):
    """{file name: (header, rows)} for every spec'd file found in a legacy CSV folder or a legacy-style workbook."""
    found = {}
    if os.path.isdir(src):
        for name in list(TABLES) + list(CELLS) + list(LAYOUT_FILES):
            p = os.path.join(src, name)
            if not os.path.exists(p):                        # a file kept beside the folder (range_anxiety_fit)
                p = os.path.join(src, os.path.basename(name))
            if os.path.exists(p):
                rows = read_csv(p)
                found[name] = (rows[0], rows[1:])
        return found
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    for name, spec in TABLES.items():
        if spec['sheet'] not in wb.sheetnames:
            continue
        vals = [[cell_text(v) for v in r] for r in wb[spec['sheet']].iter_rows(values_only=True)]
        h = spec['header_row'] - 1
        head = vals[h]
        n = max(j for j, v in enumerate(head) if v != '') + 1 if any(head) else 0
        body = [(r + [''] * n)[:n] for r in vals[h + 1:]]
        while body and not any(body[-1]):
            body.pop()
        found[name] = (head[:n], body)
    for name, spec in CELLS.items():
        if spec['sheet'] not in wb.sheetnames:
            continue
        cells = []
        for r, row in enumerate(wb[spec['sheet']].iter_rows(values_only=True), 1):
            for c, v in enumerate(row, 1):
                if v not in (None, '') and r >= spec.get('first_row', 1):
                    cells.append([str(r), str(c), cell_text(v)] + [''] * (len(spec['header']) - 3))
        found[name] = (spec['header'], cells)
    return found


def write_legacy_csv(out_dir, tables):
    """Every file into out_dir (a file the repository keeps beside the folder, e.g. ../range_anxiety_fit/, goes into
    out_dir under its own name)."""
    for name, (head, rows) in tables.items():
        write_csv(os.path.join(out_dir, os.path.basename(name)), head, rows)


def write_legacy_xlsx(path, tables):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (head, rows) in tables.items():
        if name in TABLES:
            spec = TABLES[name]
            ws = wb.create_sheet(spec['sheet'])
            h = spec['header_row']
            for j, v in enumerate(head, 1):
                if v != '':
                    ws.cell(h, j, v)
            for i, r in enumerate(rows, h + 1):
                for j, v in enumerate(r, 1):
                    if v != '':
                        ws.cell(i, j, cell_value(v))
        elif name in CELLS:
            ws = wb.create_sheet(CELLS[name]['sheet'])
            for c in rows:
                ws.cell(int(c[0]), int(c[1]), cell_value(c[2]))
    wb.save(path)


# ------------------------------------------------------------------------------------------------ tidy side I/O
def to_tidy(legacy, out):
    manifest = []
    for name, (head, rows) in legacy.items():
        lay = os.path.join(out, '_layout')
        if name in TABLES:
            spec = TABLES[name]
            thead, trows, layout, lay_rows = table_to_tidy(spec, head, rows)
            layout['source'] = name
            write_csv(os.path.join(out, f"{spec['tidy']}.csv"), thead, trows)
            write_csv(os.path.join(lay, f"{spec['tidy']}.rows.csv"),
                      ['_row'] + [head[j] for j in layout['key_pos'] + layout['attr_pos']], lay_rows)
            with open(os.path.join(lay, f"{spec['tidy']}.json"), 'w', encoding='utf-8') as f:
                json.dump(layout, f, ensure_ascii=False)
            manifest.append([spec['tidy'], name, 'table', spec['sheet'], ' '.join(thead[:len(spec['keys'])]),
                             len(trows), spec.get('proprietary', 'no')])
        elif name in CELLS:
            spec = CELLS[name]
            blocks, layout, residual = cells_to_tidy(spec, head, rows)
            layout['source'] = name
            for bname, (bh, br) in blocks.items():
                write_csv(os.path.join(out, f'{bname}.csv'), bh, br)
                manifest.append([bname, name, 'cells block', spec['sheet'], ' '.join(bh[:-2]), len(br), 'no'])
            write_csv(os.path.join(lay, f"{spec['tidy']}.cells.csv"), head, residual)
            with open(os.path.join(lay, f"{spec['tidy']}.json"), 'w', encoding='utf-8') as f:
                json.dump(layout, f, ensure_ascii=False)
            manifest.append([spec['tidy'] + ' (layout)', name, 'cells layout', spec['sheet'], 'row col', len(residual),
                             'no'])
        elif name in LAYOUT_FILES:
            write_csv(os.path.join(lay, name), head, rows)
    write_csv(os.path.join(out, '_manifest.csv'),
              ['table', 'legacy_source', 'kind', 'legacy_tab', 'keys', 'rows', 'proprietary'], sorted(manifest))
    return manifest


def from_tidy(tdir):
    """{legacy file name: (header, rows)} from a tidy folder."""
    lay = os.path.join(tdir, '_layout')
    out = {}
    for name, spec in TABLES.items():
        p = os.path.join(lay, f"{spec['tidy']}.json")
        if not os.path.exists(p):
            continue
        layout = json.load(open(p, encoding='utf-8'))
        t = read_csv(os.path.join(tdir, f"{spec['tidy']}.csv"))
        lr = read_csv(os.path.join(lay, f"{spec['tidy']}.rows.csv"))[1:]
        out[name] = table_from_tidy(spec, t[0], t[1:], layout, lr)
    for name, spec in CELLS.items():
        p = os.path.join(lay, f"{spec['tidy']}.json")
        if not os.path.exists(p):
            continue
        layout = json.load(open(p, encoding='utf-8'))
        res = read_csv(os.path.join(lay, f"{spec['tidy']}.cells.csv"))[1:]
        tables = {}
        for b in spec.get('blocks', []):
            t = read_csv(os.path.join(tdir, f"{b['name']}.csv"))
            tables[b['name']] = (t[0], t[1:])
        out[name] = cells_from_tidy(spec, tables, layout, res)
    for name in LAYOUT_FILES:
        p = os.path.join(lay, name)
        if os.path.exists(p):
            rows = read_csv(p)
            out[name] = (rows[0], rows[1:])
    return out


# ------------------------------------------------------------------------------------------------ check
def _same(x, y, numeric):
    """Two cell texts are the same; numeric: numbers equal to 1e-15 relative (the workbook form keeps 16
    significant digits, openpyxl), whole numbers equal whatever their form (2027 = 2027.0)."""
    if x == y or not numeric:
        return x == y
    try:
        fx, fy = float(x), float(y)
    except ValueError:
        return False
    return abs(fx - fy) <= 1e-15 * max(abs(fx), abs(fy))


def compare(a, b, numeric=False, width=None):
    """Differences between two {file: (header, rows)}; numeric=True for the workbook form (see _same); width: compare
    only the first `width` columns of cells files (a values-only workbook has no formula flag)."""
    diffs = []
    for name in sorted(set(a) | set(b)):
        if name not in a or name not in b:
            diffs.append(f'{name}: only in one side')
            continue
        ha, ra = a[name]
        hb, rb = b[name]
        w = width if (width and name in CELLS) else None
        if not all(_same(x, y, numeric) for x, y in zip(ha[:w], hb[:w])) or len(ha[:w]) != len(hb[:w]):
            diffs.append(f'{name}: header differs')
        if len(ra) != len(rb):
            diffs.append(f'{name}: {len(ra)} rows against {len(rb)}')
            continue
        for i, (x, y) in enumerate(zip(ra, rb)):
            x, y = x + [''] * (len(y) - len(x)), y + [''] * (len(x) - len(y))
            if not all(_same(u, v, numeric) for u, v in zip(x[:w], y[:w])):
                j = next(j for j, (u, v) in enumerate(zip(x, y)) if not _same(u, v, numeric))
                diffs.append(f'{name}: row {i + 1}, column {j + 1}: {x[j]!r} / {y[j]!r}')
                break
    return diffs


def check(src):
    legacy = read_legacy(src)
    with tempfile.TemporaryDirectory() as tmp:
        to_tidy(legacy, os.path.join(tmp, 'tidy'))
        back = from_tidy(os.path.join(tmp, 'tidy'))
        d_csv = compare(legacy, back)
        xlsx = os.path.join(tmp, 'legacy.xlsx')
        write_legacy_xlsx(xlsx, back)
        again = from_tidy_via(read_legacy(xlsx), tmp)
        d_xlsx = compare({k: v for k, v in legacy.items() if k not in LAYOUT_FILES}, again, numeric=True, width=3)
    return d_csv, d_xlsx


def from_tidy_via(legacy, tmp):
    """legacy (from a workbook) -> tidy -> legacy: the workbook round trip."""
    d = os.path.join(tmp, 'tidy2')
    to_tidy(legacy, d)
    return {k: v for k, v in from_tidy(d).items() if k not in LAYOUT_FILES}


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest='cmd', required=True)
    a = sub.add_parser('to-tidy')
    a.add_argument('--legacy', required=True, help='legacy CSV folder or legacy-style workbook (.xlsx)')
    a.add_argument('--out', required=True, help='tidy folder')
    b = sub.add_parser('to-legacy')
    b.add_argument('--tidy', required=True)
    b.add_argument('--out', help='legacy CSV folder')
    b.add_argument('--xlsx', help='legacy-style workbook')
    c = sub.add_parser('check')
    c.add_argument('--legacy', required=True)
    args = p.parse_args()
    if args.cmd == 'to-tidy':
        m = to_tidy(read_legacy(args.legacy), args.out)
        print(f'{len(m)} tidy tables written to {args.out}')
    elif args.cmd == 'to-legacy':
        back = from_tidy(args.tidy)
        if args.out:
            write_legacy_csv(args.out, back)
            print(f'{len(back)} legacy files written to {args.out}')
        if args.xlsx:
            write_legacy_xlsx(args.xlsx, back)
            print(f'legacy-style workbook written: {args.xlsx}')
    else:
        d_csv, d_xlsx = check(args.legacy)
        print('CSV round trip:', 'exact' if not d_csv else '\n  ' + '\n  '.join(d_csv))
        print('Workbook round trip:', 'equal' if not d_xlsx else '\n  ' + '\n  '.join(d_xlsx))
        sys.exit(1 if d_csv or d_xlsx else 0)


if __name__ == '__main__':
    main()
