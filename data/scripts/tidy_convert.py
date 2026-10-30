"""Convert CPAT data between tidy CSV (the source of truth) and the legacy-compatible workbook, both ways.

    tidy CSV        data/Public01/tidy/: one CSV per table, long, one observation per row, standard key names (`country` = ISO3
                    upper case, `year`, `code`, ...), plus `_layout/` (what is not data: headers, names, row order,
                    labels and notes of the legacy tabs) and `_manifest.csv`. Coded CPAT reads it.
    legacy xlsx     data/Public01/legacy_xlsx/CPAT_LegacyData.xlsx: one sheet per legacy tab, each table at its legacy header
                    row, each cells tab at its legacy addresses, so a block pastes onto legacy CPAT at the same address.
                    Written from tidy; never edited by hand.

tidy -> xlsx -> tidy gives back every value (numbers to 16 significant digits, openpyxl; whole numbers lose a trailing
`.0`): `check` proves it, and data/tests/test_tidy_roundtrip.py runs it. Rows added in tidy form (a new country) are
written after the rows the layout knows. Table specs: tidy_specs.py; cells-tab blocks: tidy_blocks.py.

Run from the repository root:
    python data/scripts/tidy_convert.py to-xlsx --tidy data/Public01/tidy --out data/Public01/legacy_xlsx/CPAT_LegacyData.xlsx
    python data/scripts/tidy_convert.py to-tidy --xlsx <file.xlsx> --out <folder>     # import a workbook (lossy, above)
    python data/scripts/tidy_convert.py check   --tidy data/Public01/tidy                      # round trip
Real (proprietary) data: point --tidy/--out at folders outside this repository (excel_ai/NORMS.md section 8).
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from tidy_specs import ASIS, CELLS, FILTER_COL, LAYOUT_FILES, LINKS, TABLES  # noqa: E402

csv.field_size_limit(10**8)


# ------------------------------------------------------------------------------------------------ small helpers
def read_csv(path):
    with open(path, encoding='utf-8', newline='') as f:
        return list(csv.reader(f))


def write_csv(path, header, rows):
    os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
    with open(path, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(header)
        w.writerows(rows)


def tidy_country(v, spec):
    return v.upper() if spec.get('country') else v


def legacy_country(v, spec):
    return v.lower() if spec.get('country', ('', ''))[1] == 'lower' else v


def cell_text(v):
    """A workbook value as the text the legacy CSV files hold (whole numbers without .0, as the extractors write)."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)
    return str(v)


def cell_value(s):
    """Text -> number where the text is a number (for the workbook form), else the text."""
    if s == '':
        return None
    try:
        if s.lstrip('-').isdigit() and not (len(s.lstrip('-')) > 1 and s.lstrip('-')[0] == '0'):
            return int(s)
        f = float(s)
        return f if repr(f) == s or str(f) == s or s.lower().count('e') == 1 else s
    except ValueError:
        return s


# ------------------------------------------------------------------------------------------------ tables
def table_to_tidy(spec, header, rows):
    """A legacy table -> (tidy header, tidy rows, layout)."""
    pos = {h: j for j, h in enumerate(header)}
    for name in list(spec['keys']) + spec.get('attrs', []):
        assert name in pos, f"{spec['tidy']}: column {name!r} not in the legacy header"
    assert len(set(header)) == len(header), f"{spec['tidy']}: duplicate legacy column names"
    key_pos = [pos[k] for k in spec['keys']]
    attr_pos = [pos[a] for a in spec.get('attrs', [])]
    val_pos = [j for j in range(len(header)) if j not in key_pos and j not in attr_pos]
    tkeys = list(spec['keys'].values())
    skip = set(spec.get('skip', []))
    melt = spec.get('melt')
    rename = spec.get('rename', {})
    if melt:
        thead = tkeys + list(melt[:-1] if len(melt) == 3 else melt) + ['value']
    else:
        thead = tkeys + [rename.get(header[j], header[j]) for j in val_pos]
    trows, lay_rows, skipped, seen = [], [], {}, {}
    for i, r in enumerate(rows):
        r = r + [''] * (len(header) - len(r))
        if r[key_pos[0]] in skip:
            skipped[i] = r
            continue
        k = [r[j] for j in key_pos]
        tk = [tidy_country(v, spec) if spec.get('country') and header[j] == spec['country'][0] else v
              for v, j in zip(k, key_pos)]
        if tuple(tk) in seen:                                # an exact repeat of an earlier row: layout only
            assert r == seen[tuple(tk)], f"{spec['tidy']}: key {tk} is not unique"
            skipped[i] = r
            continue
        seen[tuple(tk)] = r
        lay_rows.append([i] + k + [r[j] for j in attr_pos])
        if melt:
            for j in val_pos:
                if r[j] != '':
                    lab = header[j].split(melt[2], 1) if len(melt) == 3 else [header[j]]
                    trows.append(tk + lab + [r[j]])
        else:
            trows.append(tk + [r[j] for j in val_pos])
    layout = dict(kind='table', header=header, key_pos=key_pos, attr_pos=attr_pos, val_pos=val_pos,
                  skipped={str(i): r for i, r in skipped.items()}, nrows=len(rows))
    return thead, trows, layout, lay_rows


def table_from_tidy(spec, thead, trows, layout, lay_rows):
    """Tidy rows + layout -> the legacy table (header, rows)."""
    header = layout['header']
    key_pos, attr_pos, val_pos = layout['key_pos'], layout['attr_pos'], layout['val_pos']
    nk = len(key_pos)
    melt = spec.get('melt')
    rename = spec.get('rename', {})
    col_of = {}
    if melt:
        for j in val_pos:
            col_of[tuple(header[j].split(melt[2], 1)) if len(melt) == 3 else (header[j],)] = j
    else:
        tcols = thead[nk:]
        back = {rename.get(header[j], header[j]): j for j in val_pos}
        assert set(tcols) == set(back), f"{spec['tidy']}: tidy columns differ from the layout"
    values = {}                                              # tidy key -> {legacy column: text}
    order = []
    for t in trows:
        k = tuple(t[:nk])
        if k not in values:
            values[k] = {}
            order.append(k)
        if melt:
            lab = tuple(t[nk:-1])
            assert lab in col_of, f"{spec['tidy']}: {lab} is not a column of the legacy table"
            values[k][col_of[lab]] = t[-1]
        else:
            for name, v in zip(thead[nk:], t[nk:]):
                values[k][back[name]] = v
    out = [None] * layout['nrows']
    for i, r in layout['skipped'].items():
        out[int(i)] = r
    known = set()
    for lr in lay_rows:
        i, k, attrs = int(lr[0]), lr[1:1 + nk], lr[1 + nk:]
        tk = tuple(tidy_country(v, spec) if spec.get('country') and header[j] == spec['country'][0] else v
                   for v, j in zip(k, key_pos))
        known.add(tk)
        row = [''] * len(header)
        for v, j in zip(k, key_pos):
            row[j] = v
        for v, j in zip(attrs, attr_pos):
            row[j] = v
        for j, v in values.get(tk, {}).items():
            row[j] = v
        out[i] = row
    for k in order:                                          # rows added in tidy form
        if k in known:
            continue
        row = [''] * len(header)
        for v, j in zip(k, key_pos):
            row[j] = legacy_country(v, spec) if spec.get('country') and header[j] == spec['country'][0] else v
        for j, v in values[k].items():
            row[j] = v
        out.append(row)
    return header, [r for r in out if r is not None]


# ------------------------------------------------------------------------------------------------ cells tabs
def _row_key(at, r, b):
    """The row keys of a block row; a row with blank key cells is named by its position (row<r>)."""
    if not b['key_cols']:                            # a one-row block: no row keys
        return []
    key = [at[(r, kc)][2] if (r, kc) in at else '' for kc in b['key_cols']]
    return key if any(key) else [f'row{r}'] + [''] * (len(key) - 1)


def _col_label(at, c, b):
    """The column label of a block column: the block's own name for it (`labels`, where the legacy header is blank or
    repeated), else the label cell; a blank label cell is named by its position (col<c>)."""
    if c in b.get('labels', {}):
        return b['labels'][c]
    return at[(b['label_row'], c)][2] if (b['label_row'], c) in at else f'col{c}'


def cells_to_tidy(spec, head, cells):
    """A legacy tab at its addresses (cells: [row, col, value, ...extra]) -> {tidy name: (header, rows)}, layout."""
    at = {(int(c[0]), int(c[1])): c for c in cells}
    used, tables = set(), {}
    labels = set()                                   # key and label cells of every block stay in the layout
    for b in spec.get('blocks', []):
        labels |= {(r, kc) for r in range(b['rows'][0], b['rows'][1] + 1) for kc in b['key_cols']}
        labels |= {(b['label_row'], c) for c in range(b['cols'][0], b['cols'][1] + 1)}
    for b in spec.get('blocks', []):
        (r0, r1), (c0, c1) = b['rows'], b['cols']
        trows = []
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if c in b['key_cols'] or (r, c) not in at or (r, c) in labels or (r, c) in used:
                    continue
                key = _row_key(at, r, b)
                lab = _col_label(at, c, b)
                trows.append(key + [lab, at[(r, c)][2]])
                used.add((r, c))
        tables[b['name']] = (b['keys'] + [b['label'], 'value'], trows)
    residual = [c for c in cells if (int(c[0]), int(c[1])) not in used]
    extra = {f'{r},{c}': at[(r, c)][3:] for (r, c) in used if len(at[(r, c)]) > 3 and any(at[(r, c)][3:])}
    order = [[int(c[0]), int(c[1])] for c in cells]
    layout = dict(kind='cells', header=head, extra=extra,
                  order=None if order == sorted(order) else order)
    return tables, layout, residual


def cells_from_tidy(spec, tables, layout, residual):
    head = layout['header']
    width = len(head)
    at = {(int(c[0]), int(c[1])): c for c in residual}
    out = list(residual)
    for b in spec.get('blocks', []):
        (r0, r1), (c0, c1) = b['rows'], b['cols']
        row_of = {}
        for r in range(r0, r1 + 1):
            row_of.setdefault(tuple(_row_key(at, r, b)), r)
        col_of = {_col_label(at, c, b): c for c in range(c0, c1 + 1) if c not in b['key_cols']}
        _thead, trows = tables[b['name']]
        nk = len(b['key_cols'])
        for t in trows:
            r, c = row_of.get(tuple(t[:nk])), col_of.get(t[nk])
            assert r and c, f"{b['name']}: no place in the legacy tab for {t[:nk + 1]}"
            extra = layout['extra'].get(f'{r},{c}', [''] * (width - 3))
            out.append([str(r), str(c), t[-1]] + list(extra)[:width - 3])
    if layout.get('order'):
        pos = {(r, c): i for i, (r, c) in enumerate(layout['order'])}
        out.sort(key=lambda x: pos.get((int(x[0]), int(x[1])), len(pos)))
    else:
        out.sort(key=lambda x: (int(x[0]), int(x[1])))
    return head, out


# ------------------------------------------------------------------------------------------------ legacy side I/O
def read_legacy(src):
    """{file name: (header, rows)} for every spec'd file found in a legacy-style workbook, or in a folder of legacy CSV
    files (the build cache of data/scripts/legacy_csv.py, never a stored form)."""
    found = {}
    if os.path.isdir(src):
        for name in list(TABLES) + list(CELLS) + list(LAYOUT_FILES) + list(ASIS):
            p = os.path.join(src, name)
            if os.path.exists(p):
                rows = read_csv(p)
                found[name] = (rows[0], rows[1:])
        return found
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    for name, spec in TABLES.items():
        if spec['sheet'] not in wb.sheetnames:
            continue
        vals = [[cell_text(v) for v in r] for r in wb[spec['sheet']].iter_rows(values_only=True)]
        h = spec['header_row'] - 1
        head = vals[h]
        n = max(j for j, v in enumerate(head) if v != '') + 1 if any(head) else 0
        if spec['sheet'] in LINKS:                   # the linked table: drop FilterISO3, blank the Column<n> names
            n -= head[n - 1] == FILTER_COL
            head = ['' if v == f'Column{j + 1}' else v for j, v in enumerate(head)]
        body = [(r + [''] * n)[:n] for r in vals[h + 1:]]
        while body and not any(body[-1]):
            body.pop()
        found[name] = (head[:n], body)
    for name, spec in ASIS.items():
        if spec['sheet'] not in wb.sheetnames:
            continue
        vals = [_rstrip([cell_text(v) for v in r]) for r in wb[spec['sheet']].iter_rows(values_only=True)]
        found[name] = (vals[0], vals[1:])         # ragged as the CSV may be: trailing blanks dropped
    for name, spec in CELLS.items():
        if spec['sheet'] not in wb.sheetnames:
            continue
        cells = []
        for r, row in enumerate(wb[spec['sheet']].iter_rows(values_only=True), 1):
            for c, v in enumerate(row, 1):
                if v not in (None, '') and r >= spec.get('first_row', 1):
                    cells.append([str(r), str(c), cell_text(v)] + [''] * (len(spec['header']) - 3))
        found[name] = (spec['header'], cells)
    return found


def _rstrip(r):
    while r and r[-1] == '':
        r = r[:-1]
    return r


def link_table(ws, sheet, h, head, rows):
    """The Excel table L_<sheet> the prototype's Power Query reads: the legacy range plus a column FilterISO3."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    k, f = head.index(LINKS[sheet]), len(head) + 1
    for j, v in enumerate(head, 1):
        if v == '':
            ws.cell(h, j, f'Column{j}')               # an Excel table needs a name in every header cell
    ws.cell(h, f, FILTER_COL)
    for i, r in enumerate(rows, h + 1):
        iso = r[k].split('.')[0].upper() if k < len(r) else ''
        if iso:
            ws.cell(i, f, iso)
    t = Table(displayName=f'L_{sheet}', ref=f'A{h}:{get_column_letter(f)}{h + max(len(rows), 1)}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    ws.add_table(t)


def write_legacy_xlsx(path, tables):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (head, rows) in tables.items():
        if name in TABLES:
            spec = TABLES[name]
            ws = wb.create_sheet(spec['sheet'])
            h = spec['header_row']
            for j, v in enumerate(head, 1):
                if v != '':
                    ws.cell(h, j, v)
            for i, r in enumerate(rows, h + 1):
                for j, v in enumerate(r, 1):
                    if v != '':
                        ws.cell(i, j, cell_value(v))
            if spec['sheet'] in LINKS:
                link_table(ws, spec['sheet'], h, head, rows)
        elif name in ASIS:
            ws = wb.create_sheet(ASIS[name]['sheet'])
            for i, r in enumerate([head] + rows, 1):
                for j, v in enumerate(r, 1):
                    if v != '':
                        ws.cell(i, j, v if i == 1 else cell_value(v))
        elif name in CELLS:
            ws = wb.create_sheet(CELLS[name]['sheet'])
            for c in rows:
                ws.cell(int(c[0]), int(c[1]), cell_value(c[2]))
    wb.save(path)


# ------------------------------------------------------------------------------------------------ tidy side I/O
def to_tidy(legacy, out):
    manifest = []
    for name, (head, rows) in legacy.items():
        lay = os.path.join(out, '_layout')
        if name in TABLES:
            spec = TABLES[name]
            thead, trows, layout, lay_rows = table_to_tidy(spec, head, rows)
            layout['source'] = name
            write_csv(os.path.join(out, f"{spec['tidy']}.csv"), thead, trows)
            write_csv(os.path.join(lay, f"{spec['tidy']}.rows.csv"),
                      ['_row'] + [head[j] for j in layout['key_pos'] + layout['attr_pos']], lay_rows)
            with open(os.path.join(lay, f"{spec['tidy']}.json"), 'w', encoding='utf-8') as f:
                json.dump(layout, f, ensure_ascii=False)
            manifest.append([spec['tidy'], name, 'table', spec['sheet'], ' '.join(thead[:len(spec['keys'])]),
                             len(trows), spec.get('proprietary', 'no')])
        elif name in CELLS:
            spec = CELLS[name]
            blocks, layout, residual = cells_to_tidy(spec, head, rows)
            layout['source'] = name
            for bname, (bh, br) in blocks.items():
                write_csv(os.path.join(out, f'{bname}.csv'), bh, br)
                manifest.append([bname, name, 'cells block', spec['sheet'], ' '.join(bh[:-2]), len(br), 'no'])
            write_csv(os.path.join(lay, f"{spec['tidy']}.cells.csv"), head, residual)
            with open(os.path.join(lay, f"{spec['tidy']}.json"), 'w', encoding='utf-8') as f:
                json.dump(layout, f, ensure_ascii=False)
            manifest.append([spec['tidy'] + ' (layout)', name, 'cells layout', spec['sheet'], 'row col', len(residual),
                             'no'])
        elif name in ASIS:
            write_csv(os.path.join(out, f"{ASIS[name]['tidy']}.csv"), head, rows)
            manifest.append([ASIS[name]['tidy'], name, 'as is', ASIS[name]['sheet'], '', len(rows), 'no'])
        elif name in LAYOUT_FILES:
            write_csv(os.path.join(lay, name), head, rows)
    write_csv(os.path.join(out, '_manifest.csv'),
              ['table', 'legacy_source', 'kind', 'legacy_tab', 'keys', 'rows', 'proprietary'], sorted(manifest))
    return manifest


def from_tidy(tdir):
    """{legacy file name: (header, rows)} from a tidy folder."""
    lay = os.path.join(tdir, '_layout')
    out = {}
    for name, spec in TABLES.items():
        p = os.path.join(lay, f"{spec['tidy']}.json")
        if not os.path.exists(p):
            continue
        layout = json.load(open(p, encoding='utf-8'))
        t = read_csv(os.path.join(tdir, f"{spec['tidy']}.csv"))
        lr = read_csv(os.path.join(lay, f"{spec['tidy']}.rows.csv"))[1:]
        out[name] = table_from_tidy(spec, t[0], t[1:], layout, lr)
    for name, spec in CELLS.items():
        p = os.path.join(lay, f"{spec['tidy']}.json")
        if not os.path.exists(p):
            continue
        layout = json.load(open(p, encoding='utf-8'))
        res = read_csv(os.path.join(lay, f"{spec['tidy']}.cells.csv"))[1:]
        tables = {}
        for b in spec.get('blocks', []):
            t = read_csv(os.path.join(tdir, f"{b['name']}.csv"))
            tables[b['name']] = (t[0], t[1:])
        out[name] = cells_from_tidy(spec, tables, layout, res)
    for name, spec in ASIS.items():
        p = os.path.join(tdir, f"{spec['tidy']}.csv")
        if os.path.exists(p):
            rows = read_csv(p)
            out[name] = (rows[0], rows[1:])
    for name in LAYOUT_FILES:
        p = os.path.join(lay, name)
        if os.path.exists(p):
            rows = read_csv(p)
            out[name] = (rows[0], rows[1:])
    return out


# ------------------------------------------------------------------------------------------------ check
def _same(x, y, numeric):
    """Two cell texts are the same; numeric: numbers equal to 1e-15 relative (the workbook form keeps 16
    significant digits, openpyxl), whole numbers equal whatever their form (2027 = 2027.0)."""
    if x == y or not numeric:
        return x == y
    try:
        fx, fy = float(x), float(y)
    except ValueError:
        return False
    return abs(fx - fy) <= 1e-15 * max(abs(fx), abs(fy))


def compare(a, b, numeric=False, width=None):
    """Differences between two {file: (header, rows)}; numeric=True for the workbook form (see _same); width: compare
    only the first `width` columns of cells files (a values-only workbook has no formula flag)."""
    diffs = []
    for name in sorted(set(a) | set(b)):
        if name not in a or name not in b:
            diffs.append(f'{name}: only in one side')
            continue
        ha, ra = a[name]
        hb, rb = b[name]
        w = width if (width and name in CELLS) else None
        ha, hb = (_rstrip(ha), _rstrip(hb)) if numeric else (ha, hb)    # the workbook keeps no trailing blanks
        if not all(_same(x, y, numeric) for x, y in zip(ha[:w], hb[:w])) or len(ha[:w]) != len(hb[:w]):
            diffs.append(f'{name}: header differs')
        if len(ra) != len(rb):
            diffs.append(f'{name}: {len(ra)} rows against {len(rb)}')
            continue
        for i, (x, y) in enumerate(zip(ra, rb)):
            x, y = x + [''] * (len(y) - len(x)), y + [''] * (len(x) - len(y))
            if not all(_same(u, v, numeric) for u, v in zip(x[:w], y[:w])):
                j = next(j for j, (u, v) in enumerate(zip(x, y)) if not _same(u, v, numeric))
                diffs.append(f'{name}: row {i + 1}, column {j + 1}: {x[j]!r} / {y[j]!r}')
                break
    return diffs


def check(tdir):
    """Differences between the tidy folder and its round trip through the workbook (legacy form, values)."""
    back = {k: v for k, v in from_tidy(tdir).items() if k not in LAYOUT_FILES}
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = os.path.join(tmp, 'legacy.xlsx')
        write_legacy_xlsx(xlsx, back)
        again = from_tidy_via(read_legacy(xlsx), tmp)
    return compare(back, again, numeric=True, width=3)


def from_tidy_via(legacy, tmp):
    """legacy (from a workbook) -> tidy -> legacy: the workbook round trip."""
    d = os.path.join(tmp, 'tidy2')
    to_tidy(legacy, d)
    return {k: v for k, v in from_tidy(d).items() if k not in LAYOUT_FILES}


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest='cmd', required=True)
    a = sub.add_parser('to-xlsx')
    a.add_argument('--tidy', required=True, help='tidy folder')
    a.add_argument('--out', required=True, help='legacy-style workbook (.xlsx)')
    b = sub.add_parser('to-tidy')
    b.add_argument('--xlsx', required=True, help='legacy-style workbook (.xlsx)')
    b.add_argument('--out', required=True, help='tidy folder')
    c = sub.add_parser('check')
    c.add_argument('--tidy', required=True)
    args = p.parse_args()
    if args.cmd == 'to-xlsx':
        back = {k: v for k, v in from_tidy(args.tidy).items() if k not in LAYOUT_FILES}
        write_legacy_xlsx(args.out, back)
        print(f'{len(back)} legacy tabs written to {args.out}')
    elif args.cmd == 'to-tidy':
        m = to_tidy(read_legacy(args.xlsx), args.out)
        print(f'{len(m)} tidy tables written to {args.out}')
    else:
        d = check(args.tidy)
        print('Workbook round trip:', 'equal' if not d else '\n  ' + '\n  '.join(d))
        sys.exit(1 if d else 0)


if __name__ == '__main__':
    main()
