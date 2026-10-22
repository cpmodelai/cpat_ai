"""
Convert a CPAT data workbook between the Python format and the CPAT (legacy) format.

    Python format   CPAT_Data_<version>_python.xlsx, written by scripts/build_data_workbook.py: one c_ table per
                    Coded CPAT CSV, as the file is (header row 1), plus README, Index and the four L_ tables.
    CPAT format     CPAT_Data_<version>_cpat.xlsx: one sheet per legacy CPAT data sheet, with the legacy sheet name,
                    header rows and cell positions of `CPAT 1.0pre_456_NoPropData.xlsb`, so the data block pastes
                    onto the legacy sheet at the same address (Paste Special > Values). Plus:
                      README, Index   what each sheet is, its data range, the legacy columns left blank
                      py_<table>      a Python table with no single legacy table yet, carried as it is
                      pyx_<table>     the cells of a converted table that have no place in the legacy layout (e.g.
                                      the WEO Scale column), so that CPAT -> Python gives back every value

Converted (Python CSV -> legacy sheet): energy_consumption -> EnergyCons, prices_dom -> Prices_dom,
distn_hh_survey -> HHSurvey, distn_io_gtap -> IO_GTAP (as export_to_legacy.py), weo -> WEO, wdi -> WDI (second block,
header row 265), un_pop -> UNPop, distn_gdp_ratios -> GDPRatios, ghgs -> GHGs (the inventory: header rows 5-7, data
from row 8),
ch4_inputs -> CH4Inputs, oil_product_shares -> OilProductShares, ef_ghg -> EF_GHG (CO2 rows), cal_val -> CalVal,
air_data -> Air_Data (density and NetCV rows), distn_hh_elast -> HH_Elast, distn_aspire -> ASPIRE,
distn_who_cooking -> WHOCooking. The rest are carried as py_ tables (their legacy tabs hold several tables).

Every Python -> CPAT run converts back in memory and checks that the result, with the pyx_ cells, equals the input
(numbers exactly); it stops on a difference. Header layouts and the code lookups (WEO and WDI series names) come from
cpat_layouts.json, rebuilt from the legacy workbook with --rebuild-layouts (needs pyxlsb).

Run from data/:
    python scripts/convert_data_workbook.py to-cpat       # workbooks/CPAT_Data_dummy_python.xlsx -> ..._cpat.xlsx
    python scripts/convert_data_workbook.py to-python     # workbooks/CPAT_Data_dummy_cpat.xlsx -> ..._python.xlsx
    python scripts/convert_data_workbook.py to-cpat --in dummy --out x_cpat.xlsx   # from a CSV folder
    python scripts/convert_data_workbook.py to-python --in x_cpat.xlsx --out x_python.xlsx --csv-dir x_csv
Real data stay outside the repository (CPAT_Data_*.xlsx is git-ignored except the dummy files).
"""
import argparse
import csv
import datetime
import json
import math
import os
import re
import sys
import tempfile

import openpyxl
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_ROOT = os.path.normpath(os.path.join(HERE, '..'))         # data/
REPO = os.path.dirname(DATA_ROOT)
sys.path.insert(0, HERE)
import export_to_legacy as ex  # noqa: E402
import build_data_workbook as bdw  # noqa: E402

LAYOUTS_FILE = os.path.join(HERE, 'cpat_layouts.json')
WORKBOOKS = os.path.join(DATA_ROOT, 'workbooks')
DUMMY_PYTHON = os.path.join(WORKBOOKS, 'CPAT_Data_dummy_python.xlsx')
FILTER_COL = bdw.FILTER_COL
INT_COLS = {'year', 'Year', 'quant_cons', 'Latest Year'}     # whole numbers read back from Excel stay integers
BOLD = Font(bold=True)

# Python WEO indicator -> legacy WEO series (IMF data API codes in legacy; the same series under its older WEO code)
WEO_CODE_TO_LEGACY = {'BCA_BP6': 'BCA', 'BCA_GDP_BP6': 'BCA_NGDPD', 'ENDA': 'XDC_USD.PA_RT', 'PCPI_PCH': 'PCPIPCH'}
WEO_CODE_FROM_LEGACY = {v: k for k, v in WEO_CODE_TO_LEGACY.items()}
# Python WDI code -> legacy WDI 'code' (second block of the WDI sheet)
WDI_CODE_TO_LEGACY = {'enda': 'XDC_USD.PA_RT.A'}
WDI_CODE_FROM_LEGACY = {v: k for k, v in WDI_CODE_TO_LEGACY.items()}
# Air_Data code prefix and unit by series (legacy codes 'den.die.aus', 'ncv.die.aus')
AIR_SERIES = {'density': ('den', 'ton/m3'), 'NetCV': ('ncv', 'kcal/kg')}


# ---------------------------------------------------------------- layouts
def _trim(r):
    r = list(r)
    while r and r[-1] is None:
        r.pop()
    return r


def _clean(v):
    return int(v) if isinstance(v, float) and v.is_integer() and abs(v) < 1e15 else v


def rebuild_layouts():
    """Read the header rows of the legacy data sheets (rows above the data block) and the WEO/WDI code lookups into
    cpat_layouts.json. EnergyCons, Prices_dom, HHSurvey and IO_GTAP stay in legacy_layouts.json (export_to_legacy.py)."""
    from pyxlsb import open_workbook

    # sheet -> (header rows kept, column-header row, first data row)
    spec = {'WEO': (range(1, 10), 9, 10), 'WDI': (list(range(1, 10)) + [263, 265], 265, 266),
            'UNPop': (range(1, 7), 6, 7), 'GDPRatios': ([1], 1, 2), 'GHGs': (range(1, 8), 7, 8),
            'CH4Inputs': ([1], 1, 2), 'OilProductShares': ([1], 1, 2), 'EF_GHG': (range(1, 4), 3, 4),
            'CalVal': (range(1, 3), 2, 3), 'Air_Data': (range(1, 4), 3, 4), 'HH_Elast': ([1], 1, 2),
            'ASPIRE': ([1], 1, 2), 'WHOCooking': ([1], 1, 2)}
    out = {'_comment': 'Header layouts of legacy CPAT data sheets and the code lookups of WEO and WDI, read from '
                       'excel_ai/legacy/workbooks/CPAT 1.0pre_456_NoPropData.xlsb by '
                       'data/scripts/convert_data_workbook.py --rebuild-layouts. Do not edit by hand.'}
    with open_workbook(ex.LEGACY_XLSB) as wb:
        for sheet, (keep, hrow, first) in spec.items():
            rows = {}
            with wb.get_sheet(sheet) as sh:
                for row in sh.rows(sparse=True):       # rows keyed by the row index of their cells
                    if row:
                        rows[row[0].r + 1] = _trim([_clean(c.v) for c in row])
            out[sheet] = {'header': {str(r): rows.get(r, []) for r in keep if rows.get(r)},
                          'header_row': hrow, 'first_data_row': first}
            if sheet == 'WEO':          # series code -> [DATASET, INDICATOR]
                lk = {}
                for r in range(first, max(rows) + 1):
                    row = rows.get(r, []) + [None] * 6
                    if isinstance(row[3], str) and row[3].endswith('.A'):
                        lk.setdefault(row[3].split('.', 1)[1][:-2], [row[2], row[5]])
                out[sheet]['series'] = lk
            if sheet == 'WDI':          # code -> [Indicator code, Indicator]
                lk = {}
                for r in range(first, max(rows) + 1):
                    row = rows.get(r, []) + [None] * 7
                    if isinstance(row[6], str):
                        lk.setdefault(row[6], [row[2], row[5]])
                out[sheet]['series'] = lk
    with open(LAYOUTS_FILE, 'w', encoding='utf-8') as f:
        json.dump(out, f, indent=1, ensure_ascii=False)
    print(f'wrote {LAYOUTS_FILE}')


def load_layouts():
    """Every converted sheet as {'header': {row: values}, 'header_row', 'first_data_row', ...}."""
    with open(LAYOUTS_FILE, encoding='utf-8') as f:
        lay = {k: v for k, v in json.load(f).items() if not k.startswith('_')}
    for name, l in ex.load_layouts().items():
        if name.startswith('_'):
            continue
        n = len(l['header_rows'])
        lay[name] = dict(l, header={str(i): r for i, r in enumerate(l['header_rows'], 1) if r},
                         header_row=n, first_data_row=l['first_data_row'])
    return lay


def _hdr(layout):
    """The column-header row of the table."""
    return list(layout['header'][str(layout['header_row'])])


def _padded(rows, layout):
    """Data rows as lists at least as wide as the column-header row (the converters index them by column)."""
    w = len(_hdr(layout))
    return [list(r) + [None] * (w - len(r)) for r in rows]


# ---------------------------------------------------------------- values
def _null(v):
    return v is None or (isinstance(v, float) and math.isnan(v)) or (isinstance(v, str) and v == '')


def _same(a, b):
    if _null(a) or _null(b):
        return _null(a) and _null(b)
    num = (int, float)
    if isinstance(a, num) and isinstance(b, num) and not isinstance(a, bool) and not isinstance(b, bool):
        return float(a) == float(b)
    return str(a) == str(b)


def _py(v):
    """Plain Python value for openpyxl and the comparisons (NaN -> None, numpy -> Python)."""
    if v is None:
        return None
    if hasattr(v, 'item') and not isinstance(v, (str, bytes)):
        v = v.item()
    if isinstance(v, float) and math.isnan(v):
        return None
    return v


def _frame(header, rows):
    return pd.DataFrame([list(r) for r in rows], columns=list(header), dtype=object)


def _records(df):
    return [[_py(v) for v in r] for r in df.itertuples(index=False, name=None)]


# ---------------------------------------------------------------- Python tables: CSV folder or Python workbook
def read_python_tables(src):
    """{table: (header, rows)} from a folder of Coded CPAT CSVs or a Python-format workbook (its c_ tables)."""
    if os.path.isdir(src):
        return {os.path.splitext(f)[0]: bdw.read_csv(os.path.join(src, f))
                for f in sorted(os.listdir(src)) if f.lower().endswith('.csv')}
    wb = openpyxl.load_workbook(src, read_only=True)
    idx = list(wb['Index'].iter_rows(values_only=True))
    cols = {h: i for i, h in enumerate(idx[0])}
    out = {}
    for e in idx[1:]:
        table, sheet, source = e[cols['Table']], e[cols['Sheet']], e[cols['Source']]
        if not str(table).startswith('c_'):
            continue                                        # L_ tables are rebuilt from the c_ tables
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        header, data = [str(h) for h in rows[0]], rows[1:]
        if e[cols['Country filter column']] == FILTER_COL and header[-1] == FILTER_COL:
            header, data = header[:-1], [r[:-1] for r in data]
        for change in str(e[cols['Header changes and notes']] or '').split('; '):     # undo 'old -> new'
            if ' -> ' in change:
                old, new = change.split(' -> ', 1)
                if new in header:
                    header[header.index(new)] = '' if old == '(blank)' else old
        out[os.path.splitext(source)[0]] = (header, [_trim_to(r, len(header)) for r in data])
    wb.close()
    return out


def _trim_to(r, n):
    r = list(r)[:n]
    return r + [None] * (n - len(r))


def _csv_text(v):
    if _null(v):
        return ''
    if isinstance(v, float):
        return repr(v)
    return str(v)


def write_csv_dir(tables, out_dir):
    """One CSV per table, as the Coded CPAT files are (ghgs.csv with its two header rows, sector and gas)."""
    os.makedirs(out_dir, exist_ok=True)
    for name, (header, rows) in tables.items():
        with open(os.path.join(out_dir, f'{name}.csv'), 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f)
            if any(' | ' in h for h in header):
                w.writerow([h.split(' | ', 1)[0] if ' | ' in h else '' for h in header])
                w.writerow([h.split(' | ', 1)[1] if ' | ' in h else h for h in header])
            else:
                w.writerow(header)
            for r in rows:
                w.writerow([_csv_text(v) for v in r])


# ---------------------------------------------------------------- converters
# Each converter: to_cpat(tables, data_dir, layout) -> (data rows from column A, notes);
#                 to_python(data rows, layout) -> DataFrame in the Python layout.
def _legacy4(name):
    """EnergyCons, Prices_dom, HHSurvey, IO_GTAP: forward as export_to_legacy.py (which reads the CSVs)."""
    fn = {'EnergyCons': lambda d, l: ex.energy_cons_rows(d, l, None, 2022),
          'Prices_dom': lambda d, l: ex.prices_dom_rows(d, l, None),
          'HHSurvey': lambda d, l: ex.hh_survey_rows(d, l, None),
          'IO_GTAP': lambda d, l: ex.io_gtap_rows(d, l, None)}[name]
    return lambda tables, data_dir, layout: fn(data_dir, layout)


def energy_to_python(rows, layout):
    header = _hdr(layout)
    sectors = header[4:]
    legacy_to_coded = {v: k for k, v in ex.ENERGY_FUEL_TO_LEGACY.items()}
    fuels = [legacy_to_coded.get(f, f) for f, _ in layout['fuels']]
    recs = {}
    for r in rows:
        if not isinstance(r[0], str) or '.mit.enc.' not in r[0]:
            continue
        country, fuel = r[0].split('.mit.enc.')
        fuel = legacy_to_coded.get(fuel, fuel)
        for s, v in zip(sectors, r[4:]):
            if not _null(v):
                recs.setdefault((country.upper(), s), {})[fuel] = v
    return pd.DataFrame([[c, s] + [d.get(f) for f in fuels] for (c, s), d in recs.items()],
                        columns=['CountryCode', 'SectorCode'] + fuels, dtype=object)


def prices_to_python(rows, layout):
    header = _hdr(layout)
    drop = {'country_year', 'countrycode_weo', 'countrycode', ex.PRICES_END_MARKER}
    cols = [h for h in header if h not in drop]
    i_code = header.index('countrycode')
    out = [[r[i_code]] + [r[header.index(h)] for h in cols] for r in rows if not _null(r[i_code])]
    return pd.DataFrame(out, columns=['CountryCode'] + cols, dtype=object)


def _long_to_python(nkeys, var_col, val_col, all_values=False):
    """Wide legacy table (keys, one column per variable) -> long Python table. Blank cells are left out unless
    all_values (GDPRatios: the Python table has a row for every year, blank or not)."""
    def conv(rows, layout):
        header = _hdr(layout)
        keys, var = header[:nkeys], header[nkeys:]
        out = []
        for r in rows:
            if all(_null(v) for v in r[:nkeys]):
                continue
            for name, v in zip(var, r[nkeys:]):
                if all_values or not _null(v):
                    out.append(list(r[:nkeys]) + [name, v])
        return pd.DataFrame(out, columns=keys + [var_col, val_col], dtype=object)
    return conv


def _wide_to_cpat(table, nkeys, var_col, val_col, sheet):
    def conv(tables, data_dir, layout):
        header = _hdr(layout)
        df = _frame(*tables[table])
        keys = header[:nkeys]
        w, cols, missing = ex._wide(df, keys, var_col, val_col, header[nkeys:])
        return ex._rows_from_wide(w, keys, header[nkeys:], cols, missing, sheet, f'{table}.csv')
    return conv


def _rename_to_cpat(table, names=None, fixed=None):
    """Row-for-row tables: legacy column <- Python column of the same name (or names[legacy]), or fixed[legacy](rec)."""
    names, fixed = names or {}, fixed or {}

    def conv(tables, data_dir, layout):
        header, rows = tables[table]
        legacy = _hdr(layout)
        pos = {h: i for i, h in enumerate(header)}
        src = [None if h is None or h in fixed else pos.get(names.get(h, h)) for h in legacy]
        out = []
        for r in rows:
            rec = dict(zip(header, r))
            out.append([fixed[h](rec) if h in fixed else None if i is None else r[i] for h, i in zip(legacy, src)])
        blank = [h for h, i in zip(legacy, src) if h is not None and h not in fixed and i is None]
        notes = [f'{layout["sheet"]}: legacy columns {", ".join(blank)} are not in {table}.csv; left blank.']
        return out, notes if blank else []
    return conv


def _rename_to_python(names=None, skip=(), keep=None):
    """Legacy columns -> Python columns (names: legacy -> Python), leaving out the codes legacy builds (skip).
    keep(row) picks the rows that belong to the Python table."""
    names = names or {}

    def conv(rows, layout):
        legacy = _hdr(layout)
        idx = [i for i, h in enumerate(legacy) if h is not None and h not in skip]
        out = [[r[i] if i < len(r) else None for i in idx] for r in rows if keep is None or keep(r)]
        return pd.DataFrame(out, columns=[names.get(legacy[i], legacy[i]) for i in idx], dtype=object)
    return conv


def _years(header, first_col):
    return {int(h): i for i, h in enumerate(header) if i >= first_col and isinstance(h, (int, float))}


def _last_year(rec, years):
    have = [y for y in years if not _null(rec.get(str(y)))]
    return max(have) if have else None


def weo_to_cpat(tables, data_dir, layout):
    header, rows = tables['weo']
    legacy = _hdr(layout)
    years = _years(legacy, 6)
    series = layout['series']
    out = []
    for r in rows:
        rec = dict(zip(header, r))
        code = WEO_CODE_TO_LEGACY.get(rec['Indicator'], rec['Indicator'])
        dataset, label = series.get(code, ['IMF.RES:WEO', rec['Indicator']])
        row = [_last_year(rec, years), rec['CountryCode'], dataset, f'{rec["CountryCode"]}.{code}.A', rec['Country'],
               label] + [None] * (len(legacy) - 6)
        for y, i in years.items():
            row[i] = rec.get(str(y))
        out.append(row)
    renamed = sorted(set(WEO_CODE_TO_LEGACY) & {r[header.index('Indicator')] for r in rows})
    notes = [f'WEO: indicators {", ".join(renamed)} '
             'take their legacy series code; Scale has no legacy column (pyx_weo).']
    return out, notes


def weo_to_python(rows, layout):
    legacy = _hdr(layout)
    years = _years(legacy, 6)
    out = []
    for r in rows:
        if not isinstance(r[3], str) or not r[3].endswith('.A'):
            continue
        code = r[3].split('.', 1)[1][:-2]
        out.append([r[1], r[4], WEO_CODE_FROM_LEGACY.get(code, code)] + [r[i] for i in years.values()])
    return pd.DataFrame(out, columns=['CountryCode', 'Country', 'Indicator'] + [str(y) for y in years], dtype=object)


def wdi_to_cpat(tables, data_dir, layout):
    header, rows = tables['wdi']
    legacy = _hdr(layout)
    years = _years(legacy, 7)
    out = []
    for r in rows:
        rec = dict(zip(header, r))
        code = WDI_CODE_TO_LEGACY.get(rec['code'], f'{rec["code"]}.a')
        wb_code, label = layout['series'].get(code, [None, None])
        row = [_last_year(rec, years), rec['CountryCode'], wb_code, f'{rec["CountryCode"]}.{code}', rec['Country'],
               label, code] + [None] * (len(legacy) - 7)
        for y, i in years.items():
            row[i] = rec.get(str(y))
        out.append(row)
    return out, ['WDI: the Python series fill the second legacy block (header row 265); the first block (urban '
                 'population growth, rows 10-260) is not in wdi.csv and is left blank.']


def wdi_to_python(rows, layout):
    legacy = _hdr(layout)
    years = _years(legacy, 7)
    out = []
    for r in rows:
        if not isinstance(r[6], str):
            continue
        code = WDI_CODE_FROM_LEGACY.get(r[6], r[6][:-2] if r[6].endswith('.a') else r[6])
        out.append([r[1], r[4], code] + [r[i] for i in years.values()])
    return pd.DataFrame(out, columns=['CountryCode', 'Country', 'code'] + [str(y) for y in years], dtype=object)


def unpop_to_cpat(tables, data_dir, layout):
    header, rows = tables['un_pop']
    legacy = _hdr(layout)
    years = _years(legacy, 4)
    out = []
    for r in rows:
        rec = dict(zip(header, r))
        row = [None, f'{str(rec["CountryCode"]).lower()}.lp', None, rec['CountryCode']] + [None] * (len(legacy) - 4)
        for y, i in years.items():
            row[i] = rec.get(str(y))
        out.append(row)
    return out, ['UNPop: column C (Population) is not in un_pop.csv; left blank.']


def unpop_to_python(rows, layout):
    years = _years(_hdr(layout), 4)
    out = [[r[3]] + [r[i] for i in years.values()] for r in rows if not _null(r[3])]
    return pd.DataFrame(out, columns=['CountryCode'] + [str(y) for y in years], dtype=object)


GDP_KEYS = ['Country Name', 'Country Code', 'Indicator Name', 'Indicator Code', 'Latest Year']


def gdp_to_cpat(tables, data_dir, layout):
    df = _frame(*tables['distn_gdp_ratios'])
    legacy = _hdr(layout)
    years = _years(legacy, 5)
    out = []
    for keys, g in df.groupby(GDP_KEYS, sort=False, dropna=False):
        row = list(keys) + [None] * (len(legacy) - 5)
        for y, v in zip(g['year'], g['value']):
            if int(y) in years:
                row[years[int(y)]] = v
        out.append(row)
    return out, []


def ghgs_cols(layout):
    """legacy column index -> Python column ('1.  Energy | CO2'), groups carried right from header row 6."""
    groups, subs = layout['header']['6'], layout['header']['7']
    out, group = {2: 'Country', 3: 'CountryCode', 4: 'Year'}, None
    for i in range(7, len(subs)):
        group = groups[i] if i < len(groups) and groups[i] is not None else group
        if subs[i] is not None:
            out[i] = f'{group} | {subs[i]}'
    return out


def ghgs_to_cpat(tables, data_dir, layout):
    header, rows = tables['ghgs']
    pos = {h: i for i, h in enumerate(header)}
    cols = ghgs_cols(layout)
    width = max(cols) + 1
    out = [[r[pos[cols[i]]] if i in cols and cols[i] in pos else None for i in range(width)] for r in rows]
    blank = [cols[i] for i in sorted(cols) if cols[i] not in pos]
    notes = [f'GHGs: UNFCCC category (column G) and {", ".join(blank)} are not in ghgs.csv; left blank.']
    return out, notes if blank else []


def ghgs_to_python(rows, layout):
    cols = ghgs_cols(layout)
    out = [[r[i] for i in cols] for r in rows if isinstance(r[4], (int, float))]     # inventory rows have a year
    return pd.DataFrame(out, columns=list(cols.values()), dtype=object)


def ef_to_cpat(tables, data_dir, layout):
    header, rows = tables['ef_ghg']
    out = []
    for r in rows:
        rec = dict(zip(header, r))
        cc = 'wld' if rec['CountryCode'] == 'world' else rec['CountryCode']
        out.append([f'{cc}.mit.ef.{rec["SectorCode"]}.{rec["FuelCode"]}.co2', rec['CountryCode'], 'ef CO2',
                    rec['SectorCode'], rec['FuelCode'], rec['Value']])
    return out, ['EF_GHG: only the CO2 factors (series "ef CO2") are in ef_ghg.csv; the GWP/GTP rows of legacy '
                 'EF_GHG are left out.']


def calval_to_cpat(tables, data_dir, layout):
    header, rows = tables['cal_val']
    out = []
    for r in rows:
        rec = dict(zip(header, r))
        fuel = 'oil' if rec['FuelCode'] == 'oop' else rec['FuelCode']      # legacy codes say oil, column E oop
        out.append([f'{rec["CountryCode"]}.mit.ncv.{fuel}.{rec["SectorCode"]}', rec['CountryCode'],
                    rec['SectorCode'], 'ncv', rec['FuelCode'], 'kcal/kg', rec['Value']])
    return out, []


def air_to_cpat(tables, data_dir, layout):
    header, rows = tables['air_data']
    out = []
    for r in rows:
        rec = dict(zip(header, r))
        prefix, unit = AIR_SERIES.get(rec['Series Name'], (rec['Series Name'], None))
        cc = str(rec['CountryCode'])
        out.append([f'{prefix}.{rec["FuelCode"]}.{cc.lower()}', cc.upper(), rec['Series Name'], unit,
                    rec['FuelCode'], rec['Value']])
    return out, ['Air_Data: only the density and NetCV rows are in air_data.csv; the other legacy series (work '
                 'absence, death probabilities, labour share) are left out.']


def air_to_python(rows, layout):
    out = [[str(r[1]).lower(), r[2], r[4], r[5]] for r in rows if len(r) > 5 and not _null(r[4])]
    return pd.DataFrame(out, columns=['CountryCode', 'Series Name', 'FuelCode', 'Value'], dtype=object)


# table -> (legacy sheet, Python key columns, to_cpat, to_python)
CONVERTERS = {
    'energy_consumption': ('EnergyCons', ['CountryCode', 'SectorCode'], _legacy4('EnergyCons'),
                           energy_to_python),
    'prices_dom': ('Prices_dom', ['CountryCode', 'year'], _legacy4('Prices_dom'), prices_to_python),
    'distn_hh_survey': ('HHSurvey', ['code', 'variable'], _legacy4('HHSurvey'),
                        _long_to_python(7, 'variable', 'value')),
    'distn_io_gtap': ('IO_GTAP', ['dis', 'gtap_sector'], _legacy4('IO_GTAP'),
                      _long_to_python(4, 'gtap_sector', 'value')),
    'weo': ('WEO', ['CountryCode', 'Indicator'], weo_to_cpat, weo_to_python),
    'wdi': ('WDI', ['CountryCode', 'code'], wdi_to_cpat, wdi_to_python),
    'un_pop': ('UNPop', ['CountryCode'], unpop_to_cpat, unpop_to_python),
    'distn_gdp_ratios': ('GDPRatios', ['Country Code', 'Indicator Code', 'year'], gdp_to_cpat,
                         _long_to_python(5, 'year', 'value', all_values=True)),
    'ghgs': ('GHGs', ['CountryCode', 'Year'], ghgs_to_cpat, ghgs_to_python),
    'ch4_inputs': ('CH4Inputs', ['CountryCode', 'tier3'], _rename_to_cpat('ch4_inputs', {'cpat_ccode3': 'CountryCode'}),
                   _rename_to_python({'cpat_ccode3': 'CountryCode'})),
    'oil_product_shares': ('OilProductShares', ['CountryCode'],
                           _rename_to_cpat('oil_product_shares', {'countrycode': 'CountryCode'}),
                           _rename_to_python({'countrycode': 'CountryCode'})),
    'ef_ghg': ('EF_GHG', ['CountryCode', 'SectorCode', 'FuelCode'], ef_to_cpat,
               _rename_to_python({'Region/country': 'CountryCode', 'Atribute 1': 'SectorCode',
                                  'Atribute 2': 'FuelCode', 'FACTOR': 'Value'}, skip={'Code', 'Series name'},
                                 keep=lambda r: len(r) > 2 and r[2] == 'ef CO2')),
    'cal_val': ('CalVal', ['CountryCode', 'SectorCode', 'FuelCode'], calval_to_cpat,
                _rename_to_python({'Country code': 'CountryCode', 'Subsector': 'SectorCode', 'Fuel': 'FuelCode'},
                                  skip={'Code', 'Series', 'Unit'}, keep=lambda r: len(r) > 3 and r[3] == 'ncv')),
    'air_data': ('Air_Data', ['CountryCode', 'Series Name', 'FuelCode'], air_to_cpat, air_to_python),
    'distn_hh_elast': ('HH_Elast', ['code', 'variable'],
                       _wide_to_cpat('distn_hh_elast', 7, 'variable', 'value', 'HH_Elast'),
                       _long_to_python(7, 'variable', 'value')),
    'distn_aspire': ('ASPIRE', ['Series_Code'], _rename_to_cpat('distn_aspire'),
                     _rename_to_python()),
    'distn_who_cooking': ('WHOCooking', ['code', 'variable'],
                          _wide_to_cpat('distn_who_cooking', 4, 'variable', 'value', 'WHOCooking'),
                          _long_to_python(4, 'variable', 'value')),
}


# ---------------------------------------------------------------- round trip: the cells legacy has no place for
def _keyed(df, keys):
    df = df.copy()
    df['_occ'] = df.groupby(keys, sort=False, dropna=False).cumcount()
    return df.set_index(keys + ['_occ'])


def _int_cols(df):
    for c in df.columns:
        if c in INT_COLS:
            df[c] = [int(v) if isinstance(v, float) and v.is_integer() else v for v in df[c]]
    return df


def python_only_cells(orig, back, keys):
    """The cells of orig that the CPAT layout cannot hold (blank or missing in back), as a Python-layout frame
    (keys + those columns; other cells blank); raises if back holds a value that differs from orig."""
    o, b = _keyed(orig, keys), _keyed(_int_cols(back), keys)
    if not b.index.is_unique:
        raise ValueError(f'keys {keys} do not identify the rows')
    extra = pd.DataFrame(index=o.index, columns=o.columns, dtype=object)
    common = o.index.intersection(b.index, sort=False)
    problems = []
    for c in o.columns:
        ov = o[c]
        bv = b[c].reindex(o.index) if c in b.columns else pd.Series([None] * len(o), index=o.index, dtype=object)
        for ix, x, y in zip(o.index, ov, bv):
            if _null(x):
                continue
            if _null(y):
                extra.at[ix, c] = x
            elif not _same(x, y):
                problems.append(f'{ix} {c}: {x!r} -> {y!r}')
    spurious = [ix for ix in b.index.difference(common, sort=False) if any(not _null(v) for v in b.loc[ix])]
    if problems or spurious:
        raise ValueError(f'round trip differs: {problems[:5]} spurious rows {spurious[:5]}')
    extra = extra.dropna(how='all', axis=0).dropna(how='all', axis=1)
    return extra.reset_index().drop(columns='_occ') if len(extra) else None


def apply_python_only_cells(back, extra, keys):
    if extra is None or not len(extra):
        return back
    b, e = _keyed(_int_cols(back), keys), _keyed(_int_cols(extra), keys)
    for c in e.columns:
        if c not in b.columns:
            b[c] = None
    new = e.index.difference(b.index, sort=False)
    if len(new):
        b = pd.concat([b, e.loc[new].reindex(columns=b.columns)])
    for c in e.columns:
        for ix, v in zip(e.index, e[c]):
            if not _null(v) and ix not in new:
                b.at[ix, c] = v
    if b.index.isin(e.index).all():                     # the cells cover every row: their order is the Python one
        b = b.loc[e.index]
    return b.reset_index().drop(columns='_occ')


def _final_columns(df, keys, columns):
    """Python column order: as recorded in the Index, else the converter's (columns blank throughout dropped)."""
    if columns:
        return df.reindex(columns=columns)
    keep = [c for c in df.columns if c in keys or any(not _null(v) for v in df[c])]
    return df[keep]


# ---------------------------------------------------------------- CPAT workbook
def _write_cells(ws, header, first_row, rows):
    for r, vals in header.items():
        for c, v in enumerate(vals, start=1):
            if v is not None:
                ws.cell(int(r), c, v)
    for r, vals in enumerate(rows, start=first_row):
        for c, v in enumerate(vals, start=1):
            v = _py(v)
            if v is not None:
                ws.cell(r, c, v)


def _write_table(ws, header, rows):
    ws.append(list(header))
    for c in ws[1]:
        c.font = BOLD
    for r in rows:
        ws.append([_py(v) for v in r])
    ws.freeze_panes = 'A2'


def to_cpat(src, out_path, version='dummy'):
    """Python tables (CSV folder or Python workbook) -> CPAT-format workbook. Returns the Index rows."""
    tables = read_python_tables(src)
    layouts = load_layouts()
    with tempfile.TemporaryDirectory() as tmp:
        data_dir = src if os.path.isdir(src) else tmp
        if not os.path.isdir(src):
            write_csv_dir(tables, tmp)                      # export_to_legacy reads the CSVs
        wb = openpyxl.Workbook()
        readme = wb.active
        readme.title = 'README'
        index = wb.create_sheet('Index')
        used = {'readme', 'index'}
        entries, sheets_py, sheets_x = [], [], []
        for table, (sheet, keys, fwd, back) in CONVERTERS.items():
            if table not in tables:
                continue
            layout = dict(layouts[sheet], sheet=sheet)
            rows, notes = fwd(tables, data_dir, layout)
            rows = [_trim([_py(v) for v in r]) for r in rows]
            orig = _frame(*tables[table])
            extra = python_only_cells(orig, back(_padded(rows, layout), layout), keys)
            ws = wb.create_sheet(sheet)
            used.add(sheet.lower())
            _write_cells(ws, layout['header'], layout['first_data_row'], rows)
            ws.freeze_panes = ws.cell(layout['first_data_row'], 1)
            width = max([len(r) for r in rows] + [len(_hdr(layout))])
            r0 = layout['first_data_row']
            rng = f'A{r0}:{get_column_letter(width)}{r0 + len(rows) - 1}' if rows else '(no rows)'
            xsheet = ''
            if extra is not None:
                xsheet = bdw._sheet_name(f'pyx_{table}', used)
                sheets_x.append((xsheet, extra))
                notes.append(f'Python-only cells ({", ".join(c for c in extra.columns if c not in keys)}) '
                             f'in {xsheet}, {len(extra)} rows.')
            entries.append([sheet, f'{table}.csv', 'legacy sheet, legacy cell addresses', rng, len(rows), xsheet,
                            ' '.join(notes)[:2000], json.dumps(list(tables[table][0]), ensure_ascii=False)])
        for table, (header, rows) in tables.items():
            if table in CONVERTERS:
                continue
            sheet = bdw._sheet_name(f'py_{table}', used)
            sheets_py.append((sheet, header, rows))
            entries.append([sheet, f'{table}.csv', 'Python layout (no single legacy table)',
                            f'A2:{get_column_letter(len(header))}{len(rows) + 1}', len(rows), '', '', ''])
        for sheet, header, rows in sheets_py:
            _write_table(wb.create_sheet(sheet), header, rows)
        for sheet, extra in sheets_x:
            _write_table(wb.create_sheet(sheet), list(extra.columns), _records(extra))
        hdr = ['Sheet', 'Python table', 'Layout', 'Data range', 'Rows', 'Python-only cells', 'Notes',
               'Python columns']
        _write_table(index, hdr, entries)
        for col, w in zip('ABCDEFGH', (34, 38, 36, 14, 8, 26, 100, 40)):
            index.column_dimensions[col].width = w
        _readme(readme, version, src)
        wb.save(out_path)
    sys.path.insert(0, os.path.join(REPO, 'excel_ai', 'shared'))
    import excel_lint
    problems = excel_lint.lint(out_path)
    if problems:
        raise SystemExit(f'{out_path}: Excel would repair it:\n  ' + '\n  '.join(problems))
    return entries


def _readme(ws, version, src):
    lines = [
        ('CPAT data workbook, CPAT format', f'version: {version}'),
        ('Built', f'{datetime.date.today().isoformat()} by data/scripts/convert_data_workbook.py to-cpat'),
        ('Source', os.path.basename(os.path.normpath(src))),
        ('What it is', 'The tables of one data version in the layout of the legacy CPAT workbook '
                       '(CPAT 1.0pre_456_NoPropData.xlsb). The Python format of the same data is '
                       f'CPAT_Data_{version}_python.xlsx; convert_data_workbook.py converts either way.'),
        ('Dummy data', 'Fake, non-proprietary values in the real formats. Never use them for analysis.'
                       if version == 'dummy' else ''),
        ('Legacy sheets', 'Same sheet name, header rows and cell positions as legacy. Copy the data range given in '
                          'the Index and Paste Special > Values onto the legacy sheet at the same top-left cell; clear '
                          'the old data rows first if the new block is shorter.'),
        ('py_ sheets', 'Python tables whose legacy tab holds several tables (ECP, Power, Prices_int, Elasticities, '
                       'Mapping, the GHG shocks) or that legacy does not have; carried in the Python layout.'),
        ('pyx_ sheets', 'Cells of a converted table that the legacy layout has no place for (Python layout, key '
                        'columns plus those cells). Not for legacy; to-python puts them back.'),
        ('Codes', 'WEO: the Python indicators BCA_BP6, BCA_GDP_BP6, ENDA and PCPI_PCH take the legacy series codes '
                  'BCA, BCA_NGDPD, XDC_USD.PA_RT and PCPIPCH; names and datasets come from the legacy WEO sheet. '
                  'WDI: the Python codes fill the second block (header row 265); enda is XDC_USD.PA_RT.A. '
                  'EnergyCons: heat is "heat" in legacy, "hea" in Python. CalVal codes write oop as "oil", as legacy.'),
        ('Derived columns', 'WEO and WDI "last year" is the last year with data; EnergyCons Year is 2022 and its '
                            'country names come from prices_dom.csv; legacy keys (UNPop "abw.lp", EF_GHG '
                            '"afg.mit.ef.ind.coa.co2", ...) are built as legacy builds them.'),
        ('NoPropData file', 'Some lookups into the stripped sheets were shrunk to row 1 there (e.g. Mitigation AA583); '
                            'widen them before pasting into that file. The full workbook has them.'),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ws.cell(r, 1, a).font = BOLD
        ws.cell(r, 2, b)
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 140


def read_cpat_tables(src):
    """CPAT-format workbook -> {table: (header, rows)} in the Python layout."""
    layouts = load_layouts()
    wb = openpyxl.load_workbook(src, read_only=True)
    sheet_of = {v[0]: k for k, v in CONVERTERS.items()}
    entries = {}
    if 'Index' in wb.sheetnames:
        idx = list(wb['Index'].iter_rows(values_only=True))
        cols = {h: i for i, h in enumerate(idx[0])}
        entries = {e[cols['Sheet']]: {h: e[i] for h, i in cols.items()} for e in idx[1:] if e[0]}
    out = {}
    for sheet in wb.sheetnames:
        e = entries.get(sheet, {})
        if sheet in sheet_of:
            table = sheet_of[sheet]
            _, keys, _, back = CONVERTERS[table]
            layout = dict(layouts[sheet], sheet=sheet)
            rows = [_trim(r) for r in wb[sheet].iter_rows(min_row=layout['first_data_row'], values_only=True)]
            df = _int_cols(back(_padded([r for r in rows if r], layout), layout))
            xs = e.get('Python-only cells')
            if xs:
                xr = [list(r) for r in wb[xs].iter_rows(values_only=True)]
                df = apply_python_only_cells(df, _frame(xr[0], xr[1:]), keys)
            cols = json.loads(e['Python columns']) if e.get('Python columns') else None
            df = _final_columns(df, keys, cols)
            out[table] = (list(df.columns), _records(df))
        elif sheet.startswith('py_'):
            table = os.path.splitext(e['Python table'])[0] if e.get('Python table') else sheet[3:]
            rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
            header = ['' if h is None else str(h) for h in rows[0]]
            out[table] = (header, [_trim_to(r, len(header)) for r in rows[1:]])
    wb.close()
    return out


def to_python(src, out_path, version='dummy', csv_dir=None):
    """CPAT-format workbook -> Python-format workbook (via a folder of CSVs, then build_data_workbook.py)."""
    tables = read_cpat_tables(src)
    with tempfile.TemporaryDirectory() as tmp:
        d = csv_dir or tmp
        write_csv_dir(tables, d)
        return bdw.build(d, version, out_path, source=f'{os.path.basename(src)} (convert_data_workbook.py to-python)')


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('direction', nargs='?', choices=['to-cpat', 'to-python'])
    ap.add_argument('--in', dest='src', help='input: Python workbook or CSV folder (to-cpat); CPAT workbook '
                                             '(to-python). Default: the dummy workbook in data/workbooks/')
    ap.add_argument('--out', help='output workbook (default: the input name with _python / _cpat swapped)')
    ap.add_argument('--version', default='dummy', help='version name for the README (default: %(default)s)')
    ap.add_argument('--csv-dir', help='to-python: also keep the Python tables as CSVs in this folder')
    ap.add_argument('--rebuild-layouts', action='store_true',
                    help='re-read cpat_layouts.json from the legacy workbook (needs pyxlsb) and exit')
    a = ap.parse_args(argv)
    if a.rebuild_layouts:
        rebuild_layouts()
        return
    if a.direction is None:
        ap.error('give to-cpat or to-python')
    if a.direction == 'to-cpat':
        src = a.src or DUMMY_PYTHON
        out = a.out or (re.sub(r'_python(\.xlsx)$', r'_cpat\1', src) if src.endswith('_python.xlsx')
                        else os.path.join(WORKBOOKS, f'CPAT_Data_{a.version}_cpat.xlsx'))
        entries = to_cpat(src, out, a.version)
        print(f'{out}: {len(entries)} sheets ({sum(e[2].startswith("legacy") for e in entries)} legacy), '
              f'{sum(e[4] for e in entries)} data rows')
    else:
        src = a.src or DUMMY_PYTHON.replace('_python.xlsx', '_cpat.xlsx')
        out = a.out or re.sub(r'_cpat(\.xlsx)$', r'_python\1', src)
        if out == src:
            ap.error('give --out')
        entries = to_python(src, out, a.version, a.csv_dir)
        print(f'{out}: {len(entries)} tables, {sum(e[4] for e in entries)} rows')


if __name__ == '__main__':
    main()
