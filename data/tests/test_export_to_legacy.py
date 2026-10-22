"""
Tests for data/scripts/export_to_legacy.py: the exported sheets have the legacy
header layout, hold every value of the source CSVs at the right cell, and (for
the distribution tables, which are real Egypt data in both places) reproduce
the legacy workbook's HHSurvey and IO_GTAP cell for cell.
"""
import os
import sys

import numpy as np
import openpyxl
import pandas as pd
import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))   # data/
sys.path.insert(0, os.path.join(ROOT, 'scripts'))

# pylint: disable=wrong-import-position
import export_to_legacy as ex

DATA = os.path.join(ROOT, 'dummy')
# openpyxl writes floats with 16 significant digits (Excel keeps 15), so values can differ by 1 ulp
REL, ABS = 1e-15, 1e-300
# the distribution CSVs come from data_standardized, which rounds to 16 decimal places
LEGACY_ABS = 1e-15


@pytest.fixture(scope='module')
def exported(tmp_path_factory):
    out = tmp_path_factory.mktemp('legacy') / 'out.xlsx'
    ex.export(DATA, str(out))
    wb = openpyxl.load_workbook(out, read_only=True)
    sheets = {n: [list(r) for r in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames}
    wb.close()
    return sheets


def _trim(r):
    r = list(r)
    while r and r[-1] is None:
        r.pop()
    return r


def test_headers_match_legacy_layout(exported):
    layouts = ex.load_layouts()
    for name in ('EnergyCons', 'Prices_dom', 'HHSurvey', 'IO_GTAP'):
        hdr = layouts[name]['header_rows']
        assert [_trim(r) for r in exported[name][:len(hdr)]] == hdr, name


def test_energy_cons_matches_csv(exported):
    ec = pd.read_csv(os.path.join(DATA, 'energy_consumption.csv'), float_precision='round_trip').set_index(['CountryCode', 'SectorCode'])
    rows = exported['EnergyCons']
    header = rows[0]
    assert len(rows) - 1 == ec.index.get_level_values(0).nunique() * 19
    checked = 0
    for r in rows[1:]:
        country, fuel = r[0].split('.mit.enc.')
        assert r[0] == r[0].lower()
        coded_fuel = 'hea' if fuel == 'heat' else fuel
        for j, sector in enumerate(header[4:], start=4):
            if sector in ('wav', 'neu'):
                assert r[j] is None
                continue
            assert r[j] == pytest.approx(ec.at[(country.upper(), sector), coded_fuel], rel=REL, abs=ABS)
            checked += 1
    assert checked == (len(rows) - 1) * 18


def test_prices_dom_matches_csv(exported):
    p = pd.read_csv(os.path.join(DATA, 'prices_dom.csv'), float_precision='round_trip')
    rows = exported['Prices_dom']
    codes = rows[7]
    data = rows[8:]
    assert len(data) == len(p)
    for r, rec in zip(data, p.to_dict('records')):
        assert r[0] == f"{rec['CountryCode']}{rec['year']}" and r[2] == rec['CountryCode'] and r[4] == rec['year']
        assert r[codes.index('***')] == '***'
        for j, code in enumerate(codes):
            if code in p.columns and code not in ('countryname', 'year'):
                assert r[j] == pytest.approx(rec[code], rel=REL, abs=ABS), code


def _long_to_dict(df, code, var):
    return {(a, b): v for a, b, v in df[[code, var, 'value']].itertuples(index=False)}


@pytest.mark.parametrize('name, csv, var', [('HHSurvey', 'distn_hh_survey.csv', 'variable'),
                                            ('IO_GTAP', 'distn_io_gtap.csv', 'gtap_sector')])
def test_distribution_tables_match_csv(exported, name, csv, var):
    df = pd.read_csv(os.path.join(DATA, csv), float_precision='round_trip')
    rows = exported[name]
    header = rows[0]
    first_value_col = header.index(df[var].iloc[0])
    values = _long_to_dict(df, header[0], var)
    assert len(rows) - 1 == df[header[0]].nunique()
    n = 0
    for r in rows[1:]:
        for j in range(first_value_col, len(header)):
            v = values.get((r[0], header[j]))
            assert (r[j] is None) if v is None or pd.isna(v) else r[j] == pytest.approx(v, rel=REL, abs=ABS)
            n += v is not None
    assert n == len(df)


@pytest.mark.parametrize('name', ['HHSurvey', 'IO_GTAP'])
def test_distribution_tables_reproduce_legacy(exported, name):
    """Egypt HHSurvey and IO_GTAP are real data in both the CSVs and the legacy NoPropData workbook."""
    pyxlsb = pytest.importorskip('pyxlsb')
    if not os.path.exists(ex.LEGACY_XLSB):
        pytest.skip('legacy workbook not present')
    with pyxlsb.open_workbook(ex.LEGACY_XLSB) as wb, wb.get_sheet(name) as sh:
        legacy = [_trim([c.v for c in row]) for row in sh.rows()]
    legacy = [r for r in legacy if r]
    ours = [_trim(r) for r in exported[name]]
    assert len(ours) == len(legacy)
    for a, b in zip(ours, legacy):
        assert len(a) == len(b)
        for x, y in zip(a, b):
            if isinstance(y, float):
                assert x == pytest.approx(y, rel=REL, abs=LEGACY_ABS)
            else:
                assert x == y


def test_workbook_has_only_the_four_tabs(exported):
    assert list(exported) == ['EnergyCons', 'Prices_dom', 'HHSurvey', 'IO_GTAP']


def test_csvs_equal_tabs(exported, tmp_path):
    import csv
    ex.export(DATA, str(tmp_path / 'o.xlsx'), csv_dir=str(tmp_path / 'csv'))
    for name, rows in exported.items():
        with open(tmp_path / 'csv' / f'{name}.csv', encoding='utf-8') as f:
            got = [_trim([None if v == '' else v for v in r]) for r in csv.reader(f)]
        want = [_trim(r) for r in rows]
        assert len(got) == len(want), name
        for g, w in zip(got, want):
            assert len(g) == len(w)
            for x, y in zip(g, w):
                if isinstance(y, (int, float)) and not isinstance(y, bool):
                    assert float(x) == pytest.approx(y, rel=REL, abs=ABS)
                else:
                    assert x == y
