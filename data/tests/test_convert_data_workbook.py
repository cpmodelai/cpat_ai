"""
Tests for data/scripts/convert_data_workbook.py: the CPAT-format sheets have the legacy header rows, and the conversion
Python -> CPAT -> Python gives back every table (the CSVs to the 16 significant digits openpyxl writes; the Python
workbook exactly). The committed CPAT_Data_dummy_cpat.xlsx must equal a fresh conversion of CPAT_Data_dummy_python.xlsx.
"""
import os
import sys

import openpyxl
import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))   # data/
sys.path.insert(0, os.path.join(ROOT, 'scripts'))

# pylint: disable=wrong-import-position
import convert_data_workbook as cv

DATA = os.path.join(ROOT, 'dummy')
PYTHON_XLSX = os.path.join(ROOT, 'workbooks', 'CPAT_Data_dummy_python.xlsx')
CPAT_XLSX = os.path.join(ROOT, 'workbooks', 'CPAT_Data_dummy_cpat.xlsx')


def _norm(v, digits=None):
    if cv._null(v):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(f'{v:.{digits}g}') if digits else (type(v).__name__, v)
    return str(v)


def _sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    out = {n: [list(r) for r in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames}
    wb.close()
    return out


@pytest.fixture(scope='module')
def from_csv(tmp_path_factory):
    out = tmp_path_factory.mktemp('cpat') / 'x_cpat.xlsx'
    entries = cv.to_cpat(DATA, str(out))
    return str(out), entries


@pytest.fixture(scope='module')
def from_python_xlsx(tmp_path_factory):
    out = tmp_path_factory.mktemp('cpat') / 'y_cpat.xlsx'
    cv.to_cpat(PYTHON_XLSX, str(out))
    return str(out)


def test_every_table_is_in_the_cpat_workbook(from_csv):
    _, entries = from_csv
    tables = {e[1] for e in entries}
    assert tables == {f for f in os.listdir(DATA) if f.endswith('.csv')}
    legacy = {e[0] for e in entries if e[2].startswith('legacy')}
    assert legacy == {v[0] for v in cv.CONVERTERS.values()}


def test_headers_match_legacy_layout(from_csv):
    path, _ = from_csv
    sheets = _sheets(path)
    for name, layout in cv.load_layouts().items():
        for r, vals in layout['header'].items():
            got = cv._trim(sheets[name][int(r) - 1]) if int(r) <= len(sheets[name]) else []
            assert got == list(vals), (name, r)


def test_round_trip_from_csv(from_csv, tmp_path):
    """CSV -> CPAT workbook -> CSV: every table, every row, in the same order."""
    path, _ = from_csv
    cv.write_csv_dir(cv.read_cpat_tables(path), str(tmp_path))
    orig = cv.read_python_tables(DATA)
    back = cv.read_python_tables(str(tmp_path))
    assert set(back) == set(orig)
    for name, (header, rows) in orig.items():
        h2, r2 = back[name]
        assert h2 == header, name
        assert [[_norm(v, 16) for v in r] for r in r2] == [[_norm(v, 16) for v in r] for r in rows], name


def test_round_trip_from_python_workbook_is_exact(from_python_xlsx):
    """Python workbook -> CPAT workbook -> Python tables: identical values and types."""
    orig = cv.read_python_tables(PYTHON_XLSX)
    back = cv.read_cpat_tables(from_python_xlsx)
    assert set(back) == set(orig)
    for name, (header, rows) in orig.items():
        h2, r2 = back[name]
        assert h2 == header, name
        assert [[_norm(v) for v in r] for r in r2] == [[_norm(v) for v in r] for r in rows], name


def test_committed_cpat_workbook_is_current(from_python_xlsx):
    a, b = _sheets(CPAT_XLSX), _sheets(from_python_xlsx)
    assert list(a) == list(b)
    for name in a:
        if name != 'README':                            # README carries the build date
            assert a[name] == b[name], name


def test_weo_codes_take_the_legacy_series():
    layout = cv.load_layouts()['WEO']
    tables = {'weo': (['CountryCode', 'Country', 'Indicator', 'Scale', '2020'],
                      [['EGY', 'Egypt', 'ENDA', 'Units', 15.7], ['EGY', 'Egypt', 'NGDP_R', 'Billions', 1.0]])}
    rows, _ = cv.weo_to_cpat(tables, None, dict(layout, sheet='WEO'))
    assert [r[2:4] for r in rows] == [['IMF.STA:ER', 'EGY.XDC_USD.PA_RT.A'], ['IMF.RES:WEO', 'EGY.NGDP_R.A']]
    back = cv.weo_to_python(cv._padded(rows, layout), layout)
    assert list(back['Indicator']) == ['ENDA', 'NGDP_R']
    assert list(back['2020']) == [15.7, 1.0]
